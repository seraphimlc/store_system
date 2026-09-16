"""One-shot: parse 8月全月巡回最终结算.xlsx relevant sheets into JSON cache.
Usage: ./.venv/bin/python tools_dump_recon.py <out.json>
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

import json, os, sys, time
warnings_safe = True
try:
    import warnings
    warnings.filterwarnings("ignore")
    from openpyxl import load_workbook
except Exception as e:
    print("import error", e); sys.exit(1)

base = os.path.expanduser("~/Desktop/万总/")
path = base + "8月全月巡回最终结算.xlsx"
t0 = time.time()
wb = load_workbook(path, read_only=True, data_only=True)
print("sheets:", wb.sheetnames, "load_s", round(time.time() - t0, 1), flush=True)

out = {}
for sn in ["月度总览", "最终有效数据", "全部重复数据"]:
    if sn not in wb.sheetnames:
        print("MISSING", sn); continue
    ws = wb[sn]
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        # keep first 20 cols only, coerce to str-safe
        rows.append([None if c is None else (str(c) if not isinstance(c, (int, float)) else c) for c in r[:20]])
        if i % 20000 == 0:
            print(sn, "row", i, flush=True)
    out[sn] = rows
    print(sn, "done rows=", len(rows), "s=", round(time.time() - t0, 1), flush=True)
wb.close()
with open(sys.argv[1], "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False)
print("saved", sys.argv[1], "total_s", round(time.time() - t0, 1))
