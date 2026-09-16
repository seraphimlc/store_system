# -*- coding: utf-8 -*-
"""测试助手：tmp_path 内现造 xlsx；坏 dimension 手术。"""
import os
import zipfile
from openpyxl import Workbook


def write_workbook(path, sheets):
    """sheets = [(sheet_name, header_rows, data_rows)]。
    header_rows: 列表的列表（一行一个列表）；若 2 行则为双表头（第 1 行分组、第 2 行字段）。
    data_rows: 列表的列表，顺序对应最后一行表头。
    返回 path。
    """
    wb = Workbook()
    wb.remove(wb.active)
    for name, header_rows, data_rows in sheets:
        ws = wb.create_sheet(title=name)
        for r_i, hrow in enumerate(header_rows, start=1):
            for c_i, v in enumerate(hrow, start=1):
                ws.cell(r_i, c_i, v)
        for r_i, drow in enumerate(data_rows, start=len(header_rows) + 1):
            for c_i, v in enumerate(drow, start=1):
                ws.cell(r_i, c_i, v)
    wb.save(path)
    return path


def corrupt_dimension(path):
    """把 sheet1.xml 的 <dimension ref="..."> 改为 ref="A1"（模拟坏 dimension 文件 2/3）。"""
    import re
    tmp = path + ".zip"
    os.rename(path, tmp)
    zin = zipfile.ZipFile(tmp, "r")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.endswith("xl/worksheets/sheet1.xml"):
                text = data.decode("utf-8")
                text = re.sub(r'<dimension ref="[^"]*"/>', '<dimension ref="A1"/>', text, count=1)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    zin.close()
    os.remove(tmp)
    return path


# ---- 导入自动建员工账号 测试用：宽表上传文件 ----
WIDE_H1 = ["Store Information"] + [""] * 13
WIDE_H2 = ["Store ID", "Store Name-Local", "Store Name-English", "Country/Region",
           "city", "Create Time", "Modified Time", "First Visit Time",
           "Review Completion Time", "Submitter", "ISO Company", "Record ID",
           "A+ POSM Visible", "Existing A+ POSM Visible", "Deploy New A+POSM"]


def wide_xlsx_bytes(tmp_path, submitters, name="wide"):
    """submitters=[(姓名_编号文本,...)] 每人生成一行宽表数据。"""
    rows = []
    for i, sub in enumerate(submitters, start=1):
        rows.append([
            f"ID{i}", f"店{i}", "", "Japan", "Tokyo",
            "2026-07-01 08:00:00", "2026-07-05 14:21:49",
            "2026-07-01 08:30:00", "2026-07-05 15:00:00",
            sub, "MarsNavi", f"R{i}", "YES", "YES", "YES"])
    p = str(tmp_path / f"{name}.xlsx")
    write_workbook(p, [("STORE_TASK_EXCEL_SHEET", [WIDE_H1, WIDE_H2], rows)])
    with open(p, "rb") as f:
        return f.read()
