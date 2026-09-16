# -*- coding: utf-8 -*-
"""生成仓库内合成 xlsx fixtures（运行: ./.venv/bin/python tests/fixtures/gen_fixtures.py）。"""
import os, re, sys, zipfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))
H9 = ["Store ID", "Store Name-Local", "Store Name-English", "Modified Time",
      "Submitter", "Record ID", "A+ POSM Visible", "Existing A+ POSM", "NEW A+ POSM"]
H9_no_rec = ["Store ID", "Store Name-Local", "Store Name-English", "Modified Time",
             "First Visit Time", "Submitter", "A+ POSM Visible",
             "Existing A+ POSM", "NEW A+ POSM"]
W1 = ["Store Information"] + [""] * 13
W2 = ["Store ID", "Store Name-Local", "Store Name-English", "Country/Region",
      "city", "Create Time", "Modified Time", "First Visit Time",
      "Review Completion Time", "Submitter", "ISO Company", "Record ID",
      "A+ POSM Visible", "Existing A+ POSM Visible", "Deploy New A+POSM"]

def new_wb():
    wb = Workbook(); wb.remove(wb.active); return wb

def save(wb, name):
    wb.save(os.path.join(HERE, name)); print("wrote", name)

def corrupt_dimension(path):
    tmp = path + ".t"; os.rename(path, tmp)
    zin = zipfile.ZipFile(tmp)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename.endswith("xl/worksheets/sheet1.xml"):
                data = re.sub(rb'<dimension ref="[^"]*"/>', b'<dimension ref="A1"/>', data, count=1)
            zo.writestr(it, data)
    zin.close(); os.remove(tmp)

# wide50
wb = new_wb(); ws = wb.create_sheet("STORE_TASK_EXCEL_SHEET ")
ws.append(W1); ws.append(W2)
for i in range(1, 6):
    ws.append([f"ID00{i}", f"店{i}", f"en{i}", "Japan", "Tokyo",
               f"2026-07-0{i} 08:00:00", f"2026-07-0{i} 09:0{i}:00",
               f"2026-07-0{i} 08:30:00", f"2026-07-0{i} 10:00:00",
               f"吴海峰(2188240620009615)", "MarsNavi", f"R{i}",
               "YES", "YES", "YES"])
save(wb, "wide50.xlsx")

# flat9 variants
wb = new_wb(); ws = wb.create_sheet("STORE_TASK_EXCEL_SHEET ")
ws.append(H9)
ws.append(["ID1", "A店", "", "2026-07-05 14:21:49", "潘春晓(2188240606697508)",
           "R11", "YES", "YES", "YES"])
save(wb, "flat9_record_id.xlsx")

wb = new_wb(); ws = wb.create_sheet("STORE_TASK_EXCEL_SHEET ")
ws.append(H9_no_rec)
ws.append(["ID2", "B店", "", "2026-07-05 14:10:49", "2026-07-05 14:10:36",
           "陈嘉溢(2188240606650879)", "NO", "NO", ""])
save(wb, "flat9_no_record_id.xlsx")

# broken dimension（30 行宽表 + 坏 dimension）
wb = new_wb(); ws = wb.create_sheet("STORE_TASK_EXCEL_SHEET ")
ws.append(W1); ws.append(W2)
for i in range(1, 31):
    ws.append([f"BD{i:03d}", f"店{i}", "", "Japan", "Tokyo",
               f"2026-07-{i%28+1:02d} 08:00:00", f"2026-07-{i%28+1:02d} 09:00:00",
               f"2026-07-{i%28+1:02d} 08:30:00", f"2026-07-{i%28+1:02d} 10:00:00",
               "新井圭史(2188240606737868)", "MarsNavi", f"R{i}",
               "NO", "NO", ""])
p = os.path.join(HERE, "broken_dimension.xlsx")
wb.save(p); corrupt_dimension(p); print("wrote broken_dimension.xlsx")

# yan-like：主表 + Sheet5
wb = new_wb(); ws = wb.create_sheet("STORE_TASK_EXCEL_SHEET ")
ws.append(H9)
for i in range(1, 21):
    ws.append([f"Y{i:03d}", f"店{i}", "", f"2026-07-{i:02d} 09:00:00",
               "小川逸(2188240606634082)", f"RY{i}", "YES", "YES", "YES"])
ws5 = wb.create_sheet("Sheet5")
ws5.append(["作业分"] + H9)
for i in range(1, 21):
    ws5.append(["", f"Y{i:03d}", f"店{i}", "", f"2026-07-{i:02d} 09:00:00",
                "小川逸(2188240606634082)", f"RY{i}", "YES", "YES", "YES"])
save(wb, "yan_like.xlsx")
