# -*- coding: utf-8 -*-
"""线上导入三份「执行计划」→ 建成 E陈嘉溢队 / F甘子杰队 的任务。

**为什么要在临时副本上改「负责队」列**：三份文件里这一列全是「待分配」，
而导入服务的口径是"**只按文件里的负责队列**、系统不猜"（用户 2026-10-16 定稿）
→ 不改就是整份拒收 / 静默跳过。派给谁由**用户 2026-10-17 当面指定**：
千叶B、千叶C → `E陈嘉溢队`；埼玉A → `F甘子杰队`。原始文件不动，只改副本，
并且**原始与改后两份都用 `save_blob()` 归档**（审计能看出改了什么）。

对不上的 2 行按服务口径勾选跳过（不勾就一条都不写）：
  · 千叶C 行22 `幕張豊砂`（JR 京葉線）→ 库里**没有**这个站（2023 新站）
  · 埼玉A 行34 `獨協大学前〈草加松原〉` → 库里是 `獨協大学前`（id 1727）同站异写

用法（容器内）：`APPLY=1 python -` 才真写；不设 APPLY 只预览。
"""
import io
import os

from openpyxl import load_workbook

from app.db import SessionLocal
from app.models import BdTask, User
from app.services import bd_task_import as TI

JOBS = [
    ("千叶B线", "/tmp/tp/b.xlsx", "E陈嘉溢队", []),
    ("千叶C线", "/tmp/tp/c.xlsx", "E陈嘉溢队", [22]),
    ("埼玉A线", "/tmp/tp/a.xlsx", "F甘子杰队", [34]),
]
DRY = os.environ.get("APPLY") != "1"

db = SessionLocal()
admin = db.query(User).filter(User.username == "admin").first()
print("执行账号 =", getattr(admin, "username", None),
      "| 模式 =", "干跑(不写)" if DRY else "**真写**")

before = {t.id: t.team_id for t in db.query(BdTask).all()}
print("写前任务总数 =", len(before))

for label, path, team, skips in JOBS:
    print("=" * 88)
    print(label, path, "→", team, "| 跳过行:", skips or "无")
    raw_bytes = open(path, "rb").read()               # 原始文件先归档
    sha_raw, _p_raw = TI.save_blob(raw_bytes)

    layout = TI.detect_layout(path)
    cols, hrow = layout["cols"], layout["header_row"]
    wb = load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    got = [ws.cell(hrow, cols[k] + 1).value for k in ("line", "station", "team")]
    print("  表头行=%s  列(0基) line=%s station=%s team=%s → 表头文字 %s"
          % (hrow, cols["line"], cols["station"], cols["team"], got))
    assert "负责" in str(got[2]), "认到的负责队列是 %r，别写！" % (got[2],)

    tc, sc = cols["team"] + 1, cols["station"] + 1
    n = 0
    for r in range(hrow + 1, ws.max_row + 1):
        if ws.cell(r, sc).value in (None, ""):
            continue
        ws.cell(r, tc).value = team
        n += 1
    bio = io.BytesIO()
    wb.save(bio)
    sha_new, p_new = TI.save_blob(bio.getvalue())     # 改后文件也归档
    print("  改 %d 行负责队；原始存档 %s… 改后存档 %s…"
          % (n, sha_raw[:12], sha_new[:12]))

    plan = TI.build_plan(db, p_new)
    s = plan["summary"]
    print("  预览：读到 %d 行 → 新建 %d 行 / **%d 个站**；已有 %d；跳过 %d；对不上 %d"
          % (s["rows"], s["create_rows"], s["create_stations"],
             s["existing"], s["skipped"], s["bad"]))
    print("  认到的队名:", s["teams"])
    for r in plan["rows"]:
        if r.decision == "bad":
            print("    ✗ 行%d %s | %s | 近似 %s"
                  % (r.row_no, r.line, r.station_raw, "、".join(r.near or [])))
    if DRY:
        continue
    res = TI.apply_plan(db, p_new, skip_rows=skips, actor_user=admin)
    print("  >>> 新建 %d 个任务（%s）" % (
        res["created"],
        "、".join("%s %d 个" % (t["team"], t["created"]) for t in res["by_team"])))
    print("  >>> 建时又跳过(已有) %d；行级跳过 %d；已有未动 %d"
          % (res["skipped"], res["skipped_rows"], res["existing_untouched"]))

after = {t.id: t.team_id for t in db.query(BdTask).all()}
changed = [(i, before[i], after[i]) for i in before if after.get(i) != before[i]]
print("=" * 88)
print("任务总数 写前/写后 = %d / %d" % (len(before), len(after)))
print("已有任务归属被改动 =", len(changed), changed[:5])
for tid, tname in ((5, "E陈嘉溢队"), (3, "F甘子杰队")):
    n = db.query(BdTask).filter(BdTask.team_id == tid).count()
    print("  %s 现任务数 = %d" % (tname, n))
db.close()
