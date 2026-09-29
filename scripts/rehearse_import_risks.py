# -*- coding: utf-8 -*-
"""上线前的导入风险演练（第 2 组）：新月份 / 坏文件 / 布局变化。

1) 新月份导入：把真实文件日期整体 +45 天（→ 10 月）→ 验证"全新周期"能完整导入并自动生成报告
2) 坏文件：伪造一个损坏的 xlsx → 验证系统干净报错、不破坏已有数据、不影响后续上传
3) 布局变化：把表头下移一行 + 增加一列 → 验证解析器如何应对（客户文件格式微调是最现实的风险）

全程 AI mock、只动临时库。用法：./.venv/bin/python scripts/rehearse_import_risks.py
"""
import io
import os
import re
import sys
import time
from datetime import datetime, timedelta

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
SRC = "/Users/liuchang/Desktop/万总/0901-0915MarsNavi store visit.xlsx"
DB = "/tmp/rehearse_risk.db"
SHEET = "STORE_TASK_EXCEL_SHEET "


def log(m):
    print("[%s] %s" % (datetime.now().strftime("%H:%M:%S"), m), flush=True)


def derive(path, shift_days=0, header_row_shift=0, extra_col=False):
    """由真实文件派生变体：位移日期 / 下移表头 / 加一列。"""
    from openpyxl import load_workbook
    wb = load_workbook(SRC)
    ws = wb[SHEET]
    if shift_days:
        n = 0
        for row in ws.iter_rows(min_row=3, min_col=3, max_col=3):
            c = row[0]
            v = c.value
            if isinstance(v, datetime):
                c.value = v + timedelta(days=shift_days); n += 1
            elif v is not None:
                m = re.match(r"^(\d{4})-(\d{2})-(\d{2})(.*)$", str(v).strip())
                if m:
                    d = datetime(int(m.group(1)), int(m.group(2)), int(m.group(3))) \
                        + timedelta(days=shift_days)
                    c.value = d.strftime("%Y-%m-%d") + m.group(4); n += 1
        log("  派生：日期位移 %+d 天（%d 行）" % (shift_days, n))
    if extra_col:
        ws.insert_cols(8)
        ws.cell(row=2, column=8, value="Extra Col")
        log("  派生：新增第 8 列（模拟客户加列）")
    if header_row_shift:
        ws.insert_rows(1, header_row_shift)
        log("  派生：表头下移 %d 行" % header_row_shift)
    wb.save(path)
    wb.close()
    return path


def main():
    for p in (DB, DB + "-journal"):
        if os.path.exists(p):
            os.remove(p)
    os.environ["DATABASE_URL"] = "sqlite:///%s" % DB
    os.environ["VISIT_REPORT_AI"] = "1"

    from sqlalchemy import create_engine
    import app.db as appdb
    import app.models  # noqa: F401
    from app.db import Base
    Base.metadata.create_all(create_engine("sqlite:///%s" % DB, future=True))

    from app.services import ai_chat, report_ai
    ai_chat.configured = lambda: True
    ai_chat.chat = lambda *a, **k: (
        ('{"overall_comment":"彩排","accuracy_notes":"n","per_person":{}}',
         {"total_tokens": 1}) if k.get("return_usage") else
        '{"overall_comment":"彩排","accuracy_notes":"n","per_person":{}}')
    report_ai.build_prompt = lambda res, lang: "彩排"

    from app.auth import hash_password
    from app.models import ImportFile, User
    db = appdb.SessionLocal()
    db.add(User(username="admin", password_hash=hash_password("rehearse123"),
                display_name="彩排", role="admin", is_active=True))
    db.commit(); db.close()

    from fastapi.testclient import TestClient
    from app.main import create_app
    c = TestClient(create_app())
    c.post("/login", data={"username": "admin", "password": "rehearse123"},
           follow_redirects=False)

    def upload(path, label):
        t0 = time.time()
        with open(path, "rb") as fh:
            r = c.post("/files/upload",
                       data={"csrf_token": _csrf(c), "_ft": _token(c)},
                       files=[("files", (os.path.basename(path), fh.read(),
                                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))],
                       follow_redirects=False)
        log("%s → 上传 HTTP %s（%.1fs）" % (label, r.status_code, time.time() - t0))
        db = appdb.SessionLocal()
        imp = db.query(ImportFile).order_by(ImportFile.id.desc()).first()
        info = (imp.id, imp.status, str(imp.errors or [])[:90]) if imp else None
        db.close()
        return r, info

    # ---------- 1) 新月份（+45 天 → 10 月）----------
    log("① 新月份导入演练")
    f1 = derive("/tmp/oct.xlsx", shift_days=45)
    r, info = upload(f1, "   10 月文件")
    if info and info[1] == "parsed":
        t0 = time.time()
        fin = c.post("/files/%d/finalize" % info[0],
                     data={"_ft": _token(c), "csrf_token": _csrf(c)},
                     follow_redirects=False)
        log("   入表 HTTP %s（%.1fs）" % (fin.status_code, time.time() - t0))
    else:
        log("   解析未成功：%s" % (info,))
    time.sleep(2)

    # ---------- 2) 坏文件 ----------
    log("② 坏文件演练（伪造损坏 xlsx）")
    with open("/tmp/broken.xlsx", "wb") as fh:
        fh.write(b"PK\x03\x04 this is not a real xlsx")
    r, info = upload("/tmp/broken.xlsx", "   坏文件")
    log("   结果：%s（期望：干净报错，不是 500）" % (info,))

    # ---------- 3) 布局变化 ----------
    log("③ 布局变化演练（表头下移 1 行 + 加一列）")
    f3 = derive("/tmp/layout_changed.xlsx", shift_days=75,
                header_row_shift=1, extra_col=True)
    r, info = upload(f3, "   变体文件")
    if info and info[1] == "parsed":
        fin = c.post("/files/%d/finalize" % info[0],
                     data={"_ft": _token(c), "csrf_token": _csrf(c)},
                     follow_redirects=False)
        log("   入表 HTTP %s（居然能自适应？看下面汇总）" % fin.status_code)

    # ---------- 汇总 ----------
    from app.models import FormalRecord, PersonDailyStat, StaffReportAnalysis
    from sqlalchemy import func
    db = appdb.SessionLocal()
    print("\n===== 汇总 =====")
    for imp in db.query(ImportFile).order_by(ImportFile.id):
        print("  #%s %-26s status=%-8s %s"
              % (imp.id, imp.file_name[:26], imp.status,
                 str(imp.errors or [])[:60]))
    lo, hi = db.query(func.min(FormalRecord.japan_date),
                      func.max(FormalRecord.japan_date)).first()
    print("  正式表：%s ~ %s / %d 行" % (lo, hi, db.query(FormalRecord).count()))
    for d, n in (db.query(FormalRecord.japan_date, func.count())
                 .group_by(FormalRecord.japan_date)
                 .order_by(FormalRecord.japan_date).all())[:2]:
        print("    %s: %d 行" % (d, n))
    print("  person_daily_stats %d 行 / 报告 %d 份"
          % (db.query(PersonDailyStat).count(),
             db.query(StaffReportAnalysis).count()))
    for a in db.query(StaffReportAnalysis).order_by(StaffReportAnalysis.id):
        print("    报告 #%s %s~%s %s" % (a.id, a.period_start, a.period_end,
                                         a.status))
    db.close()
    return 0


def _token(c):
    m = re.search(r'name="_ft" value="([^"]+)"', c.get("/files").text)
    return m.group(1) if m else ""


def _csrf(c):
    from app.auth import SESSION_COOKIE, read_session_token
    return (read_session_token(c.cookies.get(SESSION_COOKIE)) or {}).get("csrf", "")


if __name__ == "__main__":
    sys.exit(main())
