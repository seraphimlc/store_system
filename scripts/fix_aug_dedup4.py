# -*- coding: utf-8 -*-
"""修复 8 月 4 家"该去重没去重"的店（支付宝账单只确认 1 个 store_id，我们算了 2 条）。

支付宝真值判定（Alipay8月结算数据.xlsx）→ 每组删支付宝无的那条 formal，
并把它对应的 raw 标记为 master_late（避免未来 rebuild 复活），然后重算 8 月。
"""
import os
import sys
import unicodedata
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from collections import defaultdict
from openpyxl import load_workbook
import app.db as appdb
from app.models import FormalRecord, RawRecord
from app.services import perf, period


def norm(s):
    s = unicodedata.normalize("NFKC", s or "")
    return "".join(s.casefold().split())


ALI_PATH = "/Users/liuchang/Desktop/万总/Alipay8月结算数据.xlsx"
MONTH = "2026-08"

db = appdb.SessionLocal()
# 支付宝: norm 店名 -> {store_id}
wb = load_workbook(ALI_PATH, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
ali = defaultdict(set)
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0 or not row or row[4] is None:
        continue
    ali[norm(str(row[4] or ""))].add(str(row[3] or ""))
wb.close()

# 系统 8 月 formal，按 norm 店名分组
fr = db.query(FormalRecord).filter(
    FormalRecord.japan_date >= f"{MONTH}-01",
    FormalRecord.japan_date < "2026-09-01").all()
rawmap = {r.id: r for r in db.query(RawRecord).all()}
groups = defaultdict(list)
for f in fr:
    rr = rawmap.get(f.raw_record_id)
    nm = (rr.store_name_local_raw or "") if rr else ""
    groups[norm(nm)].append(f)

# 真漏判重：norm 同名、多 store_id、支付宝只确认 1 个 store
to_del = []
for sn, flist in groups.items():
    sids = {f.store_id_raw for f in flist}
    if len(sids) < 2 or sn not in ali or len(ali[sn]) != 1:
        continue
    keep_sid = list(ali[sn])[0]
    for f in flist:
        if f.store_id_raw not in ali[sn]:
            to_del.append(f)   # 支付宝无此 store → 删
print("待删除 formal:", len(to_del))
for f in to_del:
    rr = rawmap.get(f.raw_record_id)
    print(f"  formal#{f.id} store={f.store_id_raw} 「{(rr.store_name_local_raw or '') if rr else '?'}」 {f.japan_date} {f.points}点")

# 执行：删 formal + 标记 raw
raw_updated = 0
for f in to_del:
    rr = rawmap.get(f.raw_record_id)
    if rr is not None:
        rr.clean_status = "master_late"
        rr.filter_reason = "master_late"
        rr.confirm_state = "auto_ok"
        raw_updated += 1
    db.delete(f)
db.commit()
print("标记 raw:", raw_updated, "| 已删 formal:", len(to_del))

# 重算 8 月：统计表 → 月绩效 → 薪资找平
perf.sync_month_stats(db, MONTH)
perf.sync_month_perf(db, MONTH)
period.sync_period_table(db, MONTH)

# 新基准
fr2 = db.query(FormalRecord).filter(
    FormalRecord.japan_date >= f"{MONTH}-01",
    FormalRecord.japan_date < "2026-09-01").all()
mp = perf.month_perf(db, MONTH)
print("\n修复后 8 月基准:")
print("  formal:", len(fr2), "点数:", sum(x.points for x in fr2))
print("  月绩效:", len(mp), "人 工资:", sum(x["salary"] for x in mp),
      "奖金:", sum((x["points"] // 68) * 3000 for x in mp))
rows = period.period_rows(db, MONTH)
print("  薪资找平:", len(rows), "行 分期已发:",
      sum(r["half1_amt"] + r["half2_amt"] for r in rows),
      "对账金额:", sum(r["settle_amt"] for r in rows))
db.close()