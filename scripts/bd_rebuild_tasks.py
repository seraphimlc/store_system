# -*- coding: utf-8 -*-
"""任务重建：**删掉现有全部任务**，按 `万总/team_task/` 最新 6 份分配重建。

用户 2026-10-07："任务重新派分，之前的515条分组任务全部删除。使用"万总/team_task"下最新的任务分配。"

口径（沿用 2026-10-05 定稿，与 `scripts/bd_init_515.py` 同一套）：
- **1 个物理车站（place）= 1 个任务**；只到**队伍**，**担当由队长分**（新文件里担当列全空）
- 站点匹配优先级：① 站名+线路（严格）② 站名+线路（宽松：去「N号線」前后缀包含）
  ③ 站名唯一时退化为按站名 ④ 都不行 → **列入待人工，不猜**
- **队名必须已在库里存在**（绝不新建/改名团队 —— 线上 6 个队是用户手工配的）
- 分配日期默认**重建当天**（= 重新派分的日子）；`--use-file-date` 则用文件里的日期

安全措施：
- 默认 **dry-run**（只打印，不写库）
- `--apply` 时先把现有任务/担当/进展**快照成 CSV**（`data/backup/`，可人工恢复）
- 删任务时**连带删掉指向任务的站内消息**（否则消息点进去是 404）
- `bd_log`（审计日志）**保留**（历史就是历史）

用法：
    DATABASE_URL="sqlite:///./store_settle_live.db" ./.venv/bin/python \
        scripts/bd_rebuild_tasks.py --src "/Users/liuchang/Desktop/万总/team_task"
    # 真写：加 --apply
"""
from __future__ import annotations

import argparse
import collections
import csv
import glob
import io
import os
import sys
import unicodedata
from datetime import datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

CSV_HEADER = ["任务id", "站点", "团队", "状态", "进展%", "担当", "分配日期",
              "开始日", "完成日", "店铺数"]


def norm(s) -> str:
    return unicodedata.normalize("NFKC", str(s or "")).strip()


def read_src(src: str):
    """读 6 份 〈队名〉_车站任务.xlsx → [(队名, 线路, 站点, 文件里的分配日期)]"""
    from openpyxl import load_workbook
    rows = []
    files = sorted(glob.glob(os.path.join(src, "*.xlsx")))
    for f in files:
        tname = norm(os.path.basename(f).split("_")[0])
        ws = load_workbook(f, data_only=True, read_only=True).active
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or not row:
                continue
            line = norm(row[0]) if len(row) > 0 else ""
            st = norm(row[1]) if len(row) > 1 else ""
            d = norm(row[5])[:10] if len(row) > 5 else ""
            if st:
                rows.append((tname, line, st, d))
    return files, rows


def build_place_index(db):
    """线路名 → {站名: place_id}；站名 → {place_id}（与 bd_init_515 同口径）。"""
    from sqlalchemy import text
    from app.services import bd_lines
    by_line, by_name = {}, collections.defaultdict(set)
    q = db.execute(text("""SELECT s.place_id, s.name, l.name, l.operator_short
            FROM bd_station s JOIN bd_line l ON l.id = s.line_id
            WHERE s.place_id IS NOT NULL""")).all()
    for pid, nm, lname, lshort in q:
        by_line.setdefault(bd_lines.norm_line(lname), {})[nm] = pid
        by_line.setdefault(bd_lines.norm_line(lshort + lname), {})[nm] = pid
        by_name[nm].add(pid)
    return by_line, by_name


def match_place(by_line, by_name, line_raw: str, name: str):
    """→ (place_id, 权重) 或 (None, 原因)。线路列可能是「A、B、C」多线。"""
    from app.services import bd_lines
    parts = [p for p in (line_raw or "").replace("，", "、").split("、") if norm(p)]
    for base in parts:                                   # ① 严格：挨条线路试
        hit = by_line.get(bd_lines.norm_line(base), {}).get(name)
        if hit:
            return hit, "strict"
    for base in parts:                                   # ② 宽松
        core = bd_lines.norm_line(base).lstrip("0123456789").replace("号線", "")
        found = None
        for k, m in by_line.items():
            kk = k.lstrip("0123456789").replace("号線", "")
            if name in m and (kk == core or (len(core) >= 3
                                             and (core in kk or kk in core))):
                if found and found != m[name]:
                    return None, "宽松匹配命中多个"
                found = m[name]
        if found:
            return found, "loose"
    cand = by_name.get(name) or set()                    # ③ 站名唯一
    if len(cand) == 1:
        return next(iter(cand)), "name-only"
    if len(cand) > 1:
        return None, "同名多站(%d)" % len(cand)
    return None, "车站表里找不到"


def snapshot(db, path: str) -> int:
    """把现有任务（含担当）快照成 CSV —— 万一要人工恢复。"""
    from app.models import BdTask, BdTaskAssign
    from app.services import bd_tasks
    assigns = collections.defaultdict(list)
    for a in db.query(BdTaskAssign).all():
        assigns[a.task_id].append(a.person_code)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    n = 0
    with io.open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(CSV_HEADER)
        for t in db.query(BdTask).order_by(BdTask.id).all():
            w.writerow([t.id, bd_tasks._station_name(db, t),
                        (db.get(__import__("app.models", fromlist=["BdTeam"]).BdTeam,
                                t.team_id).name if t.team_id else ""),
                        t.state, t.pct or 0, "、".join(assigns.get(t.id, [])),
                        t.assign_date or "", t.start_date or "", t.done_date or "",
                        "" if t.store_count is None else t.store_count])
            n += 1
    return n


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True, help="含 6 份 〈队名〉_车站任务.xlsx 的目录")
    ap.add_argument("--apply", action="store_true", help="真写库（默认只试算）")
    ap.add_argument("--by", default="rebuild:2026-10-07", help="创建人标记")
    ap.add_argument("--assign-date", default="", help="统一指定分配日期 YYYY-MM-DD")
    ap.add_argument("--use-file-date", action="store_true",
                    help="用文件里的分配日期（默认用重建当天）")
    ap.add_argument("--snapshot-dir", default=os.path.join(ROOT, "data", "backup"))
    args = ap.parse_args()

    from app.db import SessionLocal
    from app.models import (BdMessage, BdMessageRecipient, BdTask, BdTaskAssign,
                            BdTaskProgress, BdTeam)
    from app.services import bd_tasks
    from app.services.date_plan import jst_today

    files, rows = read_src(args.src)
    if not files:
        print("✗ 目录里没有 xlsx：%s" % args.src)
        return 2
    print("源文件 %d 份，合计 %d 行" % (len(files), len(rows)))

    db = SessionLocal()
    by_line, by_name = build_place_index(db)

    # ---- 队名必须已存在 ----
    teams = {norm(t.name): t for t in db.query(BdTeam).all()}
    src_teams = sorted({r[0] for r in rows})
    missing = [t for t in src_teams if t not in teams]
    print("\n① 队名检查（库里 %d 个队）" % len(teams))
    for t in src_teams:
        n = sum(1 for r in rows if r[0] == t)
        print("   %s %-10s %3d 行" % ("✅" if t in teams else "❌", t, n))
    if missing:
        print("   ❌ 这些队库里没有，**绝不新建**：%s" % "、".join(missing))
        print("   （先在系统里把队伍建好/改名一致，再重跑）")
        db.close()
        return 3

    # ---- 站点匹配 ----
    plan = collections.defaultdict(list)              # 队名 → [place_id]
    how = collections.Counter()
    unresolved = []
    dup_in_team = collections.Counter()
    for tname, line, st, _d in rows:
        pid, w = match_place(by_line, by_name, line, st)
        if pid is None:
            unresolved.append((tname, line, st, w))
            continue
        dup_in_team[(tname, pid)] += 1
        plan[tname].append(pid)
        how[w] += 1

    print("\n② 站点匹配")
    print("   ① 站名+线路 严格: %d" % how["strict"])
    print("   ② 站名+线路 宽松: %d" % how["loose"])
    print("   ③ 站名唯一      : %d" % how["name-only"])
    print("   ④ 对不上（不猜）: %d" % len(unresolved))
    for x in unresolved[:15]:
        print("        %s / %s（%s队）← %s" % (x[1], x[2], x[0], x[3]))
    rep = {k: v for k, v in dup_in_team.items() if v > 1}
    cross = collections.Counter()
    for tname, pids in plan.items():
        for pid in set(pids):
            cross[pid] += 1
    both = {k: v for k, v in cross.items() if v > 1}
    print("   同一队内重复站点: %d ｜ **跨队重复站点**: %d" % (len(rep), len(both)))

    print("\n③ 按队（队 → 物理车站数）")
    total = 0
    for tname in src_teams:
        uniq = list(dict.fromkeys(plan.get(tname, [])))
        total += len(uniq)
        print("   %-10s %3d 行 → %3d 个任务%s"
              % (tname, len(plan.get(tname, [])), len(uniq),
                 "" if tname in teams else "  ⚠️ 跳过"))
    print("   合计 **%d 个任务**" % total)

    # ---- 现有任务 ----
    tasks = db.query(BdTask).all()
    n_assign = db.query(BdTaskAssign).count()
    n_prog = db.query(BdTaskProgress).count()
    n_msg = (db.query(BdMessage).filter(BdMessage.ref_type == "task").count())
    with_assign = len({a.task_id for a in db.query(BdTaskAssign).all()})
    pct_gt0 = sum(1 for t in tasks if (t.pct or 0) > 0)
    done = sum(1 for t in tasks if t.state == "done")
    will = {(teams[r[0]].id, pid) for r in rows if r[0] in teams
            for pid in [match_place(by_line, by_name, r[1], r[2])[0]] if pid}
    src_place = {pid for _tid, pid in will}
    old_place = set()
    by_place = {}
    for t in tasks:
        nm = norm(bd_tasks._station_name(db, t))
        old_place.add(nm)
    dropped = [t for t in tasks if norm(bd_tasks._station_name(db, t)) not in
               {norm(r[2]) for r in rows}]
    # ---- ③b 换队统计（这次"重新派分"到底改了什么）----
    tname_by_id = {t.id: t.name for t in db.query(BdTeam).all()}
    new_team_of = {norm(r[2]): r[0] for r in rows}
    moves = collections.Counter()
    for t in tasks:
        nm = norm(bd_tasks._station_name(db, t))
        if nm not in new_team_of:
            continue
        old_t = tname_by_id.get(t.team_id, "（未派队）")
        if old_t != new_team_of[nm]:
            moves[(old_t, new_team_of[nm])] += 1
    n_move = sum(moves.values())
    n_same = sum(1 for t in tasks
                 if norm(bd_tasks._station_name(db, t)) in new_team_of
                 and tname_by_id.get(t.team_id, "（未派队）")
                 == new_team_of[norm(bd_tasks._station_name(db, t))])
    print("\n③b 重新派分的变化：队没变 %d 个 ｜ **换队 %d 个**" % (n_same, n_move))
    for (o, n), c in moves.most_common():
        print("     %-10s → %-10s %3d" % (o, n, c))

    print("\n④ 现有任务 %d 条 → 全删" % len(tasks))
    print("   带担当 %d 条（%d 条担当记录）｜ 有进展 %d 条（%d 条进展记录）｜ 已完成 %d 条"
          % (with_assign, n_assign, pct_gt0, n_prog, done))
    print("   指向任务的站内消息 %d 条会一起删（否则点进去 404）" % n_msg)
    print("   站点不在新文件里、重建后**会消失**的：%d 条" % len(dropped))
    for t in dropped[:15]:
        print("        %s" % bd_tasks._station_name(db, t))
    if len(dropped) > 15:
        print("        … 还有 %d 条" % (len(dropped) - 15))

    if not args.apply:
        print("\n（dry-run：没写库。加 --apply 才真删真建）")
        db.close()
        return 0

    # ---- 真写 ----
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    snap = os.path.join(args.snapshot_dir, "bd_tasks_before_rebuild_%s.csv" % ts)
    n = snapshot(db, snap)
    print("\n⑤ 已快照 %d 条现有任务 → %s" % (n, snap))

    n_msg_del = (db.query(BdMessage).filter(BdMessage.ref_type == "task").count())
    if n_msg_del:
        ids = [mid for (mid,) in db.query(BdMessage.id)
               .filter(BdMessage.ref_type == "task").all()]
        db.query(BdMessageRecipient).filter(
            BdMessageRecipient.message_id.in_(ids)).delete(synchronize_session=False)
        db.query(BdMessage).filter(BdMessage.id.in_(ids)).delete(
            synchronize_session=False)
    db.query(BdTaskProgress).delete(synchronize_session=False)
    db.query(BdTaskAssign).delete(synchronize_session=False)
    db.query(BdTask).delete(synchronize_session=False)
    db.commit()
    print("   已删除：任务 %d ｜ 担当 %d ｜ 进展 %d ｜ 消息 %d"
          % (len(tasks), n_assign, n_prog, n_msg_del))

    print("\n⑥ 开始重建")
    created = skipped = 0
    for tname in src_teams:
        uniq = list(dict.fromkeys(plan.get(tname, [])))
        if not uniq:
            continue
        if args.assign_date:
            d = datetime.strptime(args.assign_date, "%Y-%m-%d").date()
        elif args.use_file_date:
            dates = [r[3] for r in rows if r[0] == tname and r[3]]
            d = (datetime.strptime(collections.Counter(dates).most_common(1)[0][0],
                                   "%Y-%m-%d").date() if dates else jst_today())
        else:
            d = jst_today()
        r = bd_tasks.create_tasks_for_places(db, uniq, by=args.by,
                                             team_id=teams[tname].id,
                                             assign_date=d)
        created += r["created"]
        skipped += r["skipped"]
        print("   %-10s 新建 %3d / 跳过 %3d（分配日期 %s）"
              % (tname, r["created"], r["skipped"], d))

    after = db.query(BdTask).count()
    print("\n合计：**新建 %d 个任务**（跳过 %d），库里现有任务 %d 条"
          % (created, skipped, after))
    if after == created:
        print("✅ 数量吻合")
    db.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
