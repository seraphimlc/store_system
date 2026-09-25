# -*- coding: utf-8 -*-
"""P1 导出重构等价性基线：capture（重构前跑一次）→ compare（重构后跑一次）。

用法（cwd=项目根，用主 venv）：
    ./.venv/bin/python scripts/p1_export_baseline.py capture
    ./.venv/bin/python scripts/p1_export_baseline.py compare

capture：造 fixture 库（正式表/月绩效/分期偏差/对账任务）→ 走 4 个导出路由
（重构前的代码）→ 把每个工作簿的「sheet 名 + 逐单元格值」快照写入
mcp_service/tests/data/export_baseline.json，供 test_export_ops.py 的等价性测试使用。
compare：用同一 fixture 再造库 → 走 4 个路由（重构后）→ 单元格快照与基线逐格比对。

注意：openpyxl 输出的 xlsx 字节含时间戳（docProps/core.xml + zip header），
逐字节不可复现；等价性以「读回单元格」为准（与任务要求的比对方式一致）。
报告中「生成时间」行逐次不同，比对时归一化剔除。
"""
import io
import json
import os
import sys
from datetime import date
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

# 显式设置环境（AGENTS.md：数据脚本一律显式设 DATABASE_URL；禁 pip）
os.environ["DATABASE_URL"] = "sqlite:///./data/export_baseline/fixture.db"
os.environ["AI_API_KEY"] = ""
os.environ["AI_BASE_URL"] = ""
os.environ["AI_AUTO_INTERPRET"] = "0"

import app.db as appdb  # noqa: E402
from app.auth import hash_password  # noqa: E402
from app.db import Base  # noqa: E402
from app.main import create_app  # noqa: E402
from app.models import (  # noqa: E402
    FormalRecord, ImportFile, MonthPerfRecord, PayrollPeriodRow, Person,
    RawRecord, ReconDayRow, ReconResult, ReconTask, User,
)
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402

FIXTURE_DB = REPO_ROOT / "data" / "export_baseline" / "fixture.db"
SNAPSHOT = REPO_ROOT / "mcp_service" / "tests" / "data" / "export_baseline.json"

ENDPOINTS = [
    ("salary_half1", "/perf/export?month=2026-09&period=half1"),
    ("salary_half2", "/perf/export?month=2026-09&period=half2"),
    ("settle", "/payroll-settle/export?month=2026-09"),
    ("recon_diff", "/recon/export?task_id={task_id}"),
    ("recon_report", "/recon/report?task_id={task_id}"),
]


def seed(db) -> int:
    """最小数据：正式表(3 行) + 月绩效(2 行) + 分期偏差(2 行) + 对账任务(差异 1 行)。"""
    admin = User(username="admin", password_hash=hash_password("pw123456"),
                 display_name="管理员", role="admin", is_active=True)
    db.add(admin)
    db.flush()
    db.add_all([
        Person(code="P001", display_name="甲"),
        Person(code="P002", display_name="乙"),
    ])
    imp = ImportFile(file_name="2026-09巡店.xlsx", file_sha256="b" * 64,
                     stored_path="x", uploaded_by=admin.id)
    db.add(imp)
    db.flush()
    raws = [
        RawRecord(import_id=imp.id, sheet_name="S1", excel_row=2,
                  store_id_raw="ST-A", store_name_local_raw="店A",
                  modified_raw="2026-09-01 09:00:00",
                  submitter_raw="甲(P001)", submitter_code="P001",
                  visible_raw="YES", deploy_raw="YES"),
        RawRecord(import_id=imp.id, sheet_name="S1", excel_row=3,
                  store_id_raw="ST-B", store_name_local_raw="店B",
                  modified_raw="2026-09-20 09:00:00",
                  submitter_raw="甲(P001)", submitter_code="P001",
                  visible_raw="YES", deploy_raw="YES"),
        RawRecord(import_id=imp.id, sheet_name="S1", excel_row=4,
                  store_id_raw="ST-C", store_name_local_raw="店C",
                  modified_raw="2026-09-05 09:00:00",
                  submitter_raw="乙(P002)", submitter_code="P002",
                  visible_raw="YES", deploy_raw="NO"),
    ]
    db.add_all(raws)
    db.flush()
    db.add_all([
        FormalRecord(import_id=imp.id, raw_record_id=raws[0].id,
                     person_code="P001", japan_date=date(2026, 9, 1), points=1),
        FormalRecord(import_id=imp.id, raw_record_id=raws[1].id,
                     person_code="P001", japan_date=date(2026, 9, 20), points=2),
        FormalRecord(import_id=imp.id, raw_record_id=raws[2].id,
                     person_code="P002", japan_date=date(2026, 9, 5), points=1),
    ])
    db.add_all([
        MonthPerfRecord(month="2026-09", person_code="P001", records=2,
                        p1=1, p2=1, points=3, salary=750, per_point=250),
        MonthPerfRecord(month="2026-09", person_code="P002", records=1,
                        p1=1, p2=0, points=1, salary=250, per_point=250),
    ])
    db.add_all([
        PayrollPeriodRow(month="2026-09", person_code="P001",
                         half1_points=1, half2_points=2,
                         half1_records=1, half1_p1=1, half1_p2=0,
                         half2_records=1, half2_p1=0, half2_p2=1,
                         settle_points=3, prev_adjust_points=0, diff_points=0,
                         half1_bonus=0, half2_bonus=0,
                         half1_amount=250, half2_amount=500,
                         settle_amount=750, prev_adjust_amount=0,
                         diff_amount=0, adjust_points=0, adjust_amount=0),
        PayrollPeriodRow(month="2026-09", person_code="P002",
                         half1_points=1, half2_points=0,
                         half1_records=1, half1_p1=1, half1_p2=0,
                         half2_records=0, half2_p1=0, half2_p2=0,
                         settle_points=1, prev_adjust_points=0, diff_points=0,
                         half1_bonus=0, half2_bonus=0,
                         half1_amount=250, half2_amount=0,
                         settle_amount=250, prev_adjust_amount=0,
                         diff_amount=0, adjust_points=0, adjust_amount=0),
    ])
    task = ReconTask(kind="monthly_v3", status="done", created_by=admin.id,
                     params={"month": "2026-09", "file": "对账_2026-09.xlsx"},
                     summary={"kind": "person_points", "compared": 2,
                              "diff_count": 1, "monthly_diff": 1,
                              "sys_only": ["P002"]})
    db.add(task)
    db.flush()
    db.add(ReconResult(task_id=task.id, submitter_code="P001",
                       system_value=3, report_value=5, diff=-2,
                       status="diff", family="person_points",
                       confirmed=False, note="甲 点数差异"))
    db.add(ReconDayRow(task_id=task.id, ref_date=date(2026, 9, 1),
                       person_code="P001", sys_points=1, rep_points=2,
                       sys_count=1, rep_count=1, diff=-1, side="both",
                       note=""))
    db.commit()
    return task.id


def build_client():
    FIXTURE_DB.parent.mkdir(parents=True, exist_ok=True)
    if FIXTURE_DB.exists():
        FIXTURE_DB.unlink()
    eng = create_engine(f"sqlite:///{FIXTURE_DB}",
                        connect_args={"check_same_thread": False})
    Base.metadata.create_all(eng)
    appdb._ENGINE = eng
    appdb.SessionLocal = sessionmaker(bind=eng, future=True,
                                      expire_on_commit=False)
    db = appdb.SessionLocal()
    task_id = seed(db)
    db.close()
    app = create_app()
    c = TestClient(app)
    r = c.post("/login", data={"username": "admin", "password": "pw123456"},
               follow_redirects=False)
    assert r.status_code in (302, 200), f"login failed: {r.status_code}"
    return c, task_id


def snapshot(c, task_id) -> dict:
    out = {}
    for kind, url in ENDPOINTS:
        r = c.get(url.format(task_id=task_id))
        assert r.status_code == 200, f"{kind}: HTTP {r.status_code}"
        assert r.content[:2] == b"PK", f"{kind}: not xlsx"
        wb = load_workbook(io.BytesIO(r.content), data_only=True)
        sheets = {}
        for ws in wb.worksheets:
            rows = [[None if v is None else v for v in row]
                    for row in ws.iter_rows(values_only=True)]
            sheets[ws.title] = rows
        out[kind] = sheets
    return out


def _norm(rows):
    """剔除逐次变化的「生成时间」行后比较。"""
    return [r for r in rows if not (r and r[0] == "生成时间")]


def main() -> int:
    mode = sys.argv[1] if len(sys.argv) > 1 else "capture"
    c, task_id = build_client()
    snap = snapshot(c, task_id)
    if mode == "capture":
        SNAPSHOT.parent.mkdir(parents=True, exist_ok=True)
        SNAPSHOT.write_text(json.dumps(snap, ensure_ascii=False, indent=1),
                            encoding="utf-8")
        print(f"captured {len(snap)} kinds -> {SNAPSHOT}")
        return 0
    base = json.loads(SNAPSHOT.read_text(encoding="utf-8"))
    diffs = []
    for kind in base:
        for sn in base[kind]:
            b = snap.get(kind, {}).get(sn)
            if b is None:
                diffs.append(f"{kind}/{sn}: sheet 缺失")
                continue
            if _norm(base[kind][sn]) != _norm(b):
                diffs.append(f"{kind}/{sn}: 单元格不一致\n  before={_norm(base[kind][sn])!r}\n  after ={_norm(b)!r}")
    if diffs:
        print(f"DIFFS: {len(diffs)}")
        for d in diffs:
            print(d)
        return 1
    print("OK: 改造后路由单元格快照与基线一致（生成时间除外）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
