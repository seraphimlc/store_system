# -*- coding: utf-8 -*-
"""下半月导入演练（上线前的"彩排"）。

用**真实的上半月文件** + 由它派生的**伪造下半月文件**（日期整体 +15 天 → 覆盖 9/16~9/30，
并与上半月在 9/16 重叠），在**全新库**上跑完整链路：

    上传 → 解析 → 入正式表(finalize) → 统计/月绩效/看板 → 自动生成对比报告

验证点：
1. 两次导入都能成功（不抛异常），并打印各阶段耗时（对比 nginx 600s 超时）
2. 数字与线上上半月的一致（同文件同口径 → 同结果）
3. 9/16 重叠不会双算（同店同月判重）
4. 入表后自动触发对比报告生成（`auto_for_import`）
5. 入表后覆盖边界推进 → 员工端"已对账锁定"随之扩大

AI 全程 mock（不连外网、不烧 token）；只动临时库，不碰演示库。
用法：./.venv/bin/python scripts/rehearse_h2_import.py
"""
import io
import os
import shutil
import sys
import time
from datetime import datetime, timedelta

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

SRC_XLSX = "/Users/liuchang/Desktop/万总/0901-0915MarsNavi store visit.xlsx"
DB = "/tmp/rehearse_h2.db"


def _log(msg):
    print("[%s] %s" % (datetime.now().strftime("%H:%M:%S"), msg), flush=True)


def make_h2_file(path):
    """把上半月文件的 Modified Time 整体 +15 天 → 伪造出 9/16~9/30 的下半月文件。"""
    import re

    from openpyxl import load_workbook
    wb = load_workbook(SRC_XLSX)
    ws = wb["STORE_TASK_EXCEL_SHEET "]
    shifted = 0
    for row in ws.iter_rows(min_row=3, min_col=3, max_col=3):
        c = row[0]
        v = c.value
        if isinstance(v, datetime):
            c.value = v + timedelta(days=15)
            shifted += 1
        elif v is not None:
            # 文件里的 Modified Time 是**字符串**（如 "2026-09-16 07:30:00"）
            m = re.match(r"^(\d{4})-(\d{2})-(\d{2})(.*)$", str(v).strip())
            if m:
                d = datetime(int(m.group(1)), int(m.group(2)),
                             int(m.group(3))) + timedelta(days=15)
                c.value = d.strftime("%Y-%m-%d") + m.group(4)
                shifted += 1
    wb.save(path)
    wb.close()
    return shifted


def main():
    if not os.path.exists(SRC_XLSX):
        _log("找不到源文件：%s" % SRC_XLSX)
        return 1
    for p in (DB, DB + "-journal"):
        if os.path.exists(p):
            os.remove(p)
    os.environ["DATABASE_URL"] = "sqlite:///%s" % DB
    os.environ["VISIT_REPORT_AI"] = "1"

    from sqlalchemy import create_engine

    import app.db as appdb
    import app.models  # noqa: F401
    from app.db import Base
    Base.metadata.create_all(create_engine("sqlite:///%s" % DB, future=True))
    _log("临时库已建：%s" % DB)

    # AI 全部 mock：只验证链路与触发，不烧 token
    from app.services import ai_chat, report_ai
    from app.services import ai_chat as _ai
    _ai.configured = lambda: True
    _ai.chat = lambda *a, **k: (
        ('{"overall_comment":"彩排","accuracy_notes":"n","per_person":{}}',
         {"total_tokens": 1}) if k.get("return_usage") else
        '{"overall_comment":"彩排","accuracy_notes":"n","per_person":{}}')
    report_ai.build_prompt = lambda res, lang: "彩排 prompt"
    _log("AI 已 mock")

    from app.auth import hash_password
    from app.models import User
    db = appdb.SessionLocal()
    db.add(User(username="admin", password_hash=hash_password("rehearse123"),
                display_name="彩排管理员", role="admin", is_active=True))
    db.commit()
    db.close()

    from fastapi.testclient import TestClient
    from app.main import create_app
    client = TestClient(create_app())

    h2_path = "/tmp/h2_fabricated.xlsx"
    shifted = make_h2_file(h2_path)
    _log("已伪造下半月文件（改写了 %d 行日期）：%s" % (shifted, h2_path))

    r = client.post("/login", data={"username": "admin", "password": "rehearse123"},
                    follow_redirects=False)
    assert r.status_code == 302, ("登录失败", r.status_code)

    results = {}
    for tag, path in (("上半月(真实文件)", SRC_XLSX), ("下半月(伪造文件)", h2_path)):
        t0 = time.time()
        with open(path, "rb") as fh:
            up = client.post("/files/upload",
                             data={"csrf_token": _csrf(client),
                                   "_ft": _token(client)},
                             files=[("files", (os.path.basename(path), fh.read(),
                                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))],
                             follow_redirects=False)
        t_up = time.time() - t0
        _log("%s 上传完成：HTTP %s，耗时 %.1fs" % (tag, up.status_code, t_up))
        imp_id = _last_import_id()
        t1 = time.time()
        fin = client.post("/files/%d/finalize" % imp_id,
                          data={"_ft": _token(client), "csrf_token": _csrf(client)},
                          follow_redirects=False)
        t_fin = time.time() - t1
        _log("%s 入表完成：HTTP %s，耗时 %.1fs（nginx 超时 600s）"
             % (tag, fin.status_code, t_fin))
        results[tag] = {"upload_s": round(t_up, 1), "finalize_s": round(t_fin, 1),
                        "import_id": imp_id, "http": (up.status_code, fin.status_code)}

    time.sleep(3)                       # 等后台线程（报告生成/看板分析）
    _summary(results)
    return 0


def _token(client):
    import re
    html = client.get("/files").text
    m = re.search(r'name="_ft" value="([^"]+)"', html)
    return m.group(1) if m else ""


def _csrf(client):
    from app.auth import SESSION_COOKIE, read_session_token
    d = read_session_token(client.cookies.get(SESSION_COOKIE)) or {}
    return d.get("csrf", "")


def _last_import_id():
    import app.db as appdb
    from app.models import ImportFile
    db = appdb.SessionLocal()
    imp = db.query(ImportFile).order_by(ImportFile.id.desc()).first()
    i = imp.id if imp else 0
    db.close()
    return i


def _summary(results):
    import app.db as appdb
    from app.models import (FormalRecord, ImportFile, MonthPerfRecord,
                            PersonDailyStat, StaffReportAnalysis)
    from sqlalchemy import func
    db = appdb.SessionLocal()
    print("\n===== 演练结果 =====")
    for tag, r in results.items():
        print("  %-16s 上传 %.1fs / 入表 %.1fs / HTTP %s"
              % (tag, r["upload_s"], r["finalize_s"], r["http"]))
    print("  文件记录：")
    for imp in db.query(ImportFile).order_by(ImportFile.id):
        print("    #%s %s status=%s" % (imp.id, imp.file_name, imp.status))
    lo, hi = db.query(func.min(FormalRecord.japan_date),
                      func.max(FormalRecord.japan_date)).first()
    print("  正式表覆盖：%s ~ %s，共 %d 行"
          % (lo, hi, db.query(FormalRecord).count()))
    print("  9 月各日行数（前 3 / 后 3）：")
    rows = (db.query(FormalRecord.japan_date, func.count())
            .filter(FormalRecord.japan_date >= "2026-09-01")
            .group_by(FormalRecord.japan_date)
            .order_by(FormalRecord.japan_date).all())
    for d, n in rows[:3] + rows[-3:]:
        print("    %s: %d" % (d, n))
    print("  9/16 当日行数（两次导入重叠日，应不重复计）: %s"
          % [n for d, n in rows if str(d) == "2026-09-16"])
    print("  person_daily_stats: %d 行 / month_perf_records: %d 行"
          % (db.query(PersonDailyStat).count(), db.query(MonthPerfRecord).count()))
    print("  自动生成的对比报告：")
    for a in db.query(StaffReportAnalysis).order_by(StaffReportAnalysis.id):
        print("    #%s %s~%s status=%s" % (a.id, a.period_start, a.period_end,
                                           a.status))
    from app.services import daily_report
    print("  员工端「已对账」边界（coverage_end）: %s"
          % daily_report.coverage_end(db))
    db.close()


if __name__ == "__main__":
    sys.exit(main())
