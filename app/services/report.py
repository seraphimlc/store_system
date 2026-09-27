# -*- coding: utf-8 -*-
"""V3 文件判定明细（/files/{id}/report）：基于 raw_records.clean_status + 申诉状态。

替代旧版 analytics.import_report（旧 CleanRecord 桶模型，V3 流程已不再写入）。
"""
from typing import Optional

from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models import AppealRecord, FormalRecord, ImportFile, Person, RawRecord

PAGE = 50

STATUS_META = {
    "valid": ("有效（自动计入绩效）", "ok"),
    "master_late": ("同店跨日：该店当月已有更早记录", "warn"),
    "from_sub": ("从档归并：同店不同编号，主档保留", "warn"),
    "cross_file_dup": ("同日跨文件重复导入（同一条记录）", "err"),
    "blank": ("Visible 空白：不计有效、不判重", ""),
    "no_ref": ("无对应", ""),
}

BUCKET_OPTS = [
    ("", "全部"),
    ("valid", "有效"),
    ("master_late", "同店跨日"),
    ("from_sub", "从档"),
    ("cross_file_dup", "重复导入"),
    ("blank", "Visible 空白"),
    ("appealing", "有申诉"),
]


def import_months(db: Session, import_id: int):
    """该文件数据涉及的自然月（按 Modified 前 7 位去重，升序）。"""
    vals = db.query(RawRecord.modified_raw).filter(
        RawRecord.import_id == import_id).all()
    months = {((m or "")[:7]) for (m,) in vals if m and len(m) >= 7}
    return sorted(months)


def import_report(db: Session, import_id: int, bucket: str = "",
                  q: str = "", page: int = 1):
    """按导入文件汇总 raw 判定结果（V3 口径）。

    bucket: "" 全部 / valid / master_late / from_sub / cross_file_dup /
            blank / appealing（有申诉的行）。
    返回 dict：file/raw_total/各状态计数/rows/total/page/pages/bucket/q/months。
    """
    f = db.get(ImportFile, import_id)
    if f is None:
        return None
    qry = db.query(RawRecord).filter(RawRecord.import_id == import_id)
    if bucket == "appealing":
        qry = qry.join(AppealRecord, AppealRecord.raw_record_id == RawRecord.id)
    elif bucket == "blank":
        qry = qry.filter(RawRecord.clean_status.in_(("blank", "visible_blank")))
    elif bucket:
        qry = qry.filter(RawRecord.clean_status == bucket)
    if q:
        like = f"%{q}%"
        hit = [p.code for p in db.query(Person).filter(
            Person.display_name.like(like)).all()]
        cond = ((RawRecord.store_id_raw.like(like))
                | (RawRecord.store_name_local_raw.like(like))
                | (RawRecord.submitter_raw.like(like)))
        if hit:
            cond = cond | RawRecord.submitter_code.in_(hit)
        qry = qry.filter(cond)
    total = qry.count()
    rows_raw = (qry.order_by(RawRecord.modified_raw, RawRecord.excel_row)
                .offset((page - 1) * PAGE).limit(PAGE).all())
    names = {p.code: p.display_name for p in db.query(Person).all()}
    appeals = {ap.raw_record_id: ap for ap in db.query(AppealRecord).all()}
    formal_ids = {fr.raw_record_id for fr in db.query(FormalRecord).filter(
        FormalRecord.import_id == import_id).all()}
    out = []
    for rr in rows_raw:
        st = rr.clean_status or ""
        if st == "visible_blank":
            st = "blank"   # 旧枚举归一
        label, _pill = STATUS_META.get(st, (st, ""))
        ap = appeals.get(rr.id)
        out.append({
            "id": rr.id,
            "code": rr.submitter_code,
            "name": names.get(rr.submitter_code, rr.submitter_raw or "未识别"),
            "store_id": rr.store_id_raw, "store_name": rr.store_name_local_raw,
            "date": (rr.modified_raw or "")[:10],
            "visible": rr.visible_raw or "", "deploy": rr.deploy_raw or "",
            "status": st, "label": label,
            "excel_row": rr.excel_row, "sheet": rr.sheet_name,
            "in_formal": rr.id in formal_ids,
            "appeal": ({"status": ap.status, "reason": ap.reason or ""}
                       if ap else None),
        })
    cnt = dict(db.query(RawRecord.clean_status, func.count()).filter(
        RawRecord.import_id == import_id).group_by(
        RawRecord.clean_status).all())
    if "visible_blank" in cnt:
        cnt["blank"] = cnt.get("blank", 0) + cnt.pop("visible_blank")
    raw_total = db.query(RawRecord).filter(
        RawRecord.import_id == import_id).count()
    pages = (total + PAGE - 1) // PAGE if total else 1
    return {"file": f, "raw_total": raw_total,
            "valid": cnt.get("valid", 0),
            "master_late": cnt.get("master_late", 0),
            "from_sub": cnt.get("from_sub", 0),
            "cross_file_dup": cnt.get("cross_file_dup", 0),
            "blank": cnt.get("blank", 0),
            "no_ref": cnt.get("no_ref", 0),
            "appealing": db.query(AppealRecord).filter(
                AppealRecord.import_id == import_id,
                AppealRecord.status == "pending").count(),
            "filtered_total": sum(cnt.get(k, 0) for k in
                                  ("master_late", "from_sub",
                                   "cross_file_dup", "blank", "no_ref")),
            "rows": out, "total": total, "page": page, "pages": pages,
            "bucket": bucket, "q": q, "months": import_months(db, import_id)}


# ---------------------------------------------------------------------------
# P1 导出重构（task p1-mcpify-export）：把 4 个导出路由的工作簿构建段抽成纯函数。
# 行为保持型重构——代码自 settle_r.py 原路由逐段搬移，未改任何业务逻辑/格式；
# 路由改为调用这些函数后输出与改造前一致（见 scripts/p1_export_baseline.py 与
# mcp_service/tests/test_export_ops.py 的逐单元格等价性测试）。
# 注：openpyxl 输出的字节含时间戳，逐字节不可复现，等价性按「读回单元格」验证。
# ---------------------------------------------------------------------------

def build_payroll_workbook(db, month, period="half1"):
    """发薪表 Excel（上半月/下半月 + 每人一 sheet 日明细）。

    原 /perf/export 路由的构建段（汇总 sheet + 每人一 sheet），行为逐单元格一致。
    返回 (xlsx字节, 建议文件名)。month 为路由回落后的最终月份（"" 时与原路由
    行为一致——在月份解析处同样失败，不额外兜底）。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from app.services import perf
    from app.services import period as _payroll

    rows = perf.month_perf(db, month)
    summary = perf.company_summary(db, month)
    adj = _payroll.carry_map(db, month) if month else {}
    adj_amt_map = {k: v[1] for k, v in adj.items()}
    adj_amt_total = sum(adj_amt_map.values())

    bold = Font(bold=True)
    wb = Workbook()
    ws1 = wb.active
    # 两期发薪：half1=上半月(20日发)，half2=下半月(次月5日发)
    if period not in ("half1", "half2"):
        period = "half1"
    payroll_rows = _payroll.period_rows(db, month) if month else []
    half_map = {r["code"]: (r["half1"], r["half1_amt"], r["half1_bonus"],
                            r["half2"], r["half2_amt"], r["half2_bonus"])
                for r in payroll_rows}
    half_stats = (_payroll.half_stats_map(db, month) if month else {})
    is_h1 = period == "half1"
    ws1.title = ("上半月发薪1-15" if is_h1 else "下半月发薪16-月末") + (
        f"·{month}" if month else "")
    ws1.append(["员工编号", "姓名", "该期点数", "有效店数", "点数(1点/2点)",
                "2点成功率", "奖金(円)", "总金额(円)",
                "找平金额(円)", "应该付金额(円)"])
    for c in range(1, 11):
        ws1.cell(1, c).font = bold
    t_pay = 0
    for r in rows:
        hf = half_map.get(r["code"], (0, 0, 0, 0, 0, 0))
        hs = half_stats.get(r["code"], {"h1": (0, 0, 0),
                                        "h2": (0, 0, 0)})
        if is_h1:
            hpts, hamt, hbonus = hf[0], hf[1], hf[2]
            st = hs["h1"]
        else:
            hpts, hamt, hbonus = hf[3], hf[4], hf[5]
            st = hs["h2"]
        hp1, hp2 = st[1], st[2]
        hrate = (hp2 / (hp1 + hp2)) if (hp1 + hp2) else 0.0
        carry = adj_amt_map.get(r["code"], 0) if is_h1 else 0
        t_pay += hamt + carry
        ws1.append([r["code"], r["name"], hpts, st[0],
                    f"{hp1} / {hp2}", round(hrate, 4),
                    hbonus, hamt, carry, hamt + carry])
    ws1.append([])
    ws1.append(["合计", "", "", summary["records"],
                f"{summary['p1']} / {summary['p2']}",
                round(summary["rate37"], 4), "",
                sum(h[1 if is_h1 else 4] for h in half_map.values()),
                (adj_amt_total if is_h1 else 0), t_pay])
    for c in range(1, 11):
        ws1.cell(ws1.max_row, c).font = bold
    # 每人一个 sheet：该员工本期内所有有效巡店记录（逐条，便于对账）
    from datetime import date as _d
    y, m0 = int(month[:4]), int(month[5:7])
    if is_h1:
        lo, hi = _d(y, m0, 1), _d(y, m0, 16)      # 上半月 1~15（含15号）
    else:
        lo, hi = _d(y, m0, 16), _d(y + 1, 1, 1) if m0 == 12 else _d(y, m0 + 1, 1)
    _PERS = {p.code: p.display_name for p in db.query(Person).all()}
    by_person = {}
    for f, rr in (db.query(FormalRecord, RawRecord)
                  .join(RawRecord, FormalRecord.raw_record_id == RawRecord.id)
                  .filter(FormalRecord.japan_date >= lo,
                          FormalRecord.japan_date < hi).all()):
        by_person.setdefault(f.person_code, []).append(
            (f.japan_date, rr.store_name_local_raw or "",
             rr.modified_raw or "", rr.visible_raw or "",
             rr.deploy_raw or "", f.points or 0))
    for code in sorted(by_person):
        nm = _PERS.get(code, code) or code
        sname = f"{nm}({code})"[:31]
        ws = wb.create_sheet(sname)
        ws.append(["日期", "店铺名", "巡店时间", "S1(审核状态)", "投放",
                   "点数", "员工编号"])
        for c in range(1, 8):
            ws.cell(1, c).font = bold
        for rec in sorted(by_person[code], key=lambda x: (x[0], x[2])):
            ws.append([str(rec[0]), rec[1], rec[2], rec[3], rec[4],
                       rec[5], code])
    for ws in wb.worksheets:
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col) + 2
            ws.column_dimensions[col[0].column_letter].width = width

    import io
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    po = "上半月" if period == "half1" else "下半月"
    fname = "发薪表_%s_%s.xlsx" % (po, month or "all")
    return bio.getvalue(), fname


def build_payroll_settle_workbook(db, month):
    """月度分期对账偏差表 Excel（17 列）。原 /payroll-settle/export 构建段。

    返回 (xlsx字节, 建议文件名)。不校验 month 格式（与原路由一致，由调用方把关）。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from app.services import period as _payroll
    wb = Workbook()
    ws = wb.active
    ws.title = "月度分期对账偏差"
    ws.append(["月份", "员工编号", "员工姓名", "上半月点数", "上半月金额(円)",
               "下半月点数", "下半月金额(円)", "奖金(円)", "分期已发(円)",
               "对账点数", "对账金额(円)", "上月修正(点)", "上月修正金额(円)",
               "对账偏差(点,参考)", "对账偏差金额(円,参考)",
               "找平(点)", "找平金额(円)"])
    for c in range(1, 18):
        ws.cell(1, c).font = Font(bold=True)
    for r in _payroll.period_rows(db, month):
        ws.append([month, r["code"], r["name"], r["half1"], r["half1_amt"],
                   r["half2"], r["half2_amt"],
                   r["half1_bonus"] + r["half2_bonus"],
                   r["half1_amt"] + r["half2_amt"],
                   r["settle"], r["settle_amt"], r["prev"], r["prev_amt"],
                   r["diff"], r["diff_amt"], r["adj"], r["adj_amt"]])
    ws.append([])
    ws.append(["说明：对账偏差=系统参考值(自动刷新)；找平=人工实际执行值(默认0，"
               "金额=点×该月单价)；上月修正=上月找平递延。"])
    import io as _io
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fn = f"payroll_settle_{month}.xlsx"
    return buf.getvalue(), fn


def build_recon_diff_workbook(db, task_id):
    """对账差异 Excel：差异明细 + 反向名单。原 /recon/export 构建段。

    任务不存在（含 task_id=0）返回 None（路由据此抛 404，行为与原路由一致）。
    返回 (xlsx字节, 建议文件名)。
    """
    from app.models import Person, ReconResult, ReconTask
    from openpyxl import Workbook
    from openpyxl.styles import Font
    cur = db.get(ReconTask, task_id) if task_id else None
    if cur is None:
        return None
    rows = (db.query(ReconResult).filter(ReconResult.task_id == task_id)
            .order_by(ReconResult.submitter_code).all())
    summary = cur.summary or {}
    bold = Font(bold=True)
    wb = Workbook()
    # Sheet1 差异明细
    ws = wb.active
    ws.title = "差异明细"
    ws.append(["月份", (cur.params or {}).get("month", ""),
               "文件", (cur.params or {}).get("file", "")])
    ws.append(["比对人数", summary.get("compared", 0),
               "差异条数", summary.get("diff_count", 0)])
    ws.append([])
    if not rows:
        ws.append(["无差异：两侧点数一致"])
    else:
        ws.append(["人员", "编号", "系统点数", "对账点数", "差异(系统-对账)"])
        for c in range(1, 6):
            ws.cell(ws.max_row, c).font = bold
        for r in rows:
            ws.append([r.note.replace(" 点数差异", "") if r.note else "",
                       r.submitter_code, r.system_value, r.report_value,
                       r.diff])
    # Sheet2 反向名单
    ws2 = wb.create_sheet("系统有而对账文件无")
    so = summary.get("sys_only") or []
    if not so:
        ws2.append(["无（对账文件已覆盖系统全部有记录员工）"])
    else:
        ws2.append(["编号", "姓名"])
        ws2.cell(1, 1).font = bold
        ws2.cell(1, 2).font = bold
        persons = {p.code: p.display_name for p in db.query(Person).all()}
        for c in so:
            ws2.append([c, persons.get(c, c)])
    for wsx in (ws, ws2):
        for col in wsx.columns:
            w = max(len(str(c.value or "")) for c in col) + 2
            wsx.column_dimensions[col[0].column_letter].width = min(w, 40)
    import io
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    fname = f"对账差异_任务{task_id}.xlsx"
    return bio.getvalue(), fname


def build_recon_report_workbook(db, task_id, author_name="系统"):
    """月度对账报告 Excel。原 /recon/report 路由（复用 recon.build_report 后序列化）。

    author_name 为「生成人」（原路由传 user.display_name，故保留参数保证
    输出与改造前一致）。任务不存在返回 None。返回 (xlsx字节, 建议文件名)。
    """
    from app.services import recon
    wb = recon.build_report(db, task_id, author_name)
    if wb is None:
        return None
    import io
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    fname = f"月度对账报告_任务{task_id}.xlsx"
    return bio.getvalue(), fname
