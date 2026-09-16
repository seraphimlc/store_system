# -*- coding: utf-8 -*-
"""V3 页面路由：员工对账（我的判定确认）+ 管理员确认总览 + 入正式表。"""
from datetime import datetime
from typing import Optional

from fastapi import (APIRouter, Depends, File, Form, HTTPException,
                        Request, UploadFile)
from fastapi.responses import HTMLResponse, RedirectResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session

from app.db import get_db
from app.models import (AppealRecord, FormalRecord, ImportFile, Person,
                        RawRecord, User)
from app.routers.auth_r import csrf_ok, require_login
from app.services import v3_flow

router = APIRouter()
templates = Jinja2Templates(directory="app/templates")

_REASON_CN = {"master_late": "同主档已有更早有效（本店非首次）",
              "from_sub": "从档编号，已归并到主档店铺",
              "cross_file_dup": "同日同店跨文件重复（自动过滤）",
              "no_ref": "无店铺编号/无法归主档",
              "disputed": "员工申诉待处理"}


def _denied():
    return RedirectResponse("/login", status_code=302)


def _reason_cn(rr) -> str:
    return _REASON_CN.get(rr.filter_reason or "", rr.filter_reason or "")


# ---------------- 员工端：我的申诉 ----------------
@router.get("/my/confirm", response_class=HTMLResponse)
def my_confirm_legacy(request: Request,
                      user: Optional[User] = Depends(require_login)):
    """旧路径兼容 → 我的申诉。"""
    if user is None or user.role != "staff" or not user.person_code:
        return _denied()
    return RedirectResponse("/my/appeal", status_code=302)


@router.get("/my/appeal", response_class=HTMLResponse)
def my_appeal_page(request: Request,
                   user: Optional[User] = Depends(require_login),
                   db: Session = Depends(get_db), msg: str = ""):
    """我的申诉：被滤记录(master_late/from_sub)默认已认可滤除、不打扰；
    仅当员工认为某条被滤错时在此申诉。"""
    if user is None or user.role != "staff" or not user.person_code:
        return _denied()
    code = user.person_code
    days = v3_flow.appeal_list(db, code)
    appeal_map = v3_flow.appeal_map(db, code)
    total_appealable = sum(len(v) for _, v in days)
    return templates.TemplateResponse("my_appeal.html", {
        "request": request, "current_user": user, "msg": msg,
        "days": days, "appeal_map": appeal_map,
        "total_appealable": total_appealable,
        "reason_cn": _reason_cn})


@router.post("/my/appeal/{raw_id}")
def my_appeal(raw_id: int, request: Request, reason: str = Form(""),
              csrf_token: str = Form(...),
              user: Optional[User] = Depends(require_login),
              db: Session = Depends(get_db)):
    if user is None or user.role != "staff":
        return _denied()
    if not csrf_ok(request, csrf_token):
        return HTMLResponse("CSRF 校验失败", status_code=400)
    try:
        v3_flow.create_appeal(db, raw_id, user.person_code, reason)
    except ValueError as e:
        raise HTTPException(400, str(e))
    return RedirectResponse(f"/my/appeal?msg=已申诉，等待管理员处理",
                            status_code=303)


# ---------------- 管理员端：绩效确认（申诉处理） ----------------
@router.get("/confirm-admin", response_class=HTMLResponse)
def confirm_admin(request: Request,
                  user: Optional[User] = Depends(require_login),
                  db: Session = Depends(get_db), msg: str = "",
                  imp_id: int = 0, only: str = "", staff: str = ""):
    """绩效确认（管理端）。

    文件总览（有效/重复导入/申诉/正式表）+ 处理申诉（认可→改判有效计入
    绩效 / 驳回→维持滤除）+ 按员工查看其被滤/申诉情况。有效记录自动入
    绩效；被滤默认已认可，仅员工申诉需处理。
    """
    from app.models import AppealRecord
    if user is None or user.role != "admin":
        return _denied()
    files = db.query(ImportFile).order_by(ImportFile.id.desc()).all()
    rows = []
    for f in files:
        pend = db.query(AppealRecord).filter(
            AppealRecord.import_id == f.id,
            AppealRecord.status == "pending").count()
        n_appeal = db.query(AppealRecord).filter(
            AppealRecord.import_id == f.id).count()
        formal = db.query(FormalRecord).filter(
            FormalRecord.import_id == f.id).count()
        n_valid = db.query(RawRecord).filter(
            RawRecord.import_id == f.id,
            RawRecord.clean_status == "valid").count()
        n_filtered = db.query(RawRecord).filter(
            RawRecord.import_id == f.id,
            RawRecord.clean_status.in_(("master_late", "from_sub"))).count()
        n_dup = db.query(RawRecord).filter(
            RawRecord.import_id == f.id,
            RawRecord.clean_status == "cross_file_dup").count()
        rows.append({"file": f, "pend": pend, "appeals": n_appeal,
                     "formal": formal, "valid": n_valid,
                     "filtered": n_filtered, "dup": n_dup})
    # 申诉列表（选中文件后展示，含处理按钮）
    appeal_rows = []
    if imp_id:
        appeal_rows = (db.query(AppealRecord, RawRecord)
                       .join(RawRecord, RawRecord.id == AppealRecord.raw_record_id)
                       .filter(AppealRecord.import_id == imp_id)
                       .order_by(AppealRecord.created_at).all())
    # 按员工查看：该员工被滤(master_late/from_sub)与申诉状态
    staff_list = (db.query(Person.code, Person.display_name)
                  .join(RawRecord, RawRecord.submitter_code == Person.code)
                  .distinct().order_by(Person.display_name).all())
    staff_name = ""
    staff_rows = []
    if staff:
        pn = db.query(Person.display_name).filter(
            Person.code == staff).first()
        staff_name = pn[0] if pn else staff
        staff_rows = (db.query(RawRecord)
                      .filter(RawRecord.submitter_code == staff,
                              RawRecord.clean_status.in_(
                                  ("master_late", "from_sub")))
                      .order_by(RawRecord.modified_raw.desc()).all())
    appeal_map = v3_flow.appeal_map(db, staff or "")
    return templates.TemplateResponse("confirm_admin.html", {
        "request": request, "current_user": user, "msg": msg,
        "rows": rows, "appeal_rows": appeal_rows, "imp_id": imp_id,
        "reason_cn": _reason_cn, "staff": staff, "staff_name": staff_name,
        "staff_list": staff_list, "staff_rows": staff_rows,
        "appeal_map": appeal_map})


@router.post("/files/{fid}/finalize")
def finalize_file(fid: int, request: Request,
                  csrf_token: str = Form(...),
                  user: Optional[User] = Depends(require_login),
                  db: Session = Depends(get_db)):
    if user is None or user.role != "admin":
        return _denied()
    if not csrf_ok(request, csrf_token):
        return HTMLResponse("CSRF 校验失败", status_code=400)
    res = v3_flow.finalize_import(db, fid, user.id)
    if not res["ok"]:
        raise HTTPException(400, res.get("msg", "无法入正式表"))
    return RedirectResponse(f"/confirm-admin?imp_id={fid}"
                            f"&msg=已入正式表 {res['added']} 条", status_code=303)


@router.post("/files/{fid}/appeals/{appeal_id}/resolve")
def resolve_appeal(fid: int, appeal_id: int, request: Request,
                   decision: str = Form("accept"),
                   csrf_token: str = Form(...),
                   user: Optional[User] = Depends(require_login),
                   db: Session = Depends(get_db)):
    if user is None or user.role != "admin":
        return _denied()
    if not csrf_ok(request, csrf_token):
        return HTMLResponse("CSRF 校验失败", status_code=400)
    res = v3_flow.resolve_appeal(db, appeal_id, decision, user.id, file_id=fid)
    if not res["ok"]:
        raise HTTPException(404, res.get("msg", "申诉不存在或已处理"))
    return RedirectResponse(f"/confirm-admin?imp_id={fid}"
                            f"&msg=已处理申诉", status_code=303)


# ---------------- 员工端：我的绩效（日/月报表） ----------------
@router.get("/my/perf", response_class=HTMLResponse)
def my_perf(request: Request,
            user: Optional[User] = Depends(require_login),
            db: Session = Depends(get_db), month: str = "",
            date: str = ""):
    """员工看自己当月的绩效：日明细（日期可选）+ 月汇总（只含正式表有效店）。"""
    if user is None or user.role != "staff" or not user.person_code:
        return _denied()
    from app.services import v3_perf
    code = user.person_code
    mine = [r for r in db.query(FormalRecord).filter(
        FormalRecord.person_code == code).all()]
    months = sorted({(str(r.japan_date or ""))[:7] for r in mine if r.japan_date})
    if month not in months:
        month = months[-1] if months else ""
    daily = [d for d in v3_perf.daily_perf(db, month)
             if d["code"] == code]
    dates = sorted({str(d["date"]) for d in daily})
    if not date or date not in dates:
        date = dates[-1] if dates else ""
    if date:
        daily = [d for d in daily if str(d["date"]) == date]
    me = next((m for m in v3_perf.month_perf(db, month)
               if m["code"] == code), None)
    summary = None
    if me:
        adj = v3_perf.adjust_map(db, month).get(code, 0)
        summary = {"records": me["records"], "p1": me["p1"], "p2": me["p2"],
                   "points": me["points"], "amount": me["amount"],
                   "adjust": adj, "payable": me["amount"] + adj,
                   "rate37": me["rate37"]}
    return templates.TemplateResponse("my_perf.html", {
        "request": request, "current_user": user, "month": month,
        "months": months, "daily": daily, "summary": summary,
        "date": date, "dates": dates})


# ---------------- V3 绩效 / 工资 ----------------
@router.get("/perf", response_class=HTMLResponse)
@router.get("/v3/perf", response_class=HTMLResponse)
def v3_perf(request: Request,
            user: Optional[User] = Depends(require_login),
            db: Session = Depends(get_db), month: str = "",
            date: str = "", staff: str = "", period: str = "half1"):
    if user is None or user.role != "admin":
        return _denied()
    from app.services import v3_perf
    months = sorted({(str(r.japan_date or ""))[:7]
                     for r in db.query(FormalRecord).all()
                     if r.japan_date})
    if month not in months:
        month = months[-1] if months else ""
    rows = v3_perf.month_perf(db, month)
    summary = v3_perf.company_summary(db, month)
    # 上月找平 = 上月未找平余量（diff−adjust，同一张表 payroll_period_rows）
    from app.services import v3_period
    adj = v3_period.carry_map(db, month) if month else {}
    adj_map = {k: v[0] for k, v in adj.items()}          # 余量(点)
    adj_amt_map = {k: v[1] for k, v in adj.items()}      # 余量金额(×上月单价)
    adj_total = sum(adj_map.values())
    adj_amt_total = sum(adj_amt_map.values())
    per_point = month and v3_perf.month_per_point(db, month) or 250
    payroll_rows = v3_period.period_rows(db, month) if month else []
    payroll_total = {
        "settle_amt": sum(r["settle_amt"] for r in payroll_rows),
        "diff_amt": sum(r["diff_amt"] for r in payroll_rows),
        "diff": sum(r["diff"] for r in payroll_rows),
        "half1": sum(r["half1"] for r in payroll_rows),
        "half2": sum(r["half2"] for r in payroll_rows),
        "half1_amt": sum(r["half1_amt"] for r in payroll_rows),
        "half2_amt": sum(r["half2_amt"] for r in payroll_rows),
    }
    # 本月找平（实时读薪资找平同一张表 payroll_period_rows，保存后立即同步）
    cur_adj = {r["code"]: (r["adj"], r["adj_amt"]) for r in payroll_rows}
    # 两期发薪视图（月内发两次工资照此算）：
    # code -> (上半月点, 上半月金额(含该期奖金), 上半月奖金, 下半月点, 下半月金额, 下半月奖金)
    half_map = {}
    for r in payroll_rows:
        half_map[r["code"]] = (r["half1"], r["half1_amt"], r["half1_bonus"],
                               r["half2"], r["half2_amt"], r["half2_bonus"])
    # 按期的 有效店/1点/2点（期视图这三列只看该期，避免全月数据误导）
    half_stats = v3_period.half_stats_map(db, month) if month else {}
    # 两期应付（金额找平）：上半月应付=上半月金额+上月金额余量(正补负扣)；
    # 若上半月不够扣(为负)→上半月发0、剩余负数转到下半月继续扣。
    pay_map = {}
    for r in rows:
        hf = half_map.get(r["code"], (0, 0, 0, 0, 0, 0))
        carry = adj_amt_map.get(r["code"], 0)
        p1 = hf[1] + carry
        if p1 < 0:
            pay_map[r["code"]] = (0, hf[4] + p1)
        else:
            pay_map[r["code"]] = (p1, hf[4])
    if period not in ("half1", "half2"):
        period = "half1"
    bonus_g, bonus_a = v3_perf.bonus_params()
    page_state = {
        "page": "perf", "month": month, "period": period,
        "employees": len(rows),
        "total_points": summary.get("total_points", 0),
        "salary_total": summary.get("total_amount", 0),
        "payroll_total": payroll_total.get("half1_amt", 0)
        + payroll_total.get("half2_amt", 0),
        "payroll_rows": len(payroll_rows),
        "per_point": per_point,
        "bonus_group": bonus_g, "bonus_amount": bonus_a,
    }
    return templates.TemplateResponse("v3_perf.html", {
        "request": request, "current_user": user, "month": month,
        "months": months, "rows": rows, "period": period,
        "summary": summary, "adj_map": adj_map, "adj_total": adj_total,
        "adj_amt_map": adj_amt_map, "adj_amt_total": adj_amt_total,
        "per_point": per_point, "half_stats": half_stats,
        "cur_adj": cur_adj, "half_map": half_map,
        "bonus_g": bonus_g, "bonus_a": bonus_a,
        "pay_map": pay_map,
        "payroll_rows": payroll_rows, "payroll_total": payroll_total,
        "page_state": page_state})


@router.post("/perf/per-point")
def v3_set_per_point(request: Request, month: str = Form(""),
                     per_point: int = Form(250),
                     csrf_token: str = Form(...),
                     user: Optional[User] = Depends(require_login),
                     db: Session = Depends(get_db)):
    """设置该月点数单价（円/点）：月绩效工资与薪资找平金额按新单价重算；
    已确认找平记录锁存当时单价，不受影响。"""
    if user is None or user.role != "admin":
        return _denied()
    if not csrf_ok(request, csrf_token):
        return HTMLResponse("CSRF 校验失败", status_code=400)
    from app.services import v3_perf
    from urllib.parse import quote as _q
    if per_point <= 0 or not month:
        return RedirectResponse(
            f"/perf?month={month}&msg={_q('单价需为正整数')}",
            status_code=303)
    v3_perf.set_month_per_point(db, month, per_point)
    return RedirectResponse(
        f"/perf?month={month}&msg="
        f"{_q(f'{month} 点数单价已设为 {per_point}円/点（工资/找平金额已重算；已确认找平按当时单价不变）')}",
        status_code=303)


@router.get("/perf/daily", response_class=HTMLResponse)
@router.get("/v3/perf/daily", response_class=HTMLResponse)
def v3_perf_daily(request: Request,
                  user: Optional[User] = Depends(require_login),
                  db: Session = Depends(get_db), month: str = "",
                  date: str = "", staff: str = ""):
    """日绩效明细（独立页面，从绩效工资点「详细」进入）。"""
    if user is None or user.role != "admin":
        return _denied()
    from app.services import v3_perf
    months = sorted({(str(r.japan_date or ""))[:7]
                     for r in db.query(FormalRecord).all()
                     if r.japan_date})
    if month not in months:
        month = months[-1] if months else ""
    daily = v3_perf.daily_perf(db, month)
    dates = sorted({str(d["date"]) for d in daily})
    staff_opts = sorted({(d["code"], d["name"]) for d in daily})
    if date and date not in dates:
        date = ""
    if not date:
        if staff:
            date = ""
        else:
            date = dates[-1] if dates else ""
    if date:
        daily = [d for d in daily if str(d["date"]) == date]
    if staff:
        daily = [d for d in daily if d["code"] == staff]
    return templates.TemplateResponse("daily_detail.html", {
        "request": request, "current_user": user, "month": month,
        "months": months, "daily": daily, "date": date, "staff": staff,
        "dates": dates, "staff_opts": staff_opts})


@router.get("/perf/export")
@router.get("/v3/perf/export")
def v3_perf_export(request: Request, user: Optional[User] =
                   Depends(require_login),
                   db: Session = Depends(get_db), month: str = "",
                   period: str = "half1"):
    """导出发薪表 Excel（按期：上半月/下半月 + 日明细）。"""
    if user is None or user.role != "admin":
        return _denied()
    from app.services import v3_perf
    from openpyxl import Workbook
    from openpyxl.styles import Font
    months = sorted({(str(r.japan_date or ""))[:7]
                     for r in db.query(FormalRecord).all()
                     if r.japan_date})
    if month not in months:
        month = months[-1] if months else ""
    rows = v3_perf.month_perf(db, month)
    daily = v3_perf.daily_perf(db, month)
    summary = v3_perf.company_summary(db, month)
    from app.services import v3_period
    adj = v3_period.carry_map(db, month) if month else {}
    adj_map = {k: v[0] for k, v in adj.items()}
    adj_amt_map = {k: v[1] for k, v in adj.items()}
    adj_total = sum(adj_map.values())
    adj_amt_total = sum(adj_amt_map.values())

    bold = Font(bold=True)
    wb = Workbook()
    ws1 = wb.active
    # 两期发薪：half1=上半月(20日发)，half2=下半月(次月5日发)
    if period not in ("half1", "half2"):
        period = "half1"
    payroll_rows = v3_period.period_rows(db, month) if month else []
    half_map = {r["code"]: (r["half1"], r["half1_amt"], r["half1_bonus"],
                            r["half2"], r["half2_amt"], r["half2_bonus"])
                for r in payroll_rows}
    half_stats = (v3_period.half_stats_map(db, month) if month else {})
    is_h1 = period == "half1"
    ws1.title = ("上半月发薪1-15" if is_h1 else "下半月发薪16-月末") + (
        f"·{month}" if month else "")
    ws1.append(["员工编号", "姓名", "该期点数", "有效店数", "点数(1点/2点)",
                "2点成功率", "奖金(円)", "总金额(円)",
                "找平金额(円)", "应该付金额(円)"])
    for c in range(1, 11):
        ws1.cell(1, c).font = bold
    t_pay = 0
    for r in rows:
        hf = half_map.get(r["code"], (0, 0, 0, 0, 0, 0))
        hs = half_stats.get(r["code"], {"h1": (0, 0, 0),
                                        "h2": (0, 0, 0)})
        if is_h1:
            hpts, hamt, hbonus = hf[0], hf[1], hf[2]
            st = hs["h1"]
        else:
            hpts, hamt, hbonus = hf[3], hf[4], hf[5]
            st = hs["h2"]
        hp1, hp2 = st[1], st[2]
        hrate = (hp2 / (hp1 + hp2)) if (hp1 + hp2) else 0.0
        carry = adj_amt_map.get(r["code"], 0) if is_h1 else 0
        t_pay += hamt + carry
        ws1.append([r["code"], r["name"], hpts, st[0],
                    f"{hp1} / {hp2}", round(hrate, 4),
                    hbonus, hamt, carry, hamt + carry])
    ws1.append([])
    ws1.append(["合计", "", "", summary["records"],
                f"{summary['p1']} / {summary['p2']}",
                round(summary["rate37"], 4), "",
                sum(h[1 if is_h1 else 4] for h in half_map.values()),
                (adj_amt_total if is_h1 else 0), t_pay])
    for c in range(1, 11):
        ws1.cell(ws1.max_row, c).font = bold
    ws1.append([])
    ws1.append(["说明：总金额=该期点数×单价＋该期奖金(满68奖3000当月一次)；"
                "找平金额=上月余量×上月单价(仅上半月，正补负扣)；"
                "应该付金额=总金额＋找平金额，照此发薪。"])

    ws2 = wb.create_sheet("日绩效明细")
    ws2.append(["日期", "员工编号", "姓名", "有效店数", "1点店", "2点店",
                "当日点"])
    for c in range(1, 8):
        ws2.cell(1, c).font = bold
    for d in daily:
        ws2.append([str(d["date"]), d["code"], d["name"], d["records"],
                    d["p1"], d["p2"], d["points"]])
    for ws in (ws1, ws2):
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col) + 2
            ws.column_dimensions[col[0].column_letter].width = width

    from fastapi.responses import StreamingResponse
    import io
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    from urllib.parse import quote
    po = "上半月" if period == "half1" else "下半月"
    ascii_name = "wage_%s_%s.xlsx" % (month or "all", period)
    fname = "发薪表_%s_%s.xlsx" % (po, month or "all")
    cd = "attachment; filename=" + ascii_name + "; filename*=UTF-8''" + quote(fname)
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": cd})


# ---------------- V3 月度对账 ----------------
@router.get("/recon", response_class=HTMLResponse)
@router.get("/v3/recon", response_class=HTMLResponse)
def v3_recon_page(request: Request,
                  user: Optional[User] = Depends(require_login),
                  db: Session = Depends(get_db), msg: str = "",
                  task_id: int = 0):
    if user is None or user.role != "admin":
        return _denied()
    from app.models import ReconDayRow, ReconResult, ReconTask
    tasks = (db.query(ReconTask)
             .filter(ReconTask.kind == "monthly_v3")
             .order_by(ReconTask.id.desc()).all())
    cur = db.get(ReconTask, task_id) if task_id else None
    diffs = (db.query(ReconResult).filter(ReconResult.task_id == task_id)
             .order_by(ReconResult.submitter_code).all()) if cur else []
    day_rows = []
    if cur:
        from app.models import Person
        _names = {p.code: p.display_name for p in db.query(Person).all()}
        day_rows = [{
            "ref_date": d.ref_date, "person_code": d.person_code,
            "name": _names.get(d.person_code, ""),
            "sys_points": d.sys_points or 0, "rep_points": d.rep_points or 0,
            "diff": d.diff or 0, "side": d.side, "note": d.note or ""}
            for d in (db.query(ReconDayRow)
                      .filter(ReconDayRow.task_id == task_id)
                      .order_by(ReconDayRow.ref_date,
                                ReconDayRow.person_code).all())]
    sys_rows = []
    diff_rows = []
    next_month = ""
    if cur:
        from app.services import v3_recon as vr
        amap = vr.task_adjust_map(db, task_id)
        next_month = vr._next_month((cur.params or {}).get("month", ""))
        for d in diffs:
            diff_rows.append({
                "code": d.submitter_code,
                "name": (d.note or "").replace(" 点数差异", "") or
                        d.submitter_code,
                "sys": d.system_value or 0, "rep": d.report_value or 0,
                "diff": d.diff or 0,
                "amount": vr._diff_amount(d.system_value or 0,
                                          d.report_value or 0),
                "adjusted": amap.get(d.submitter_code)})
        if (cur.summary or {}).get("sys_only"):
            from app.models import Person
            persons = {p.code: p.display_name for p in db.query(Person).all()}
            sys_rows = [{"code": c, "name": persons.get(c, c)}
                        for c in (cur.summary or {}).get("sys_only", [])]
    kind = (cur.summary or {}).get("kind") if cur else ""
    page_state = {
        "page": "recon", "cur_task": cur.id if cur else None,
        "cur_status": cur.status if cur else None,
        "month": (cur.params or {}).get("month", "") if cur else "",
        "tasks": [{"id": t.id, "month": (t.params or {}).get("month", ""),
                   "status": t.status,
                   "replaced_by": (t.params or {}).get("replaced_by")}
                  for t in tasks],
        "compared": (cur.summary or {}).get("compared", 0) if cur else 0,
        "diff_count": (cur.summary or {}).get("diff_count", 0) if cur else 0,
        "ai_ready": bool(cur and (cur.summary or {}).get("ai_interpret")),
    }
    return templates.TemplateResponse("v3_recon.html", {
        "request": request, "current_user": user, "msg": msg,
        "tasks": tasks, "cur": cur, "diffs": diff_rows,
        "sys_rows": sys_rows, "next_month": next_month,
        "day_rows": day_rows, "kind": kind, "page_state": page_state})


@router.post("/recon/upload2")
@router.post("/v3/recon/upload2")
async def v3_recon_upload2(request: Request, month: str = Form(""),
                           csrf_token: str = Form(...),
                           file=File(...),
                           user: Optional[User] = Depends(require_login),
                           db: Session = Depends(get_db)):
    if user is None or user.role != "admin":
        return _denied()
    if not csrf_ok(request, csrf_token):
        return HTMLResponse("CSRF 校验失败", status_code=400)
    content = await file.read()
    import os as _os
    try:
        from app.services import v3_recon
        t, older = v3_recon.submit_task(
            db, month, file.filename, content, user.id,
            sync=_os.environ.get("RECON_SYNC") == "1")
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, f"对账文件无法登记任务: {type(e).__name__}: {e}")
    msg = f"对账任务 # {t.id} 已提交，正在离线处理（完成后可下载结果）"
    if older:
        msg = (f"检测到 {month} 已有对账版本（{'、'.join('#'+str(x) for x in older)}），"
               f"已标记为「上一版」保留可查，本次为新对账；" + msg)
    if _os.environ.get("RECON_SYNC") == "1":
        msg = msg.replace("正在离线处理（完成后可下载结果）",
                          "已同步处理完成")
    from urllib.parse import quote as _quote
    return RedirectResponse(f"/v3/recon?msg={_quote(msg)}&task_id={t.id}",
                            status_code=303)



@router.post("/recon/{task_id}/interpret")
@router.post("/v3/recon/{task_id}/interpret")
def recon_interpret(task_id: int, request: Request,
                    csrf_token: str = Form(...),
                    user: Optional[User] = Depends(require_login),
                    db: Session = Depends(get_db)):
    """AI 对账解读（可选；未配置模型时提示，不影响对账）。"""
    if user is None or user.role != "admin":
        return _denied()
    if not csrf_ok(request, csrf_token):
        return HTMLResponse("CSRF 校验失败", status_code=400)
    from app.services import v3_recon
    res = v3_recon.interpret_task(db, task_id)
    msg = res.get("msg", "") or ("已生成 AI 解读" if res.get("ok") else "失败")
    return RedirectResponse(f"/v3/recon?task_id={task_id}&msg={msg}",
                            status_code=303)


@router.post("/recon/{task_id}/adjust/{code}")
@router.post("/v3/recon/{task_id}/adjust/{code}")
def recon_adjust(task_id: int, code: str, request: Request,
                 action: str = Form("add"),
                 csrf_token: str = Form(...),
                 user: Optional[User] = Depends(require_login),
                 db: Session = Depends(get_db)):
    """找平：对账差异确认 → 记入下月工资（加/减）；action=remove 取消。"""
    if user is None or user.role != "admin":
        return _denied()
    if not csrf_ok(request, csrf_token):
        return HTMLResponse("CSRF 校验失败", status_code=400)
    from app.services import v3_recon as vr
    if action == "add":
        res = vr.confirm_adjust(db, task_id, code, user.id)
        msg = (f"已确认找平：{res['amount']:+d} 円记入下月工资"
               if res["ok"] else f"无法确认：{res.get('msg', '?')}")
    elif action == "remove":
        vr.cancel_adjust(db, task_id, code)
        msg = "已取消该员工的找平确认"
    else:
        raise HTTPException(400, f"未知动作: {action}")
    return RedirectResponse(f"/v3/recon?task_id={task_id}&msg={msg}",
                            status_code=303)


@router.get("/recon/export")
@router.get("/v3/recon/export")
def v3_recon_export(request: Request, task_id: int = 0,
                    user: Optional[User] = Depends(require_login),
                    db: Session = Depends(get_db)):
    """导出对账差异 Excel：差异明细 + 反向名单（系统有而对账文件无）。"""
    if user is None or user.role != "admin":
        return _denied()
    from app.models import Person, ReconResult, ReconTask
    from openpyxl import Workbook
    from openpyxl.styles import Font
    cur = db.get(ReconTask, task_id) if task_id else None
    if cur is None:
        raise HTTPException(404, "对账任务不存在")
    rows = (db.query(ReconResult).filter(ReconResult.task_id == task_id)
            .order_by(ReconResult.submitter_code).all())
    summary = cur.summary or {}
    bold = Font(bold=True)
    wb = Workbook()
    # Sheet1 差异明细
    ws = wb.active
    ws.title = "差异明细"
    ws.append(["月份", (cur.params or {}).get("month", ""),
               "文件", (cur.params or {}).get("file", "")])
    ws.append(["比对人数", summary.get("compared", 0),
               "差异条数", summary.get("diff_count", 0)])
    ws.append([])
    if not rows:
        ws.append(["无差异：两侧点数一致"])
    else:
        ws.append(["人员", "编号", "系统点数", "对账点数", "差异(系统-对账)"])
        for c in range(1, 6):
            ws.cell(ws.max_row, c).font = bold
        for r in rows:
            ws.append([r.note.replace(" 点数差异", "") if r.note else "",
                       r.submitter_code, r.system_value, r.report_value,
                       r.diff])
    # Sheet2 反向名单
    ws2 = wb.create_sheet("系统有而对账文件无")
    so = summary.get("sys_only") or []
    if not so:
        ws2.append(["无（对账文件已覆盖系统全部有记录员工）"])
    else:
        ws2.append(["编号", "姓名"])
        ws2.cell(1, 1).font = bold
        ws2.cell(1, 2).font = bold
        persons = {p.code: p.display_name for p in db.query(Person).all()}
        for c in so:
            ws2.append([c, persons.get(c, c)])
    for wsx in (ws, ws2):
        for col in wsx.columns:
            w = max(len(str(c.value or "")) for c in col) + 2
            wsx.column_dimensions[col[0].column_letter].width = min(w, 40)
    from fastapi.responses import StreamingResponse
    import io
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    from urllib.parse import quote
    ascii_name = f"recon_diff_{task_id}.xlsx"
    fname = f"对账差异_任务{task_id}.xlsx"
    cd = "attachment; filename=" + ascii_name + "; filename*=UTF-8''" + quote(fname)
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": cd})


@router.get("/recon/report")
@router.get("/v3/recon/report")
def v3_recon_report(request: Request, task_id: int = 0,
                    user: Optional[User] = Depends(require_login),
                    db: Session = Depends(get_db)):
    """一键生成《月度对账报告.xlsx》（摘要+差异+反向名单+找平留痕）。"""
    if user is None or user.role != "admin":
        return _denied()
    from app.services import v3_recon
    wb = v3_recon.build_report(db, task_id, user.display_name)
    if wb is None:
        raise HTTPException(404, "对账任务不存在")
    from fastapi.responses import StreamingResponse
    import io
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    from urllib.parse import quote
    ascii_name = f"recon_report_{task_id}.xlsx"
    fname = f"月度对账报告_任务{task_id}.xlsx"
    cd = "attachment; filename=" + ascii_name + "; filename*=UTF-8''" + quote(fname)
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": cd})


@router.get("/recon/result")
@router.get("/v3/recon/result")
def v3_recon_result(request: Request, task_id: int = 0,
                    user: Optional[User] = Depends(require_login),
                    db: Session = Depends(get_db)):
    """下载对账任务产物 Excel（离线任务完成后落盘；历史任务兜底现算）。"""
    if user is None or user.role != "admin":
        return _denied()
    from app.models import ReconTask
    from fastapi.responses import StreamingResponse
    import io, os
    t = db.get(ReconTask, task_id) if task_id else None
    if t is None:
        raise HTTPException(404, "对账任务不存在")
    if t.status not in ("done", "parsed"):
        return RedirectResponse(f"/v3/recon?task_id={task_id}"
                                "&msg=任务尚未完成，请稍后刷新再下载",
                                status_code=303)
    bio = io.BytesIO()
    fp = (t.params or {}).get("result_path", "")
    if fp and os.path.exists(fp):
        with open(fp, "rb") as f:
            bio.write(f.read())
    else:
        from app.services import v3_recon
        wb = v3_recon.build_report(db, task_id, user.display_name)
        if wb is None:
            raise HTTPException(404, "对账任务不存在")
        wb.save(bio)
    bio.seek(0)
    from urllib.parse import quote
    ascii_name = f"recon_result_{task_id}.xlsx"
    fname = f"对账结果_任务{task_id}.xlsx"
    cd = "attachment; filename=" + ascii_name + "; filename*=UTF-8''" + quote(fname)
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": cd})


# ---------------- 管理端：月度重算（防同月补传双算） ----------------
def _month_of_modified(rr):
    return (rr.modified_raw or "")[:7]


@router.get("/month", response_class=HTMLResponse)
@router.get("/v3/month", response_class=HTMLResponse)
def v3_month_page(request: Request,
                  user: Optional[User] = Depends(require_login),
                  db: Session = Depends(get_db), month: str = "",
                  msg: str = "", err: str = ""):
    """月度数据健康度：该月 raw/有效/入表情况 + 补传警示 + 重建入口。"""
    if user is None or user.role != "admin":
        return _denied()
    from sqlalchemy import func
    from app.models import AppealRecord
    from app.services import v3_perf
    months = sorted({r[0] for r in db.query(
        func.substr(RawRecord.modified_raw, 1, 7)).all() if r[0]})
    if month not in months:
        month = months[-1] if months else ""
    info = {"raw": 0, "valid": 0, "filtered": 0, "dup": 0,
            "pend": 0, "formal": 0, "points": 0, "wage": 0, "files": []}
    if month:
        q = db.query(RawRecord).filter(
            RawRecord.modified_raw.like(month + "%"))
        info["raw"] = q.count()
        info["valid"] = q.filter(RawRecord.clean_status == "valid").count()
        info["filtered"] = q.filter(RawRecord.clean_status.in_(
            ("master_late", "from_sub"))).count()
        info["dup"] = q.filter(
            RawRecord.clean_status == "cross_file_dup").count()
        info["pend"] = (db.query(AppealRecord)
                        .join(RawRecord,
                              RawRecord.id == AppealRecord.raw_record_id)
                        .filter(AppealRecord.status == "pending",
                                RawRecord.modified_raw.like(month + "%"))
                        .count())
        frs = [f for f in db.query(FormalRecord).all()
               if (str(f.japan_date or ""))[:7] == month]
        info["formal"] = len(frs)
        info["points"] = sum(f.points or 0 for f in frs)
        mp = v3_perf.month_perf(db, month)
        info["wage"] = sum(v3_perf.salary_for(m["points"]) for m in mp)
        # 涉及文件
        fids = [r[0] for r in db.query(RawRecord.import_id).filter(
            RawRecord.modified_raw.like(month + "%")).distinct().all()]
        for fid in fids:
            f = db.get(ImportFile, fid)
            if f is None:
                continue
            vm = db.query(RawRecord).filter(
                RawRecord.import_id == fid,
                RawRecord.clean_status == "valid",
                RawRecord.modified_raw.like(month + "%")).count()
            fm = sum(1 for x in frs if x.import_id == fid)
            info["files"].append({"id": f.id, "name": f.file_name,
                                  "valid": vm, "formal": fm,
                                  "synced": vm == fm,
                                  "formalized": fm > 0})
        info["need_rebuild"] = any(
            not x["synced"] for x in info["files"])
    page_state = {"page": "month", "month": month, "info": info,
                  "need_rebuild": info.get("need_rebuild", False)}
    return templates.TemplateResponse("v3_month.html", {
        "request": request, "current_user": user, "month": month,
        "months": months, "info": info, "msg": msg, "err": err,
        "page_state": page_state})


@router.post("/month/rebuild")
@router.post("/v3/month/rebuild")
def v3_month_rebuild(request: Request, month: str = Form(""),
                     csrf_token: str = Form(...),
                     user: Optional[User] = Depends(require_login),
                     db: Session = Depends(get_db)):
    """按结算月重算并重建正式表（幂等；pending 申诉存在时拒绝）。"""
    if user is None or user.role != "admin":
        return _denied()
    if not csrf_ok(request, csrf_token):
        return HTMLResponse("CSRF 校验失败", status_code=400)
    from urllib.parse import quote as _q
    # 月份格式（含年份范围）校验的唯一关口在 v3_flow.rebuild_month，此处只消费结果
    res = v3_flow.rebuild_month(db, month, user.id)
    if not res["ok"]:
        return RedirectResponse(f"/v3/month?month={_q(month)}"
                                f"&err={_q(res.get('msg', '无法重算'))}",
                                status_code=303)
    return RedirectResponse(
        f"/v3/month?month={month}"
        f"&msg=已重算 {len(res['files'])} 个文件并重建正式表："
        f"条数 {res['formal_before']} → {res['formal_after']}，"
        f"点数 {res['points_before']} → {res['points_after']}",
        status_code=303)


# ---------------- 月度分期对账偏差表（薪资找平，独立功能页） ----------------
@router.get("/payroll-settle", response_class=HTMLResponse)
def payroll_settle_page(request: Request, user: Optional[User] =
                        Depends(require_login),
                        db: Session = Depends(get_db), month: str = "",
                        staff: str = ""):
    """薪资找平：月 × 员工 分期对账偏差表（独立页面，可按员工定位）。"""
    if user is None or user.role != "admin":
        return _denied()
    from app.services import v3_period
    from app.models import PayrollPeriodRow, PersonDailyStat
    months = sorted(
        {(str(r.ref_date or ""))[:7] for r in db.query(PersonDailyStat).all()
         if r.ref_date and (str(r.ref_date))[:7] >= "2026-07"}
        | {r.month for r in db.query(PayrollPeriodRow).all()
           if r.month >= "2026-07"})
    if month not in months:
        month = months[-1] if months else ""
    rows = v3_period.period_rows(db, month) if month else []
    # 员工定位：从绩效行「找平」按钮带着 staff 参数进来 → 该员工排最前（不隐藏其他人）
    if staff and rows:
        rows = sorted(rows, key=lambda r: (r["code"] != staff, r["code"]))
    total = {
        "half1": sum(r["half1"] for r in rows),
        "half2": sum(r["half2"] for r in rows),
        "bonus": sum(r["half1_bonus"] + r["half2_bonus"] for r in rows),
        "half_amt": sum(r["half1_amt"] + r["half2_amt"] for r in rows),
        "settle": sum(r["settle"] for r in rows),
        "settle_amt": sum(r["settle_amt"] for r in rows),
        "prev": sum(r["prev"] for r in rows),
        "prev_amt": sum(r["prev_amt"] for r in rows),
        "diff": sum(r["diff"] for r in rows),
        "diff_amt": sum(r["diff_amt"] for r in rows),
    }
    page_state = {"page": "payroll_settle", "month": month,
                  "rows": len(rows), "total": total,
                  "adjusted": sum(1 for r in rows if r["adj"])}
    from app.services import v3_perf as _vp
    bonus_g, bonus_a = _vp.bonus_params()
    return templates.TemplateResponse("payroll_settle.html", {
        "request": request, "current_user": user, "month": month,
        "months": months, "rows": rows, "total": total, "staff": staff,
        "bonus_g": bonus_g, "bonus_a": bonus_a,
        "page_state": page_state})


@router.post("/payroll-settle/generate")
def payroll_settle_generate(request: Request, month: str = Form(""),
                            csrf_token: str = Form(...),
                            user: Optional[User] = Depends(require_login),
                            db: Session = Depends(get_db)):
    if user is None or user.role != "admin":
        return _denied()
    if not csrf_ok(request, csrf_token):
        return HTMLResponse("CSRF 校验失败", status_code=400)
    from app.services import v3_period
    from urllib.parse import quote as _q
    try:
        res = v3_period.sync_period_table(db, month)
        return RedirectResponse(
            f"/payroll-settle?month={month}&msg={_q('已生成/更新 ' + str(res.get('rows', 0)) + ' 人')}",
            status_code=303)
    except Exception as e:  # noqa: BLE001
        return RedirectResponse(
            f"/payroll-settle?month={month}&msg={_q('生成失败: ' + str(e)[:80])}",
            status_code=303)


@router.post("/payroll-settle/{month}/{person_code}/update")
def payroll_settle_update(month: str, person_code: str, request: Request,
                          half1_amount: int = Form(0),
                          half2_amount: int = Form(0),
                          adjust_delta: int = Form(0),
                          csrf_token: str = Form(...),
                          user: Optional[User] = Depends(require_login),
                          db: Session = Depends(get_db)):
    if user is None or user.role != "admin":
        return _denied()
    if not csrf_ok(request, csrf_token):
        return HTMLResponse("CSRF 校验失败", status_code=400)
    from app.services import v3_period
    from urllib.parse import quote as _q
    res = v3_period.set_period_values(
        db, month, person_code, half1_amount, half2_amount,
        user_id=user.id, adjust_delta=adjust_delta)
    return RedirectResponse(
        f"/payroll-settle?month={month}&msg={_q(res.get('msg', '已保存'))}",
        status_code=303)


@router.get("/payroll-settle/export")
def payroll_settle_export(request: Request,
                          user: Optional[User] = Depends(require_login),
                          db: Session = Depends(get_db), month: str = ""):
    if user is None or user.role != "admin":
        return _denied()
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from fastapi.responses import StreamingResponse
    from app.services import v3_period
    wb = Workbook()
    ws = wb.active
    ws.title = "月度分期对账偏差"
    ws.append(["月份", "员工编号", "员工姓名", "上半月点数", "上半月金额(円)",
               "下半月点数", "下半月金额(円)", "奖金(円)", "分期已发(円)",
               "对账点数", "对账金额(円)", "上月修正(点)", "上月修正金额(円)",
               "对账偏差(点,参考)", "对账偏差金额(円,参考)",
               "找平(点)", "找平金额(円)"])
    for c in range(1, 18):
        ws.cell(1, c).font = Font(bold=True)
    for r in v3_period.period_rows(db, month):
        ws.append([month, r["code"], r["name"], r["half1"], r["half1_amt"],
                   r["half2"], r["half2_amt"],
                   r["half1_bonus"] + r["half2_bonus"],
                   r["half1_amt"] + r["half2_amt"],
                   r["settle"], r["settle_amt"], r["prev"], r["prev_amt"],
                   r["diff"], r["diff_amt"], r["adj"], r["adj_amt"]])
    ws.append([])
    ws.append(["说明：对账偏差=系统参考值(自动刷新)；找平=人工实际执行值(默认0，"
               "金额=点×该月单价)；上月修正=上月找平递延。"])
    import io as _io
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fn = f"payroll_settle_{month}.xlsx"
    return StreamingResponse(
        buf, media_type=("application/vnd.openxmlformats-officedocument"
                         ".spreadsheetml.sheet"),
        headers={"Content-Disposition": f'attachment; filename="{fn}"'})
