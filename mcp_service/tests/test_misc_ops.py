# -*- coding: utf-8 -*-
"""misc_ops 能力层测试（file_layout / product_doc / staff_set_status /
store_ai_run / export_recon_result）。

临时 SQLite（Base.metadata.create_all）造数据，绝不碰真实库。
覆盖：能力函数正常路径 + 错误/权限/确认语/不存在；信封层（经 scenario_ops 暴露）
UNAUTHORIZED / FORBIDDEN_TOOL / CONFIRM_REQUIRED / NOT_FOUND / INTERNAL。
"""
import asyncio
import base64
import io
import os
import types
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.models import (AiRun, Base, ImportFile, ReconTask, StoreEntity,
                        StorePair, User)
from app.services import store_master as _sm
from mcp_service import guards, misc_ops, scenario_ops, tokens

W = tokens.Actor(uid=1, role="admin", scopes=["read", "write"], token_id=1)
R = tokens.Actor(uid=1, role="admin", scopes=["read"], token_id=1)


def fake_ctx():
    """无 Authorization 头（stdio 风格）→ actor_from_ctx 直接返回 None。"""
    return types.SimpleNamespace(headers={})


def _seed(s, tmp_path):
    """基础种子：账号 / 文件（含解析布局）/ 对账任务（含落盘产物）/ B 组候选对。"""
    s.add_all([
        User(id=1, username="admin", password_hash="h1", display_name="管理员",
             role="admin", status="active", is_active=True),
        User(id=2, username="zhangsan", password_hash="h2", display_name="张三",
             role="staff", status="active", is_active=True, person_code="P001"),
        User(id=3, username="lisi", password_hash="h3", display_name="李四",
             role="staff", status="leave", is_active=True, person_code="P002"),
    ])
    s.add_all([
        ImportFile(id=1, file_name="2026-09巡店.xlsx", file_sha256="a" * 64,
                   file_size=10, stored_path="/tmp/a.xlsx", uploaded_by=1,
                   status="parsed", format="wide", header_row=2,
                   data_start_row=3, parsed_rows=3, total_rows=5,
                   layout={
                       "header_row": 2,
                       "cols": {"store_id": 1, "store_name": 2,
                                "modified_time": 3, "submitter": 4,
                                "visible": 6, "deploy": 7},
                       "value_map": {"visible": {"YES": "candidate"},
                                     "deploy": {"YES": "YES", "NO": "NO"}},
                       "point_rules": [
                           {"visible": ["YES"], "deploy": ["YES"], "points": 2},
                           {"visible": ["YES"], "deploy": ["NO"], "points": 1},
                       ],
                       "source": "ai",
                   },
                   parsed_sheets=["S1"], ignored_sheets=["S2"],
                   warnings=["表头第1行跳过"], errors=[]),
        ImportFile(id=2, file_name="2026-10空.xlsx", file_sha256="b" * 64,
                   file_size=10, stored_path="/tmp/b.xlsx", uploaded_by=1,
                   status="uploaded", format="wide"),
    ])
    # B 组候选：实体 1/2 为 fuzzy pending（供 ai_run 启动）
    s.add_all([
        StoreEntity(id=1, store_id_raw="S001", name_local="店A",
                    name_norm=_sm.norm_name("店A"), city="", master_id=1),
        StoreEntity(id=2, store_id_raw="S002", name_local="店A ",
                    name_norm=_sm.norm_name("店A "), city="", master_id=2),
    ])
    s.add(StorePair(id=1, entity_a=1, entity_b=2, kind="fuzzy",
                    status="pending"))
    # 对账任务：1=done 且有落盘产物；2=running（未完成不可下载）
    artifact = tmp_path / "artifact" / "recon_1.xlsx"
    artifact.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "对账结果"
    ws.append(["对账结果", "任务1"])
    wb.save(artifact)
    s.add_all([
        ReconTask(id=1, kind="person_points", status="done", created_by=1,
                  params={"month": "2026-09", "result_path": str(artifact)},
                  summary={"compared": 1}),
        ReconTask(id=2, kind="person_points", status="running", created_by=1,
                  params={"month": "2026-09"}),
    ])
    s.commit()


@pytest.fixture()
def db(tmp_path, monkeypatch):
    """能力层测试：临时 SQLite session（落盘目录重定向到 tmp）。"""
    monkeypatch.setenv("VISIT_MCP_EXPORT_DIR", str(tmp_path / "exports"))
    eng = create_engine(f"sqlite:///{tmp_path}/m.db",
                        connect_args={"check_same_thread": False})
    Base.metadata.create_all(eng)
    S = sessionmaker(bind=eng, expire_on_commit=False)
    s = S()
    _seed(s, tmp_path)
    yield s
    s.close()
    eng.dispose()


@pytest.fixture()
def factory(tmp_path, monkeypatch):
    """信封层测试：临时库 factory + 把 app.db.SessionLocal / scenario_ops.SessionLocal
    重定向过去。"""
    monkeypatch.setenv("VISIT_MCP_EXPORT_DIR", str(tmp_path / "exports"))
    eng = create_engine(f"sqlite:///{tmp_path}/env.db",
                        connect_args={"check_same_thread": False})
    Base.metadata.create_all(eng)
    F = sessionmaker(bind=eng, expire_on_commit=False)
    s = F()
    _seed(s, tmp_path)
    s.close()
    monkeypatch.setattr("app.db.SessionLocal", F)
    monkeypatch.setattr(scenario_ops, "SessionLocal", F)
    yield F
    eng.dispose()


# ---------------------------------------------------------------------------
# visit_file_layout
# ---------------------------------------------------------------------------

def test_file_layout_ok(db):
    d = misc_ops.file_layout(db, 1)
    assert d["file_id"] == 1 and d["file_name"] == "2026-09巡店.xlsx"
    assert d["status"] == "parsed" and d["format"] == "wide"
    assert d["header_row"] == 2 and d["data_start_row"] == 3
    assert d["layout_source"] == "ai"
    assert d["cols"]["visible"] == 6 and d["cols"]["deploy"] == 7
    assert d["value_map"]["visible"] == {"YES": "candidate"}
    assert d["value_map_text"]["visible"] == "YES=candidate"
    assert d["point_rules"][0] == {"visible": ["YES"], "deploy": ["YES"],
                                   "points": 2}
    assert d["point_rules_text"].startswith("YES&YES=2")
    assert d["parsed_sheets"] == ["S1"] and d["ignored_sheets"] == ["S2"]
    assert d["warnings"] == ["表头第1行跳过"] and d["errors"] == []
    assert "hint" not in d


def test_file_layout_no_layout_ok(db):
    """文件存在但无布局 → 合法结果 + hint（不是错误）。"""
    d = misc_ops.file_layout(db, 2)
    assert d["file_id"] == 2
    assert d["layout"] is None and d["point_rules"] is None
    assert d["parsed_sheets"] == [] and d["errors"] == []
    assert "hint" in d and "人工纠正" in d["hint"]


def test_file_layout_not_found_envelope(db):
    got = misc_ops._call(db, lambda db: misc_ops.file_layout(db, 999))
    assert got["ok"] is False
    assert got["error"]["code"] == "NOT_FOUND"
    assert "visit_files" in got["error"]["hint"]


# ---------------------------------------------------------------------------
# visit_product_doc
# ---------------------------------------------------------------------------

def test_product_doc_ok_real_file(db):
    """真实 app/product_doc.md（9005 字节 > 8000）→ 截断标注。"""
    d = misc_ops.product_doc(db)
    assert isinstance(d["doc"], str) and len(d["doc"]) > 0
    assert d["source"].endswith("product_doc.md")
    assert d["truncated"] is (d["total_chars"] > d["char_limit"])
    assert d["char_limit"] == 8000
    if d["truncated"]:
        assert "截断" in d["hint"]
        assert len(d["doc"]) == 8000


def test_product_doc_short_not_truncated(db, monkeypatch, tmp_path):
    p = tmp_path / "short.md"
    p.write_text("产品说明短文本", encoding="utf-8")
    monkeypatch.setattr(misc_ops, "_DOC_PATH", p)
    d = misc_ops.product_doc(db)
    assert d["doc"] == "产品说明短文本"
    assert d["truncated"] is False and d["total_chars"] == 7
    assert "hint" not in d


def test_product_doc_long_truncated(db, monkeypatch, tmp_path):
    p = tmp_path / "long.md"
    p.write_text("a" * 20000, encoding="utf-8")
    monkeypatch.setattr(misc_ops, "_DOC_PATH", p)
    d = misc_ops.product_doc(db)
    assert d["truncated"] is True
    assert len(d["doc"]) == 8000 and d["total_chars"] == 20000
    assert "截断" in d["hint"]


def test_product_doc_missing_file_not_found(db, monkeypatch, tmp_path):
    monkeypatch.setattr(misc_ops, "_DOC_PATH", tmp_path / "nope.md")
    with pytest.raises(misc_ops.MiscError) as e:
        misc_ops.product_doc(db)
    assert e.value.code == "NOT_FOUND"


# ---------------------------------------------------------------------------
# visit_staff_set_status（能力层）
# ---------------------------------------------------------------------------

def test_staff_set_status_ok(db):
    res = misc_ops.staff_set_status(db, W, user_id=2, new_status="leave",
                                    confirm_text="确认改状态 zhangsan leave")
    assert res["ok"] is True
    d = res["data"]
    assert d["user_id"] == 2 and d["username"] == "zhangsan"
    assert d["new_status"] == "leave" and d["status_label"] == "请假"
    assert d["is_active"] is True
    u = db.get(User, 2)
    assert u.status == "leave" and u.is_active is True


def test_staff_set_status_disabled_resigned_inactive(db):
    for st in ("disabled", "resigned"):
        res = misc_ops.staff_set_status(db, W, user_id=2, new_status=st,
                                        confirm_text=f"确认改状态 zhangsan {st}")
        assert res["data"]["new_status"] == st
        assert res["data"]["is_active"] is False
        assert db.get(User, 2).is_active is False


def test_staff_set_status_leave_reactivates(db):
    """disabled → leave：重新可登录（is_active 联动，与路由一致）。"""
    misc_ops.staff_set_status(db, W, user_id=2, new_status="disabled",
                              confirm_text="确认改状态 zhangsan disabled")
    assert db.get(User, 2).is_active is False
    misc_ops.staff_set_status(db, W, user_id=2, new_status="leave",
                              confirm_text="确认改状态 zhangsan leave")
    assert db.get(User, 2).is_active is True


def test_staff_set_status_idempotent(db):
    for _ in range(2):
        res = misc_ops.staff_set_status(db, W, user_id=2, new_status="leave",
                                        confirm_text="确认改状态 zhangsan leave")
        assert res["ok"] is True


def test_staff_set_status_bad_status(db):
    with pytest.raises(guards.GuardError) as e:
        misc_ops.staff_set_status(db, W, user_id=2, new_status="fired",
                                  confirm_text="确认改状态 zhangsan fired")
    assert e.value.code == "BAD_PARAM"
    assert "active" in e.value.hint


def test_staff_set_status_confirm_required(db):
    for text in (None, "改吧", "确认改状态 lisi leave"):
        with pytest.raises(guards.GuardError) as e:
            misc_ops.staff_set_status(db, W, user_id=2, new_status="leave",
                                      confirm_text=text)
        assert e.value.code == "CONFIRM_REQUIRED"
        # hint 给出准确确认语（含库里取出的 username）
        assert "确认改状态 zhangsan leave" in e.value.hint


def test_staff_set_status_not_found(db):
    with pytest.raises(guards.GuardError) as e:
        misc_ops.staff_set_status(db, W, user_id=999, new_status="leave",
                                  confirm_text="确认改状态 x leave")
    assert e.value.code == "NOT_FOUND"
    # admin 账号不可改（与网页路由 role!=staff → 404 同口径）
    with pytest.raises(guards.GuardError) as e:
        misc_ops.staff_set_status(db, W, user_id=1, new_status="leave",
                                  confirm_text="确认改状态 admin leave")
    assert e.value.code == "NOT_FOUND"


def test_staff_set_status_requires_write(db):
    with pytest.raises(guards.GuardError) as e:
        misc_ops.staff_set_status(db, R, user_id=2, new_status="leave",
                                  confirm_text="确认改状态 zhangsan leave")
    assert e.value.code == "FORBIDDEN_TOOL"
    with pytest.raises(guards.GuardError) as e:
        misc_ops.staff_set_status(db, None, user_id=2, new_status="leave",
                                  confirm_text="确认改状态 zhangsan leave")
    assert e.value.code == "UNAUTHORIZED"


# ---------------------------------------------------------------------------
# visit_store_ai_run（能力层）
# ---------------------------------------------------------------------------

def test_store_ai_run_started(db, monkeypatch):
    monkeypatch.setattr("app.services.ai_batch.configured", lambda: True)
    res = misc_ops.store_ai_run(db, W)
    assert res["ok"] is True
    d = res["data"]
    assert d["started"] is True and d["run_id"] is not None
    assert d["total_pairs"] == 1 and d["status"] == "running"
    assert "visit_store" in d["hint"]
    run = db.get(AiRun, d["run_id"])
    assert run is not None and run.status == "running"
    assert run.total_pairs == 1 and run.created_by == 1


def test_store_ai_run_already_running_rejected(db, monkeypatch):
    monkeypatch.setattr("app.services.ai_batch.configured", lambda: True)
    db.add(AiRun(id=99, total_pairs=5, created_by=1))   # status 默认 running
    db.commit()
    with pytest.raises(guards.GuardError) as e:
        misc_ops.store_ai_run(db, W)
    assert e.value.code == "BAD_PARAM"
    assert "已有" in e.value.message
    assert "运行" in e.value.hint and "done" in e.value.hint
    assert db.query(AiRun).count() == 1                 # 未新增


def test_store_ai_run_not_configured(db, monkeypatch):
    monkeypatch.setattr("app.services.ai_batch.configured", lambda: False)
    res = misc_ops.store_ai_run(db, W)
    assert res["ok"] is True
    d = res["data"]
    assert d["started"] is False and d["reason"] == "ai_not_configured"
    assert "AI_API_KEY" in d["hint"]
    assert db.query(AiRun).count() == 0


def test_store_ai_run_no_pending(db, monkeypatch):
    monkeypatch.setattr("app.services.ai_batch.configured", lambda: True)
    db.query(StorePair).delete()
    db.commit()
    res = misc_ops.store_ai_run(db, W)
    assert res["ok"] is True
    assert res["data"]["started"] is False
    assert res["data"]["reason"] == "no_pending_pairs"
    assert db.query(AiRun).count() == 0


def test_store_ai_run_requires_write(db):
    with pytest.raises(guards.GuardError) as e:
        misc_ops.store_ai_run(db, R)
    assert e.value.code == "FORBIDDEN_TOOL"


# ---------------------------------------------------------------------------
# visit_export_recon_result（能力层）
# ---------------------------------------------------------------------------

def test_export_recon_result_artifact(db):
    d = misc_ops.export_recon_result(db, 1)
    assert d["filename"] == "对账结果_任务1.xlsx"
    assert d["size"] > 0
    raw = base64.b64decode(d["content_base64"])
    assert len(raw) == d["size"]
    src = (db.get(ReconTask, 1).params or {}).get("result_path")
    assert d["source_path"] == str(Path(src).resolve())
    with open(src, "rb") as f:
        assert f.read() == raw
    p = Path(d["saved_path"])
    assert p.is_absolute() and p.exists()
    assert p.read_bytes() == raw
    assert "content_base64" in d["hint"]


def test_export_recon_result_recompute(db):
    """产物缺失 → 按路由逻辑现算（recon.build_report），source_path=None。"""
    src = (db.get(ReconTask, 1).params or {}).get("result_path")
    os.remove(src)
    d = misc_ops.export_recon_result(db, 1)
    assert d["source_path"] is None
    assert "现算" in d["hint"]
    raw = base64.b64decode(d["content_base64"])
    assert len(raw) == d["size"] and len(raw) > 0
    wb = load_workbook(io.BytesIO(raw))
    assert wb.sheetnames[0] == "对账报告摘要"
    p = Path(d["saved_path"])
    assert p.exists() and p.read_bytes() == raw


def test_export_recon_result_task_not_found(db):
    with pytest.raises(misc_ops.MiscError) as e:
        misc_ops.export_recon_result(db, 999)
    assert e.value.code == "NOT_FOUND"


def test_export_recon_result_not_done(db):
    """任务未完成（status=running）→ NOT_FOUND + 尚未完成。"""
    with pytest.raises(misc_ops.MiscError) as e:
        misc_ops.export_recon_result(db, 2)
    assert e.value.code == "NOT_FOUND"
    assert "尚未完成" in e.value.message


def test_export_recon_result_cannot_recompute(db, monkeypatch):
    """产物缺失且现算失败 → NOT_FOUND + hint（无法现算）。"""
    src = (db.get(ReconTask, 1).params or {}).get("result_path")
    os.remove(src)
    monkeypatch.setattr(misc_ops, "_recompute",
                        lambda db, task_id, author_name="": None)
    with pytest.raises(misc_ops.MiscError) as e:
        misc_ops.export_recon_result(db, 1)
    assert e.value.code == "NOT_FOUND"
    assert "无法现算" in e.value.message


# ---------------------------------------------------------------------------
# 信封层（经 _write_call / 只读包装）：UNAUTHORIZED / FORBIDDEN_TOOL /
# CONFIRM_REQUIRED / NOT_FOUND / ok / INTERNAL / INTERNAL 映射
# ---------------------------------------------------------------------------

def test_envelope_unauthorized_without_token(factory):
    # 写工具：无身份 → UNAUTHORIZED
    got = scenario_ops.visit_staff(fake_ctx(), action="set_status",
                                   username="zhangsan", status="leave",
                                   confirm_text="确认改状态 zhangsan leave")
    assert got["ok"] is False and got["error"]["code"] == "UNAUTHORIZED"
    # 只读工具：无身份（stdio 风格）放行
    got = scenario_ops.visit_files(fake_ctx(), view="layout", import_id=1)
    assert got["ok"] is True
    assert got["data"]["header_row"] == 2
    actor, err = misc_ops._authorize(fake_ctx())
    assert actor is None and err["ok"] is False
    assert err["error"]["code"] == "UNAUTHORIZED"


def test_envelope_forbidden_for_read_only_token(factory, monkeypatch):
    monkeypatch.setattr("mcp_service.tools.actor_from_ctx", lambda ctx: R)
    got = scenario_ops.visit_staff(fake_ctx(), action="set_status",
                                   username="zhangsan", status="leave",
                                   confirm_text="确认改状态 zhangsan leave")
    assert got["ok"] is False and got["error"]["code"] == "FORBIDDEN_TOOL"


def test_envelope_staff_set_status_ok(factory, monkeypatch):
    monkeypatch.setattr("mcp_service.tools.actor_from_ctx", lambda ctx: W)
    got = scenario_ops.visit_staff(fake_ctx(), action="set_status",
                                   username="zhangsan", status="leave",
                                   confirm_text="确认改状态 zhangsan leave")
    assert got["ok"] is True
    assert got["data"]["new_status"] == "leave"
    s = factory()
    assert s.get(User, 2).status == "leave"
    assert s.get(User, 2).is_active is True
    s.close()


def test_envelope_staff_set_status_confirm_required(factory, monkeypatch):
    monkeypatch.setattr("mcp_service.tools.actor_from_ctx", lambda ctx: W)
    got = scenario_ops.visit_staff(fake_ctx(), action="set_status",
                                   username="zhangsan", status="leave",
                                   confirm_text="改吧")
    assert got["ok"] is False and got["error"]["code"] == "CONFIRM_REQUIRED"
    assert "确认改状态 zhangsan leave" in got["error"]["hint"]


def test_envelope_staff_set_status_unknown_username(factory, monkeypatch):
    monkeypatch.setattr("mcp_service.tools.actor_from_ctx", lambda ctx: W)
    got = scenario_ops.visit_staff(fake_ctx(), action="set_status",
                                   username="nobody", status="leave",
                                   confirm_text="确认改状态 nobody leave")
    assert got["ok"] is False and got["error"]["code"] == "NOT_FOUND"


def test_envelope_write_internal_retryable(factory, monkeypatch):
    """写工具 retryable=True → 意外异常必须 INTERNAL（可重试），绝不 INTERNAL_WRITE。"""
    monkeypatch.setattr("mcp_service.tools.actor_from_ctx", lambda ctx: W)

    def boom(*a, **k):
        raise RuntimeError("boom-db-down")

    monkeypatch.setattr(guards, "require_write", boom)
    got = scenario_ops.visit_staff(fake_ctx(), action="set_status",
                                   username="zhangsan", status="leave",
                                   confirm_text="确认改状态 zhangsan leave")
    assert got["ok"] is False and got["error"]["code"] == "INTERNAL"
    assert "boom-db-down" in got["error"]["message"]
    assert "可重试" in got["error"]["hint"]


def test_envelope_file_layout_ok(factory, monkeypatch):
    monkeypatch.setattr("mcp_service.tools.actor_from_ctx", lambda ctx: W)
    got = scenario_ops.visit_files(fake_ctx(), view="layout", import_id=1)
    assert got["ok"] is True
    assert got["data"]["header_row"] == 2
    assert got["data"]["value_map_text"]["visible"] == "YES=candidate"


def test_envelope_recon_export_detail_ok(factory, monkeypatch):
    monkeypatch.setattr(scenario_ops, "actor_from_ctx", lambda ctx: W)
    got = scenario_ops.visit_recon_export(fake_ctx(), kind="detail", task_id=1)
    assert got["ok"] is True
    d = got["data"]
    assert d["filename"] == "对账结果_任务1.xlsx"
    raw = base64.b64decode(d["content_base64"])
    assert len(raw) == d["size"]
    assert Path(d["saved_path"]).exists()


def test_envelope_recon_export_detail_not_found(factory, monkeypatch):
    monkeypatch.setattr(scenario_ops, "actor_from_ctx", lambda ctx: W)
    got = scenario_ops.visit_recon_export(fake_ctx(), kind="detail", task_id=999)
    assert got["ok"] is False and got["error"]["code"] == "NOT_FOUND"
    assert "hint" in got["error"]


def test_call_maps_unexpected_to_internal(db):
    def boom(_db):
        raise RuntimeError("boom-x")

    got = misc_ops._call(db, boom)
    assert got["ok"] is False and got["error"]["code"] == "INTERNAL"
    assert "boom-x" in got["error"]["message"]
    assert "可重试" in got["error"]["hint"]

    def misc(_db):
        raise misc_ops.MiscError("BAD_PARAM", "参数错", "h")

    got = misc_ops._call(db, misc)
    assert got["ok"] is False and got["error"]["code"] == "BAD_PARAM"


# ---------------------------------------------------------------------------
# 注册（能力层不注册工具；工具注册在 scenario_ops）
# ---------------------------------------------------------------------------

def test_no_register_in_misc_ops():
    assert not hasattr(misc_ops, "register")
