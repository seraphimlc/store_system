# -*- coding: utf-8 -*-
"""P1 导出 MCP 工具（visit_export_*）与 report 层重构等价性测试。

- 工具成功路径：返回的 content_base64 能解码成合法 xlsx（openpyxl 可打开且 sheet ≥1）、
  filename 合理、saved_path 文件真实存在且大小一致；
- 错误路径：BAD_MONTH / NOT_FOUND / UNAUTHORIZED；
- 重构等价性：build_* 函数输出与「改造前路由」的逐单元格快照一致
  （快照由 scripts/p1_export_baseline.py capture 生成；报告「生成时间」行逐次不同，
  比对时归一化剔除——与任务给定的比对方式一致）。

fixture 用本地 create_engine（不碰 app.db 全局引擎，见 test_month_summary 注释）。
"""
import base64
import io
import json
import types
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.auth import hash_password
from app.models import (Base, FormalRecord, ImportFile, MonthPerfRecord,
                        PayrollPeriodRow, Person, RawRecord, ReconDayRow,
                        ReconResult, ReconTask, User)

SNAPSHOT = Path(__file__).resolve().parent / "data" / "export_baseline.json"
BASELINE = json.loads(SNAPSHOT.read_text(encoding="utf-8"))
MONTH = "2026-09"


def _seed(db) -> int:
    """与 scripts/p1_export_baseline.py 完全相同的种子（等价性测试依赖一致性）。"""
    admin = User(username="admin", password_hash=hash_password("pw123456"),
                 display_name="管理员", role="admin", is_active=True)
    db.add(admin)
    db.flush()
    db.add_all([
        Person(code="P001", display_name="甲"),
        Person(code="P002", display_name="乙"),
    ])
    imp = ImportFile(file_name="2026-09巡店.xlsx", file_sha256="b" * 64,
                     stored_path="x", uploaded_by=admin.id)
    db.add(imp)
    db.flush()
    raws = [
        RawRecord(import_id=imp.id, sheet_name="S1", excel_row=2,
                  store_id_raw="ST-A", store_name_local_raw="店A",
                  modified_raw="2026-09-01 09:00:00",
                  submitter_raw="甲(P001)", submitter_code="P001",
                  visible_raw="YES", deploy_raw="YES"),
        RawRecord(import_id=imp.id, sheet_name="S1", excel_row=3,
                  store_id_raw="ST-B", store_name_local_raw="店B",
                  modified_raw="2026-09-20 09:00:00",
                  submitter_raw="甲(P001)", submitter_code="P001",
                  visible_raw="YES", deploy_raw="YES"),
        RawRecord(import_id=imp.id, sheet_name="S1", excel_row=4,
                  store_id_raw="ST-C", store_name_local_raw="店C",
                  modified_raw="2026-09-05 09:00:00",
                  submitter_raw="乙(P002)", submitter_code="P002",
                  visible_raw="YES", deploy_raw="NO"),
    ]
    db.add_all(raws)
    db.flush()
    db.add_all([
        FormalRecord(import_id=imp.id, raw_record_id=raws[0].id,
                     person_code="P001", japan_date=date(2026, 9, 1), points=1),
        FormalRecord(import_id=imp.id, raw_record_id=raws[1].id,
                     person_code="P001", japan_date=date(2026, 9, 20), points=2),
        FormalRecord(import_id=imp.id, raw_record_id=raws[2].id,
                     person_code="P002", japan_date=date(2026, 9, 5), points=1),
    ])
    db.add_all([
        MonthPerfRecord(month="2026-09", person_code="P001", records=2,
                        p1=1, p2=1, points=3, salary=750, per_point=250),
        MonthPerfRecord(month="2026-09", person_code="P002", records=1,
                        p1=1, p2=0, points=1, salary=250, per_point=250),
    ])
    db.add_all([
        PayrollPeriodRow(month="2026-09", person_code="P001",
                         half1_points=1, half2_points=2,
                         half1_records=1, half1_p1=1, half1_p2=0,
                         half2_records=1, half2_p1=0, half2_p2=1,
                         settle_points=3, prev_adjust_points=0, diff_points=0,
                         half1_bonus=0, half2_bonus=0,
                         half1_amount=250, half2_amount=500,
                         settle_amount=750, prev_adjust_amount=0,
                         diff_amount=0, adjust_points=0, adjust_amount=0),
        PayrollPeriodRow(month="2026-09", person_code="P002",
                         half1_points=1, half2_points=0,
                         half1_records=1, half1_p1=1, half1_p2=0,
                         half2_records=0, half2_p1=0, half2_p2=0,
                         settle_points=1, prev_adjust_points=0, diff_points=0,
                         half1_bonus=0, half2_bonus=0,
                         half1_amount=250, half2_amount=0,
                         settle_amount=250, prev_adjust_amount=0,
                         diff_amount=0, adjust_points=0, adjust_amount=0),
    ])
    task = ReconTask(kind="monthly_v3", status="done", created_by=admin.id,
                     params={"month": "2026-09", "file": "对账_2026-09.xlsx"},
                     summary={"kind": "person_points", "compared": 2,
                              "diff_count": 1, "monthly_diff": 1,
                              "sys_only": ["P002"]})
    db.add(task)
    db.flush()
    db.add(ReconResult(task_id=task.id, submitter_code="P001",
                       system_value=3, report_value=5, diff=-2,
                       status="diff", family="person_points",
                       confirmed=False, note="甲 点数差异"))
    db.add(ReconDayRow(task_id=task.id, ref_date=date(2026, 9, 1),
                       person_code="P001", sys_points=1, rep_points=2,
                       sys_count=1, rep_count=1, diff=-1, side="both",
                       note=""))
    db.commit()
    return task.id


@pytest.fixture()
def db(tmp_path, monkeypatch):
    """临时 SQLite（全量 schema）+ 最小数据；落盘目录重定向到 tmp。"""
    monkeypatch.setenv("VISIT_MCP_EXPORT_DIR", str(tmp_path / "exports"))
    eng = create_engine(f"sqlite:///{tmp_path}/e.db",
                        connect_args={"check_same_thread": False})
    Base.metadata.create_all(eng)
    S = sessionmaker(bind=eng, expire_on_commit=False)
    s = S()
    s.task_id = _seed(s)
    yield s
    s.close()
    eng.dispose()


def _cells(raw: bytes) -> dict:
    wb = load_workbook(io.BytesIO(raw), data_only=True)
    return {ws.title: [[None if v is None else v for v in row]
                       for row in ws.iter_rows(values_only=True)]
            for ws in wb.worksheets}


def _norm(rows):
    """报告「生成时间」行逐次不同 → 归一化后比较（与基线脚本一致）。"""
    return [r for r in rows if not (r and r[0] == "生成时间")]


# ---------------- 工具成功路径 ----------------

def test_salary_export_payload_ok(db):
    from mcp_service import export_ops
    res = export_ops._run_export(db, "salary", month=MONTH, period="half1")
    assert res["ok"] is True
    d = res["data"]
    assert d["filename"] == "发薪表_上半月_2026-09.xlsx"
    assert d["size"] > 0
    raw = base64.b64decode(d["content_base64"])
    assert len(raw) == d["size"]
    wb = load_workbook(io.BytesIO(raw))
    assert len(wb.sheetnames) >= 1
    assert wb.sheetnames[0] == "上半月发薪1-15·2026-09"
    p = Path(d["saved_path"])
    assert p.is_absolute() and p.exists()
    assert p.read_bytes() == raw
    assert p.stat().st_size == d["size"]
    assert "content_base64" in d["hint"]


def test_salary_export_half2_filename_and_title(db):
    from mcp_service import export_ops
    d = export_ops._run_export(db, "salary", month=MONTH,
                               period="half2")["data"]
    assert d["filename"] == "发薪表_下半月_2026-09.xlsx"
    wb = load_workbook(io.BytesIO(base64.b64decode(d["content_base64"])))
    assert wb.sheetnames[0] == "下半月发薪16-月末·2026-09"


def test_settle_export_payload_ok(db):
    from mcp_service import export_ops
    res = export_ops._run_export(db, "payroll_settle", month=MONTH)
    assert res["ok"] is True
    d = res["data"]
    assert d["filename"] == "payroll_settle_2026-09.xlsx"
    raw = base64.b64decode(d["content_base64"])
    wb = load_workbook(io.BytesIO(raw))
    assert wb.sheetnames == ["月度分期对账偏差"]
    ws = wb["月度分期对账偏差"]
    # 表头(1) + 数据(2) + 空行 + 说明 = 5 行 × 17 列
    assert ws.max_row == 5 and ws.max_column == 17
    assert ws["A2"].value == "2026-09" and ws["B2"].value == "P001"
    p = Path(d["saved_path"])
    assert p.exists() and p.stat().st_size == d["size"]


def test_recon_diff_export_payload_ok(db):
    from mcp_service import export_ops
    res = export_ops._run_export(db, "recon_diff", task_id=db.task_id)
    assert res["ok"] is True
    d = res["data"]
    assert d["filename"] == f"对账差异_任务{db.task_id}.xlsx"
    wb = load_workbook(io.BytesIO(base64.b64decode(d["content_base64"])))
    assert wb.sheetnames == ["差异明细", "系统有而对账文件无"]
    ws = wb["差异明细"]
    cells = [str(c) for row in ws.iter_rows(values_only=True) for c in row]
    assert "甲" in cells and "-2" in cells
    assert Path(d["saved_path"]).exists()


def test_recon_report_export_payload_ok(db):
    from mcp_service import export_ops
    res = export_ops._run_export(db, "recon_report", task_id=db.task_id)
    assert res["ok"] is True
    d = res["data"]
    assert d["filename"] == f"月度对账报告_任务{db.task_id}.xlsx"
    wb = load_workbook(io.BytesIO(base64.b64decode(d["content_base64"])))
    assert set(wb.sheetnames) == {
        "对账报告摘要", "员工×日对账明细", "差异明细",
        "系统有而对账文件无", "找平确认"}
    assert Path(d["saved_path"]).exists()


# ---------------- 错误路径 ----------------

def test_bad_month_returns_bad_month_envelope(db):
    from mcp_service import export_ops
    for kind in ("salary", "payroll_settle"):
        for bad in ("2026-9", "２０２６-09", "2026-13", "2026-09-01", ""):
            res = export_ops._run_export(db, kind, month=bad)
            assert res["ok"] is False
            assert res["error"]["code"] == "BAD_MONTH", (kind, bad)


def test_month_without_data_returns_not_found(db):
    from mcp_service import export_ops
    res = export_ops._run_export(db, "salary", month="2026-08")
    assert res["ok"] is False and res["error"]["code"] == "NOT_FOUND"
    res = export_ops._run_export(db, "payroll_settle", month="2026-08")
    assert res["ok"] is False and res["error"]["code"] == "NOT_FOUND"


def test_missing_task_returns_not_found(db):
    from mcp_service import export_ops
    for kind in ("recon_diff", "recon_report"):
        res = export_ops._run_export(db, kind, task_id=999999)
        assert res["ok"] is False
        assert res["error"]["code"] == "NOT_FOUND", kind


def test_unauthorized_envelope_without_bearer():
    from mcp_service import export_ops
    fake_ctx = types.SimpleNamespace(headers={})
    actor, err = export_ops._authorize(fake_ctx)
    assert actor is None
    assert err["ok"] is False and err["error"]["code"] == "UNAUTHORIZED"


# ---------------- 重构等价性（build_* 输出 vs 改造前路由快照） ----------------

def test_build_functions_match_pre_refactor_baseline(db):
    """对同一份数据，build_* 的单元格输出必须与改造前路由完全一致。"""
    from app.services import report
    cases = [
        ("salary_half1",
         report.build_payroll_workbook(db, MONTH, "half1")[0]),
        ("salary_half2",
         report.build_payroll_workbook(db, MONTH, "half2")[0]),
        ("settle",
         report.build_payroll_settle_workbook(db, MONTH)[0]),
        ("recon_diff",
         report.build_recon_diff_workbook(db, db.task_id)[0]),
        # 路由传 user.display_name（基线的「生成人」= 管理员）
        ("recon_report",
         report.build_recon_report_workbook(db, db.task_id, "管理员")[0]),
    ]
    for kind, raw in cases:
        got = _cells(raw)
        exp = BASELINE[kind]
        assert set(got) == set(exp), f"{kind}: sheet 名不一致 {set(got)} vs {set(exp)}"
        for sn in exp:
            assert _norm(got[sn]) == _norm(exp[sn]), (
                f"{kind}/{sn}: 单元格与改造前不一致\n"
                f"  before={_norm(exp[sn])!r}\n  after ={_norm(got[sn])!r}")


def test_build_functions_return_none_for_missing_task(db):
    from app.services import report
    assert report.build_recon_diff_workbook(db, 999999) is None
    assert report.build_recon_report_workbook(db, 999999) is None
