# -*- coding: utf-8 -*-
"""员工登录名导出 + 「首登强制改密」流程（发号前的整套检查）。"""
from datetime import date

import app.db as appdb
from app.auth import SESSION_COOKIE, hash_password, read_session_token
from app.models import Person, User
from tests.helpers import form_token


def _mk(client_list, username, name, code, role="staff", status="active",
        active=True, must_change=False, password="demo123"):
    db = appdb.SessionLocal()
    if code and db.get(Person, code) is None:
        db.add(Person(code=code, display_name=name))
    if db.query(User).filter(User.username == username).first() is None:
        db.add(User(username=username, password_hash=hash_password(password),
                    display_name=name, role=role, person_code=code,
                    status=status, is_active=active,
                    must_change_password=must_change, lang=None))
    db.commit()
    db.close()


def _login(client, username, password="demo123"):
    return client.post("/login", data={"username": username, "password": password},
                       follow_redirects=False)


def _csrf(client):
    return (read_session_token(client.cookies.get(SESSION_COOKIE)) or {}).get("csrf", "")


def _read_xlsx(data):
    import io as _io

    from openpyxl import load_workbook
    return load_workbook(_io.BytesIO(data))


# ---------- 导出 ----------

def test_export_accounts_admin_only(client):
    _mk(client, "admin", "管理员", None, role="admin")
    _mk(client, "empA", "员工A", "PA")
    _login(client, "admin")
    r = client.get("/staff-admin/accounts-export")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    wb = _read_xlsx(r.content)
    assert wb.sheetnames == ["员工登录名", "使用说明"]
    rows = list(wb["员工登录名"].values)
    assert rows[0][:4] == ("登录名", "姓名", "员工编号", "在职状态")
    assert ("empA", "员工A", "PA") == tuple(rows[1][:3])


def test_export_only_active_variant(client):
    """active=1 → 只导在岗且启用的；离职员工不出现。"""
    _mk(client, "admin", "管理员", None, role="admin")
    _mk(client, "empA", "在岗甲", "PA")
    _mk(client, "empB", "离职乙", "PB", status="resigned", active=False)
    _login(client, "admin")
    all_names = [r[0] for r in list(_read_xlsx(client.get(
        "/staff-admin/accounts-export").content)["员工登录名"].values)][1:]
    act_names = [r[0] for r in list(_read_xlsx(client.get(
        "/staff-admin/accounts-export?active=1").content)["员工登录名"].values)][1:]
    assert "empA" in all_names and "empB" in all_names
    assert "empA" in act_names and "empB" not in act_names


def test_export_writes_password_only_when_must_change(client):
    """初始口令列：待改密的写默认口令；已自设密码的留空。"""
    _mk(client, "admin", "管理员", None, role="admin")
    _mk(client, "empMust", "待改密", "PM", must_change=True)
    _mk(client, "empDone", "已自设", "PD", must_change=False)
    _login(client, "admin")
    rows = {r[0]: r for r in list(_read_xlsx(client.get(
        "/staff-admin/accounts-export").content)["员工登录名"].values)[1:]}
    assert rows["empMust"][5] == "是" and rows["empMust"][6] == "demo123"
    assert not rows["empDone"][5] and not rows["empDone"][6]   # 空（openpyxl 读出 None）


def test_export_blocked_for_staff(client):
    _mk(client, "empA", "员工A", "PA")
    _login(client, "empA")
    r = client.get("/staff-admin/accounts-export", follow_redirects=False)
    # 员工首页 = 待填报出勤计划 /my/plan，否则每日自报 /my/report（date_plan.staff_home）
    assert r.status_code == 302 and r.headers["location"].startswith(
        ("/my/plan", "/my/report", "/my/perf"))


# ---------- 首登强制改密 ----------

def test_first_login_forces_password_change(client):
    """重置后（must_change=1）：除改密页与登出外一律被拦到改密页；改完恢复正常。"""
    _mk(client, "empA", "员工A", "PA", must_change=True)
    assert _login(client, "empA").status_code == 302
    for path in ("/my/perf", "/my/report", "/my/report/feedback"):
        r = client.get(path, follow_redirects=False)
        assert r.status_code == 302 and r.headers["location"].startswith(
            "/my/password?must=1"), path
    # 旧口令不对 → 留在改密页
    bad = client.post("/my/password", data={"old_password": "wrong",
                                            "new_password": "newpass123",
                                            "csrf_token": _csrf(client),
                                            "_ft": form_token(client)},
                      follow_redirects=False)
    assert bad.status_code == 400 and "原密码不对" in bad.text   # 原密码错误：400 + 提示
    # 正确改密 → 放行
    ok = client.post("/my/password", data={"old_password": "demo123",
                                           "new_password": "newpass123",
                                           "csrf_token": _csrf(client),
                                           "_ft": form_token(client)},
                     follow_redirects=False)
    assert ok.status_code == 303
    db = appdb.SessionLocal()
    u = db.query(User).filter(User.username == "empA").one()
    assert u.must_change_password is False
    db.close()
    assert client.get("/my/perf", follow_redirects=False).status_code == 200
    # 新口令可登录、旧口令不行
    client.get("/logout")
    assert _login(client, "empA", "newpass123").status_code == 302
    client.get("/logout")
    assert _login(client, "empA", "demo123").status_code == 401   # 老口令已失效（401 + 登录页）


def test_reset_all_sets_default_and_must_change(client):
    """批量重置：全部员工口令 → 默认口令 + must_change=True；管理员不受影响。"""
    _mk(client, "admin", "管理员", None, role="admin", password="adminpw")
    _mk(client, "empA", "甲", "PA", must_change=False, password="ownpass")
    _mk(client, "empB", "乙", "PB", must_change=False, password="ownpass")
    _login(client, "admin", "adminpw")
    r = client.post("/staff-admin/reset-all",
                    data={"csrf_token": _csrf(client), "_ft": form_token(client),
                          "password": "demo123"}, follow_redirects=False)
    assert r.status_code == 303 and "msg=" in r.headers["location"]
    db = appdb.SessionLocal()
    for name in ("empA", "empB"):
        u = db.query(User).filter(User.username == name).one()
        assert u.must_change_password is True
    assert db.query(User).filter(User.username == "admin").one() \
        .must_change_password is False               # 管理员不动
    db.close()
    # 用默认口令能登录，且被强制去改密
    assert _login(client, "empA", "demo123").status_code == 302
    r2 = client.get("/my/perf", follow_redirects=False)
    assert r2.status_code == 302 and r2.headers["location"].startswith("/my/password")


def test_export_login_url_prefers_public_issuer(client, monkeypatch):
    """登录地址要用公开地址（应用在 nginx 后面，request.base_url 会给出 http://）。"""
    from types import SimpleNamespace

    import app.config as app_config
    _mk(client, "admin", "管理员", None, role="admin")
    _mk(client, "empA", "员工A", "PA")
    _login(client, "admin")
    # 服务里是函数内 `from app.config import get_settings` → 打到真正的来源上
    monkeypatch.setattr(app_config, "get_settings", lambda: SimpleNamespace(
        default_staff_password="demo123",
        visit_oauth_issuer="https://store.visitworld.me"))
    rows = list(_read_xlsx(client.get(
        "/staff-admin/accounts-export").content)["员工登录名"].values)[1:]
    assert all(r[8].startswith("https://store.visitworld.me/") for r in rows)


def test_cli_create_employee_is_idempotent_by_code(client):
    """同一编号只能有一个账号（历史 bug：CLI 只按登录名查重 → 同一人两条账号）。"""
    from app.cli import create_employee
    _mk(client, "xiaochuanyi", "小川逸", "2188240606634082")   # 模拟已有账号

    # 用一个"不同的登录名"再建一次同编号 → 必须复用，不新建
    u = create_employee("ogawa", "2188240606634082")
    db = appdb.SessionLocal()
    n = (db.query(User)
         .filter(User.person_code == "2188240606634082").count())
    assert n == 1, "同一编号出现了 %d 条账号" % n
    assert u.username == "xiaochuanyi"                   # 复用了已有账号
    db.close()


def test_delete_file_purges_empty_month_derived_data(client):
    """删文件后：该月已无正式记录 → 派生数据（月绩效/日统计/应发工资/看板/分析）必须清掉。

    历史坑（2026-09-29 线上实测）：删掉测试文件后，11 月的派生行仍在库里。
    """
    from app.models import (DashMetric, FormalRecord, ImportFile, MonthPerfRecord,
                            PayrollPeriodRow, PersonDailyStat, RawRecord,
                            StaffAnalysis, StoreEntity)
    from app.services import importer
    db = appdb.SessionLocal()
    imp = ImportFile(file_name="t.xlsx", file_sha256="sha-purge", file_size=1,
                     stored_path="/tmp/t.xlsx", uploaded_by=1, status="parsed")
    db.add(imp)
    db.commit()
    raw = RawRecord(import_id=imp.id, sheet_name="s", excel_row=2,
                    store_id_raw="TESTSTORE1", submitter_raw="甲(PA)",
                    submitter_code="PA")
    db.add(raw)
    db.commit()
    db.add(FormalRecord(import_id=imp.id, raw_record_id=raw.id, person_code="PA",
                        store_id_raw="TESTSTORE1", japan_date=date(2026, 11, 1),
                        points=1))
    db.add(StoreEntity(store_id_raw="TESTSTORE1", name_local="测试店",
                       name_norm="测试店", master_id=0,
                       first_seen_import_id=imp.id))
    # 该月派生数据（模拟上传时自动同步出来的）
    db.add(PersonDailyStat(person_code="PA", ref_date=date(2026, 11, 1),
                           records=1, p1=1, p2=0, points=1))
    db.add(MonthPerfRecord(month="2026-11", person_code="PA", records=1, p1=1,
                           p2=0, points=1, salary=250))
    db.add(PayrollPeriodRow(month="2026-11", person_code="PA"))
    db.add(DashMetric(month="2026-11", metric="checkins", value=1))
    db.add(StaffAnalysis(month="2026-11", person_code="PA", content="{}"))
    db.commit()
    ent_id = db.query(StoreEntity.id).filter(
        StoreEntity.store_id_raw == "TESTSTORE1").scalar()
    imp_id = imp.id
    db.close()

    db = appdb.SessionLocal()
    imp = db.get(ImportFile, imp_id)
    importer.delete_file(imp, db)
    db.close()

    db = appdb.SessionLocal()
    for model, args in ((PersonDailyStat, {}), (MonthPerfRecord, {}),
                        (PayrollPeriodRow, {}), (DashMetric, {}),
                        (StaffAnalysis, {})):
        assert db.query(model).count() == 0, model.__name__ + " 还有残留"
    assert db.query(StoreEntity).filter(StoreEntity.id == ent_id).count() == 0
    assert db.query(FormalRecord).count() == 0
    db.close()


def test_export_includes_leaders(client):
    """⚠️ 2026-10-06：账号表导出原来只导 `role == 'staff'` → **队长不在里面**。

    管理端「员工管理」页早就是 `("staff","leader")`（能看到并管理队长），导出漏了 ——
    用户口径："队长不也是staff嘛"。
    """
    from io import BytesIO
    from openpyxl import load_workbook
    _mk(None, "admin2", "管理员", "ADM2", role="admin")
    _mk(None, "emp9", "普通员工", "E9")
    _mk(None, "lead9", "陈嘉溢", "L9", role="leader")
    _login(client, "admin2")
    r = client.get("/staff-admin/accounts-export")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert "员工登录名" in wb.sheetnames
    vals = [list(x) for x in wb["员工登录名"].iter_rows(values_only=True)]
    text = " ".join(str(c) for row in vals for c in row if c)
    assert "陈嘉溢" in text, "队长必须出现在账号表里：%s" % vals[:3]
    assert "普通员工" in text
