# -*- coding: utf-8 -*-
"""AI 布局解析（巡店表头通用能力）与 loader 坐标模式测试。"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import pytest
from openpyxl import Workbook


def _make(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "任意sheet名"
    for r in rows:
        ws.append(r)
    wb.save(path)


@pytest.fixture()
def variant_file(tmp_path):
    """变体表头：标题行 + 日文列名乱序 + 非模板 sheet。"""
    p = str(tmp_path / "variant.xlsx")
    _make(p, [
        ["2026年9月 巡店数据", "", "", "", "", "", ""],
        ["POP展開", "店名", "店舗ID", "担当者", "巡店日時", "審査ステータス", "メモ"],
        ["YES", "A店", "01010470920260915001", "山田(111)",
         "2026-09-03 09:15:00", "AUDIT_SUCCESS", ""],
        ["", "B店", "01010470920260915002", "李四(222)",
         "2026-09-05 10:00:00", "OTHER", ""],
        ["YES", "C店", "01010470920260915003", "王五(333)",
         "2026-09-08 11:30:00", "AUDIT_FAILED", ""],
    ])
    return p


def test_loader_ai_layout_mode(variant_file):
    """loader AI 坐标模式：不要求表头含 'Store ID'，按 AI 行列号 + 值语义解析。"""
    from store_settle.loader import load_workbook
    layout = {"header_row": 2,
              "cols": {"store_id": 3, "store_name": 2, "modified_time": 5,
                       "submitter": 4, "visible": 6, "deploy": 1},
              "value_map": {"visible": {"AUDIT_SUCCESS": "candidate",
                                        "OTHER": "candidate",
                                        "AUDIT_FAILED": "blank"},
                            "deploy": {"YES": 2, "": 1}}}
    res = load_workbook(variant_file, col_override=layout)
    assert not res.failed and res.parsed_rows == 3
    r0 = res.rows[0]
    assert r0.store_id_raw == "01010470920260915001"
    assert r0.store_name_local_raw == "A店"
    assert r0.modified_raw == "2026-09-03 09:15:00"
    assert r0.submitter_code == "111"
    assert r0.visible_raw == "AUDIT_SUCCESS"  # raw 保留原值
    assert r0.deploy_raw == "YES"
    assert res.rows[1].visible_raw == "OTHER"
    assert res.rows[2].visible_raw == "AUDIT_FAILED"


def test_loader_value_map_controls_visibility(variant_file):
    """值语义完全由 value_map 决定（不写死）：改成 candidate 则 FAILED 也算有效。"""
    from store_settle.loader import load_workbook
    layout = {"header_row": 2,
              "cols": {"store_id": 3, "store_name": 2, "modified_time": 5,
                       "submitter": 4, "visible": 6, "deploy": 1},
              "value_map": {"visible": {"AUDIT_SUCCESS": "candidate",
                                        "OTHER": "candidate",
                                        "AUDIT_FAILED": "candidate"},
                            "deploy": {"YES": 2, "": 1}}}
    res = load_workbook(variant_file, col_override=layout)
    assert res.rows[2].visible_raw == "AUDIT_FAILED"  # 原值保留
    # 候选判定由 value_map 决定（judge 阶段）
    from store_settle.rules import visible_is_candidate
    vm_blank = {"AUDIT_SUCCESS": "candidate", "OTHER": "candidate",
                "AUDIT_FAILED": "blank"}
    assert visible_is_candidate("AUDIT_SUCCESS", vm_blank) is True
    assert visible_is_candidate("AUDIT_FAILED", vm_blank) is False
    vm_all = {"AUDIT_SUCCESS": "candidate", "OTHER": "candidate",
              "AUDIT_FAILED": "candidate"}
    assert visible_is_candidate("AUDIT_FAILED", vm_all) is True   # 口径可配置


def test_ai_parse_visit_layout_monkeypatched(variant_file, monkeypatch):
    """ai_parse_visit_layout：模型返回布局 → 结构校验/可信校验；失败回退 {}。"""
    import app.services.ai_visit as av
    from app.services.ai_visit import ai_parse_visit_layout

    monkeypatch.setattr(av, "_ai_configured", lambda: True)
    monkeypatch.setattr(av, "_chat", lambda prompt, **kw: (
        '{"header_row": 1, "store_id": 2, "store_name": 1, '
        '"modified_time": 4, "submitter": 3, "visible": 5, "deploy": 0, '
        '"record_id": null}'))
    out = ai_parse_visit_layout(variant_file)
    assert out["header_row"] == 2
    assert out["cols"] == {"store_id": 3, "store_name": 2, "modified_time": 5,
                           "submitter": 4, "visible": 6, "deploy": 1}
    assert out["source"] == "ai"

    # 模型漏报必需字段 → 不可信 → {}
    monkeypatch.setattr(av, "_chat", lambda prompt, **kw: (
        '{"header_row": 1, "store_id": 2, "store_name": 1, '
        '"modified_time": null, "submitter": null, "visible": 5, '
        '"deploy": 0}'))
    assert ai_parse_visit_layout(variant_file) == {}

    # 模型返回乱码 → {}
    monkeypatch.setattr(av, "_chat", lambda prompt, **kw: "not json")
    assert ai_parse_visit_layout(variant_file) == {}

    # 未配置 AI → {}
    monkeypatch.setattr(av, "_ai_configured", lambda: False)
    assert ai_parse_visit_layout(variant_file) == {}


def test_rule_fallback_still_works(tmp_path):
    """规则兜底（无 AI）：既有 8 月表头 + Review status 变体仍可解析。"""
    from store_settle.loader import load_workbook
    p = str(tmp_path / "review.xlsx")
    _make(p, [
        ["Store ID", "Store Name-Local", "Modified Time", "Submitter",
         "Review status", "Deploy New A+POSM"],
        ["01010470920260915001", "店A", "2026-09-03 09:15:00", "甲(111)",
         "AUDIT_SUCCESS", "YES"],
    ])
    res = load_workbook(p)
    assert not res.failed and res.parsed_rows == 1
    assert res.rows[0].visible_raw == "AUDIT_SUCCESS"  # 原值保留（候选在 judge 判定）


def test_point_for_combination_rules():
    """点数组合规则：AUDIT_FAILED+非YES 不计成绩；FAILED+YES 仍 2 点。"""
    from datetime import date
    from store_settle.rules import point_for
    b = date(2026, 7, 9)
    rules = [
        {"visible": ["OTHER", "AUDIT_SUCCESS"], "deploy": ["YES"], "points": 2},
        {"visible": ["OTHER", "AUDIT_SUCCESS"], "deploy": ["NO", ""], "points": 1},
        {"visible": ["AUDIT_FAILED"], "deploy": ["YES"], "points": 2},
        {"visible": ["AUDIT_FAILED"], "deploy": ["NO", ""], "points": 0},
    ]
    assert point_for(date(2026, 9, 1), "YES", b, "AUDIT_SUCCESS", rules) == 2
    assert point_for(date(2026, 9, 1), "", b, "AUDIT_SUCCESS", rules) == 1
    assert point_for(date(2026, 9, 1), "YES", b, "OTHER", rules) == 2
    assert point_for(date(2026, 9, 1), "YES", b, "AUDIT_FAILED", rules) == 2
    assert point_for(date(2026, 9, 1), "NO", b, "AUDIT_FAILED", rules) == 0
    assert point_for(date(2026, 9, 1), "", b, "AUDIT_FAILED", rules) == 0
    # 未覆盖组合 → 0
    assert point_for(date(2026, 9, 1), "?", b, "OTHER", rules) == 0
    # 无规则 → 默认口径（boundary 后 deploy=YES→2）
    assert point_for(date(2026, 9, 1), "YES", b) == 2
    assert point_for(date(2026, 9, 1), "NO", b) == 1


def test_delete_import_removes_formal_rows(client):
    """删除已入表的文件：正式表行一并删除（PG 外键约束下的正确行为）。"""
    import app.db as appdb
    from app.models import FormalRecord, ImportFile, RawRecord
    from app.services import importer
    db = appdb.SessionLocal()
    imp = ImportFile(file_name="t.xlsx", file_sha256="sha-del-test",
                     file_size=1, stored_path="/tmp/none.xlsx", uploaded_by=1,
                     status="parsed", parsed_sheets=[], ignored_sheets=[],
                     warnings=[], errors=[])
    db.add(imp)
    db.commit()
    rr = RawRecord(import_id=imp.id, sheet_name="s", excel_row=2,
                   store_id_raw="S1", store_name_local_raw="店A",
                   modified_raw="2026-09-01 10:00:00",
                   submitter_raw="甲(111)", submitter_code="111",
                   clean_status="valid")
    db.add(rr)
    db.commit()
    from datetime import date as _date
    db.add(FormalRecord(import_id=imp.id, raw_record_id=rr.id,
                        person_code="111", store_id_raw="S1",
                        japan_date=_date(2026, 9, 1), points=1))
    db.commit()
    importer.delete_file(imp, db)
    assert db.query(RawRecord).filter(RawRecord.import_id == imp.id).count() == 0
    assert db.query(FormalRecord).filter(
        FormalRecord.import_id == imp.id).count() == 0
    assert db.get(ImportFile, imp.id) is None
    db.close()


def test_zero_point_rows_not_counted(client):
    """0 点行（规则"不计成绩"）：只计店数、不计点，月绩效/日统计一致。"""
    import app.db as appdb
    from datetime import date as _d
    from app.models import (FormalRecord, ImportFile, MonthPerfRecord,
                            Person, PersonDailyStat, RawRecord)
    from app.services import perf
    db = appdb.SessionLocal()
    db.add(Person(code="911", display_name="零点儿"))
    imp = ImportFile(file_name="z.xlsx", file_sha256="sha-zero-pt",
                     file_size=1, stored_path="/tmp/none.xlsx", uploaded_by=1,
                     status="parsed", parsed_sheets=[], ignored_sheets=[],
                     warnings=[], errors=[])
    db.add(imp)
    db.commit()
    for i, pts in enumerate((2, 1, 0, 0), start=1):
        rr = RawRecord(import_id=imp.id, sheet_name="s", excel_row=i,
                       store_id_raw=f"S{i}", store_name_local_raw=f"店{i}",
                       modified_raw=f"2026-09-0{i} 10:00:00",
                       submitter_raw="甲(911)", submitter_code="911",
                       clean_status="valid")
        db.add(rr)
        db.commit()
        db.add(FormalRecord(import_id=imp.id, raw_record_id=rr.id,
                            person_code="911", store_id_raw=f"S{i}",
                            japan_date=_d(2026, 9, i), points=pts))
    db.commit()
    perf.sync_month_stats(db, "2026-09")
    mp = db.query(MonthPerfRecord).filter(
        MonthPerfRecord.month == "2026-09",
        MonthPerfRecord.person_code == "911").first()
    assert mp is not None
    assert mp.records == 4          # 店数含 0 点行
    assert mp.points == 2 + 1       # 点数不含 0 点行
    st = db.query(PersonDailyStat).filter(
        PersonDailyStat.person_code == "911").all()
    assert sum(s.points for s in st) == 3
    db.close()


def test_extract_json_tolerant():
    """模型输出容错解析：代码块/前后杂文/推理模型思考文本里的 JSON。"""
    from app.services.ai_chat import extract_json
    assert extract_json('{"a": 1}') == {"a": 1}
    assert extract_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert extract_json('好的，结果如下：\n{"a": 1}\n以上') == {"a": 1}
    # 推理模型：思考里带 JSON（前后有自然语言）
    assert extract_json('我们根据要求返回JSON。 header_row=0 ... {"header_row": 1, "store_id": 0} 完成') \
        == {"header_row": 1, "store_id": 0}
    assert extract_json("") is None
    assert extract_json("没有 JSON") is None


def test_ai_visit_reasoning_model_output(monkeypatch, tmp_path):
    """推理模型（content 空 / 思考里带 JSON）也能解析出布局。"""
    from openpyxl import Workbook
    import app.services.ai_visit as av
    p = str(tmp_path / "t.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append(["Store ID", "Store Name-Local", "Modified Time", "Submitter",
               "Review status", "Deploy New A+POSM"])
    ws.append(["S1", "店A", "2026-09-01 10:00:00", "甲(111)",
               "AUDIT_SUCCESS", "YES"])
    wb.save(p)
    monkeypatch.setattr(av, "_ai_configured", lambda: True)
    calls = {"n": 0}

    def fake_chat(prompt, **kw):
        calls["n"] += 1
        if "header_row" in prompt:      # 第一步：列识别
            return ('先分析…… {"header_row": 0, "store_id": 0, "store_name": 1, '
                    '"modified_time": 2, "submitter": 3, "visible": 4, "deploy": 5}')
        return ('{"point_rules": [{"visible": ["AUDIT_SUCCESS"], '
                '"deploy": ["YES"], "points": 2}]}')

    monkeypatch.setattr(av, "_chat", fake_chat)
    out = av.ai_parse_visit_layout(p)
    assert calls["n"] >= 1                  # 列识别至少一次（该文件无规则页）
    assert out["cols"]["store_id"] == 1
    assert out["cols"]["visible"] == 5
    assert out["source"] == "ai"


def test_rules_sheet_prose_reaches_the_prompt(tmp_path):
    """规则页写成「计2点 / 暂不计分」时，整段都要进提示词。

    2026-10 文件的规则页不再用等号。截短或只挑含「点/=」的行，会把
    「AUDITING 暂不计分」这条切掉，模型就套系统默认口径。
    """
    from openpyxl import Workbook
    from app.services.ai_visit import _rules_prompt
    from app.services.recon import _sheet_preview
    p = str(tmp_path / "oct.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "STORE_TASK_EXCEL_SHEET "
    ws.append(["Store Information"])
    ws.append(["Store ID", "Store Name-Local", "Modified Time", "Submitter",
               "Sales Activity", "Review status", "Deploy New A+POSM"])
    rule = wb.create_sheet("规则")
    rule.append(["请读取第一个Sheet的数据，第2行为表头，第3行起为店铺记录。"])
    rule.append([None])
    rule.append(["审核状态使用“Review status”列，新物料投放使用“Deploy New A+POSM”列，按以下规则计算："])
    rule.append(["1. Review status为AUDIT_SUCCESS，且Deploy New A+POSM为YES的，计2点。"])
    rule.append(["4. Review status为AUDITING、NOT_REQUEST、空白或其他未列明状态的，暂不计分，列为待复核。"])
    wb.save(p)
    prompt = _rules_prompt(_sheet_preview(p))
    assert "计2点" in prompt
    assert "暂不计分" in prompt
    assert "Deploy New A+POSM" in prompt


def test_rules_from_sheet_overrides_model_guess(tmp_path, monkeypatch):
    """文件规则页写全四条时，以页上条文为准，模型少给的不能覆盖。"""
    from datetime import date
    from openpyxl import Workbook
    import app.services.ai_visit as av
    from store_settle.rules import point_for
    p = str(tmp_path / "oct.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append(["Store ID", "Store Name-Local", "Modified Time", "Submitter",
               "Sales Activity", "Review status", "Deploy New A+POSM"])
    ws.append(["S1", "店A", "2026-10-01 10:00:00", "甲(111)",
               "POSM", "AUDIT_SUCCESS", "YES"])
    rule = wb.create_sheet("规则")
    rule.append(["1. Review status为AUDIT_SUCCESS，且Deploy New A+POSM为YES的，计2点。"])
    rule.append(["2. Review status为AUDIT_SUCCESS，且Deploy New A+POSM为NO或空白的，计1点。"])
    rule.append(["3. Review status为AUDIT_FAILED或OTHER的，无论是否投放新物料，均计0点。"])
    rule.append(["4. Review status为AUDITING、NOT_REQUEST、空白或其他未列明状态的，暂不计分。"])
    wb.save(p)
    monkeypatch.setattr(av, "_ai_configured", lambda: True)

    def fake_chat(prompt, **kw):
        if "header_row" in prompt:
            return ('{"header_row": 0, "store_id": 0, "store_name": 1, '
                    '"modified_time": 2, "submitter": 3, "visible": 5, "deploy": 6}')
        # 模型漏掉「暂不计分」，还把 OTHER+YES 猜成 2 点
        return ('{"point_rules": [{"visible": ["AUDIT_SUCCESS", "OTHER"], '
                '"deploy": ["YES"], "points": 2}]}')

    monkeypatch.setattr(av, "_chat", fake_chat)
    out = av.ai_parse_visit_layout(p)
    assert out["rules_source"] == "sheet"
    b = date(2026, 7, 9)
    rules = out["point_rules"]
    assert point_for(date(2026, 10, 1), "YES", b, "AUDIT_SUCCESS", rules) == 2
    assert point_for(date(2026, 10, 1), "", b, "AUDIT_SUCCESS", rules) == 1
    assert point_for(date(2026, 10, 1), "YES", b, "OTHER", rules) == 0
    assert point_for(date(2026, 10, 1), "NO", b, "AUDITING", rules) == 0
    # 半截规则页不采用，避免盖掉默认口径
    rule["A4"] = "4. 其他状态另议。"
    wb.save(p)
    assert av.rules_from_sheet(p) == []


def test_rules_from_sheet_reads_points_from_text_not_hardcoded(tmp_path):
    """**点数必须从条文里读**，不能写死（2026-10-16 线上事故）。

    真实 10 月文件的规则页把「AUDIT_SUCCESS + YES」从 9 月的 2 点改成了 **1 点**。
    老实现把「计2点」硬写进正则 → 对这份文件**永远读不出来** → 返回 [] →
    回退给模型现猜。模型每次猜的规则可能不同，而 0 点的行不进正式表，
    于是同一份文件导两次「有效条数」就差十几条（用户上报的现象）。
    """
    from openpyxl import Workbook
    from datetime import date
    import app.services.ai_visit as av
    from store_settle.rules import point_for
    p = str(tmp_path / "real_oct.xlsx")
    wb = Workbook()
    wb.active.append(["Store ID", "Store Name-Local", "Modified Time"])
    rule = wb.create_sheet("规则")
    # ↓↓↓ 与线上文件逐字一致（含「计1点」）
    rule.append(["请读取第一个Sheet的数据，第2行为表头，第3行起为店铺记录。"])
    rule.append([])
    rule.append(["审核状态使用“Review status”列，新物料投放使用“Deploy New A+POSM”列，"
                 "按以下规则计算："])
    rule.append(["1. Review status为AUDIT_SUCCESS，且Deploy New A+POSM为YES的，计1点。"])
    rule.append(["2. Review status为AUDIT_SUCCESS，且Deploy New A+POSM为NO或空白的，计1点。"])
    rule.append(["3. Review status为AUDIT_FAILED或OTHER的，无论是否投放新物料，"
                 "均计0点，不计入有效成绩及报酬。"])
    rule.append(["4. Review status为AUDITING、NOT_REQUEST、空白或其他未列明状态的，"
                 "暂不计分，列为待复核。"])
    wb.save(p)

    rules = av.rules_from_sheet(p)
    assert rules, "四条齐全的真实规则页必须能读出来（老实现读不出）"
    assert point_for(date(2026, 10, 1), "YES", date(2026, 7, 9), "AUDIT_SUCCESS", rules) == 1
    assert point_for(date(2026, 10, 1), "NO", date(2026, 7, 9), "AUDIT_SUCCESS", rules) == 1
    assert point_for(date(2026, 10, 1), "", date(2026, 7, 9), "AUDIT_SUCCESS", rules) == 1
    # AUDIT_FAILED / OTHER 无论投放都是 0 点 → 不进正式表
    for dep in ("YES", "NO", ""):
        assert point_for(date(2026, 10, 1), dep, date(2026, 7, 9),
                         "AUDIT_FAILED", rules) == 0
        assert point_for(date(2026, 10, 1), dep, date(2026, 7, 9),
                         "OTHER", rules) == 0
    # 待复核的状态不在规则里 → 0 点
    assert point_for(date(2026, 10, 1), "YES", date(2026, 7, 9), "AUDITING", rules) == 0
    # 规则里**不能**出现 2 点（这份文件就是 1 点）
    assert all(r["points"] != 2 for r in rules), rules


def test_default_rules_user_口径():
    """用户口径固化：新物料YES=2点；S1∈{OTHER,AUDIT_SUCCESS}=1点；其它不计点。"""
    from datetime import date
    from store_settle.rules import point_for
    from app.services.ai_visit import default_layout
    rules = default_layout().get("point_rules")
    assert rules, "默认口径应内建"
    b = date(2026, 7, 9)
    d = date(2026, 9, 1)
    # ① 新物料=YES → 2 点（任意 S1）
    assert point_for(d, "YES", b, "AUDIT_SUCCESS", rules) == 2
    assert point_for(d, "YES", b, "OTHER", rules) == 2
    assert point_for(d, "YES", b, "AUDIT_FAILED", rules) == 2
    assert point_for(d, "YES", b, "NOT_REQUEST", rules) == 2
    # ② S1 ∈ {OTHER, AUDIT_SUCCESS} 且非 YES → 1 点
    assert point_for(d, "NO", b, "AUDIT_SUCCESS", rules) == 1
    assert point_for(d, "", b, "OTHER", rules) == 1
    # ③ 其它 S1（FAILED/NOT_REQUEST/未知）且非 YES → 不计点
    assert point_for(d, "NO", b, "AUDIT_FAILED", rules) == 0
    assert point_for(d, "", b, "AUDIT_FAILED", rules) == 0
    assert point_for(d, "", b, "NOT_REQUEST", rules) == 0
    assert point_for(d, "NO", b, "WHATEVER", rules) == 0


def test_norm_name_dedup(client):
    """9月起按归一化店名判重：同店不同写法合并为一家（只留 YES 那条）。"""
    import app.db as appdb
    from app.models import ImportFile, Person, RawRecord
    from app.services import flow
    db = appdb.SessionLocal()
    db.add(Person(code="771", display_name="甲"))
    imp = ImportFile(file_name="n.xlsx", file_sha256="sha-norm-dedup",
                     file_size=1, stored_path="/tmp/none.xlsx", uploaded_by=1,
                     status="parsed", parsed_sheets=[], ignored_sheets=[],
                     warnings=[], errors=[])
    db.add(imp)
    db.commit()
    # 同店两种写法：较早无投放、较晚 YES
    for i, (nm, dep, mt) in enumerate([
            ("WineShop Sommelier", "", "2026-09-02 10:00:00"),
            ("Wine shop sommelier", "YES", "2026-09-05 10:00:00")], start=1):
        db.add(RawRecord(import_id=imp.id, sheet_name="s", excel_row=i,
                         store_id_raw=f"S{i}", store_name_local_raw=nm,
                         modified_raw=mt, submitter_raw="甲(771)",
                         submitter_code="771", visible_raw="AUDIT_SUCCESS",
                         deploy_raw=dep, clean_status="pending"))
    db.commit()
    flow.judge_import(db, imp)
    db.commit()
    rows = db.query(RawRecord).filter(RawRecord.import_id == imp.id).all()
    valid = [r for r in rows if r.clean_status == "valid"]
    assert len(valid) == 1                      # 归一化后同一家店只留一条
    assert (valid[0].deploy_raw or "") == "YES"  # 且留的是 YES 那条（2 点）
    db.close()
