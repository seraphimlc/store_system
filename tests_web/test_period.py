# -*- coding: utf-8 -*-
"""月度分期对账偏差表（薪资找平）：生成公式/上月修正递延/手改保留/页面。"""
from datetime import date, datetime

import app.db as appdb
from app.models import (PayrollAdjust, PayrollPayment, PayrollPeriodRow,
                        Person, PersonDailyStat, ReconDataRow, ReconTask)
from app.services import period


def test_sync_formula_and_prev_adjust(client):
    """偏差 = 对账 − (上半月+下半月) + 上月修正；找平自动=金额差；上月修正=链式余额。"""
    db = appdb.SessionLocal()
    db.add(Person(code="111", display_name="甲"))
    # 上月（7月）结转余额 -500（要扣，上月工资没扣完）；金额差=500
    db.add(PayrollPeriodRow(month="2026-07", person_code="111",
                            half1_points=0, half2_points=0,
                            settle_points=0, prev_adjust_points=0,
                            diff_points=2, adjust_points=2,
                            diff_amount=500, adjust_amount=500,
                            prev_adjust_amount=-500,
                            updated_at=datetime.utcnow()))
    # 结转的权威来源=账表（period_rows 的「上月修正」读它，不再读链子）
    db.add(PayrollAdjust(source_month="2026-07", person_code="111",
                         adjust_amount=-500, settled_amount=0, remaining=-500,
                         status="in_progress", updated_at=datetime.utcnow()))
    # 8月统计表：8/1=1点(上半月)，8/16=2点+8/20=1点(下半月3点)
    for d, pts in ((date(2026, 8, 1), 1), (date(2026, 8, 16), 2),
                   (date(2026, 8, 20), 1)):
        db.add(PersonDailyStat(person_code="111", ref_date=d, points=pts))
    # 8月对账任务（当前）+ 对账数据表 甲=10 点
    t = ReconTask(kind="monthly_v3", status="done", created_by=1,
                  params={"month": "2026-08", "file": "t.xlsx"})
    db.add(t)
    db.commit()
    db.add(ReconDataRow(task_id=t.id, ref_date=date(2026, 8, 1),
                        person_code="111", person_name="甲", points=10, cnt=1))
    db.commit()
    res = period.sync_period_table(db, "2026-08")
    assert res["rows"] == 1
    rows = period.period_rows(db, "2026-08")
    r = rows[0]
    assert r["half1"] == 1 and r["half2"] == 3      # 上半月1 / 下半月3
    assert r["settle"] == 10                          # 对账点数（全量）
    assert r["prev"] == 0                       # 上月余量点数(参考；已自动找平→0)
    assert r["prev_amt"] == -500                 # 上月结转余额-500(要扣)经链式递延
    assert r["diff"] == 10 - (1 + 3) + 0 == 6   # 偏差点数(参考；上月余量0)
    assert r["diff_amt"] == 2500 - (250 + 750) == 1500  # 金额差(含奖金)=对账−系统
    assert r["adj_amt"] == r["diff_amt"] == 1500  # 找平自动=金额差(无点击)
    db.close()


def test_prev_adjust_chain_carry(client):
    """链式递延：本月结转 = −本月金额差 + 本月两期工资吸收上月结转后的剩余(<0才递延)。"""
    db = appdb.SessionLocal()
    db.add(Person(code="A", display_name="甲"))
    db.add(Person(code="B", display_name="乙"))
    # 上月结转：A=-25,000(要扣,本月工资吸收不完)、B=0(无结转)
    db.add(PayrollPeriodRow(month="2026-07", person_code="A",
                            diff_points=100, adjust_points=100,
                            diff_amount=100 * 250, adjust_amount=100 * 250,
                            half1_amount=10000, half2_amount=0,
                            prev_adjust_amount=-25000,
                            updated_at=datetime.utcnow()))
    db.add(PayrollAdjust(source_month="2026-07", person_code="A",
                         adjust_amount=-25000, settled_amount=0,
                         remaining=-25000, status="in_progress",
                         updated_at=datetime.utcnow()))
    db.add(PayrollPeriodRow(month="2026-07", person_code="B",
                            diff_points=100, adjust_points=100,
                            diff_amount=100 * 250, adjust_amount=100 * 250,
                            half1_amount=10000, half2_amount=20000,
                            prev_adjust_amount=0,
                            updated_at=datetime.utcnow()))
    for code in ("A", "B"):
        db.add(PersonDailyStat(person_code=code,
                               ref_date=date(2026, 8, 10), points=10))
    db.commit()
    period.sync_period_table(db, "2026-08")
    rows = {r["code"]: r for r in period.period_rows(db, "2026-08")}
    # 8月无对账任务 → 本月金额差=0（2026-09-25 修复，不再全额扣）；
    # prev_amt = 上月结转列（carry-in，来自 7 月行），与本月差异无关。
    assert rows["A"]["prev_amt"] == -25000        # 上月结转输入
    assert rows["B"]["prev_amt"] == 0             # B 无结转
    assert rows["A"]["adj_amt"] == rows["A"]["diff_amt"] == 0  # 无对账=0
    # carry_map（发薪表上月找平列）：上月结转余额（正补/负扣）
    cm = period.carry_map(db, "2026-08")
    assert cm["A"] == [0, -25000] and cm.get("B") is None
    db.close()


def test_adjust_auto_equals_diff_amount(client):
    """找平自动=金额差（无点击操作）：**有对账任务**时 adj_amt==diff_amt。

    无对账任务时差异记 0（不产生"全额扣"幻影差异，2026-09-25 修复）。
    """
    from app.models import ReconTask, ReconDataRow
    db = appdb.SessionLocal()
    db.add(Person(code="111", display_name="甲"))
    db.add(PersonDailyStat(person_code="111", ref_date=date(2026, 8, 1),
                           points=5))
    # 先看无对账：差异 0（旧行为是 −5×250 = 全额扣，已修）
    db.commit()
    period.sync_period_table(db, "2026-08")
    rows = period.period_rows(db, "2026-08")
    assert rows[0]["diff_amt"] == 0 and rows[0]["adj_amt"] == 0

    # 再建对账任务（对账 3 点 vs 系统 5 点 → 差 −2 点 = −500 円）
    db.add(ReconTask(id=1, kind="monthly_v3", status="done", created_by=1,
                     params={"month": "2026-08", "kind": "daily_records",
                             "version": 1, "current": True},
                     summary={}, created_at=datetime.utcnow()))
    db.add(ReconDataRow(task_id=1, ref_date=date(2026, 8, 1),
                        person_code="111", person_name="甲",
                        points=3, cnt=3))
    db.commit()
    period.sync_period_table(db, "2026-08")
    rows = period.period_rows(db, "2026-08")
    assert rows[0]["diff_amt"] == rows[0]["adj_amt"] == -2 * 250  # 负=扣款
    assert rows[0]["diff"] == -2                       # 偏差点数=参考值
    period.sync_period_table(db, "2026-08")         # 重新生成保持
    rows = period.period_rows(db, "2026-08")
    assert rows[0]["adj_amt"] == rows[0]["diff_amt"] == -2 * 250
    db.close()


def test_payroll_settle_page_readonly(client):
    """薪资找平页面只读：找平自动=金额差，无输入框/保存（去掉点击操作）。"""
    from tests_web.test_flow import _seed_admin
    _seed_admin(client)
    db = appdb.SessionLocal()
    db.add(Person(code="111", display_name="甲"))
    db.add(PersonDailyStat(person_code="111", ref_date=date(2026, 8, 1),
                           points=5))
    db.commit()
    period.sync_period_table(db, "2026-08")
    db.close()
    client.post("/login", data={"username": "admin", "password": "pw123456"},
                follow_redirects=False)
    p = client.get("/payroll-settle?month=2026-08").text
    assert "薪资找平" in p and "本月对账偏差(参考)=找平" in p and "生成/更新" in p
    assert "本月对账偏差(参考)" in p
    assert 'name="adjust_delta"' not in p               # 无找平输入框（自动）
    assert "即自动找平金额" in p                        # 自动找平说明
    pp = client.get("/perf?month=2026-08").text
    assert "奖金合计(円)" not in pp and "导出薪资找平 Excel" not in pp
    db.close()


def test_per_point_lock_and_change(client):
    """单价按月可变：设置本月单价只影响本月；已确认找平按当时单价不变。"""
    import app.db as appdb
    from app.models import (Person, PersonDailyStat, AdjustRecord)
    from app.services import period, perf
    db = appdb.SessionLocal()
    from app.models import MonthPerfRecord
    db.add(Person(code="111", display_name="甲"))
    db.add(PersonDailyStat(person_code="111", ref_date=date(2026, 8, 1),
                           points=80))
    # 直插一条月绩效行（绕开 formal 依赖），作为该月单价载体
    db.add(MonthPerfRecord(month="2026-08", person_code="111",
                           records=40, p1=40, p2=0, points=80,
                           salary=perf.salary_for(80), per_point=250))
    db.commit()
    sync1 = period.sync_period_table(db, "2026-08")
    rows = period.period_rows(db, "2026-08")
    r = rows[0]
    # 默认单价 250：下半月 0，上半月 80 点 → 分期=80×250+3000=23000
    assert r["half1_amt"] == 80 * 250 + 3000
    # 8月确认一条找平（锁存当时单价250）
    per_point = perf.month_per_point(db, "2026-08")
    assert per_point == 250
    rec = AdjustRecord(month="2026-08", applied_to_month="2026-09",
                       person_code="111", amount=10, per_point=per_point,
                       source_task_id=0, reason="测试")
    db.add(rec); db.commit()
    # 设置 8 月单价为 300：月绩效工资与薪资找平重算，但找平记录仍 250
    perf.set_month_per_point(db, "2026-08", 300)
    rows = period.period_rows(db, "2026-08")
    assert rows[0]["half1_amt"] == 80 * 300 + 3000   # 按新单价重算
    mp = db.query(MonthPerfRecord).filter(
        MonthPerfRecord.month == "2026-08").first()
    assert mp.per_point == 300 and mp.salary == 80 * 300 + 3000
    rec = db.query(AdjustRecord).first()
    assert rec.per_point == 250 and rec.amount == 10   # 锁存不动
    assert perf.adjust_map(db, "2026-09") == {"111": [10, 10 * 250]}
    db.close()


def test_bonus_configurable(client, monkeypatch):
    """奖金门槛/金额可配置：改 BONUS_GROUP/BONUS_AMOUNT 后工资按新值算。"""
    import app.config as cfg
    from app.services import perf
    monkeypatch.setenv("BONUS_GROUP", "50")
    monkeypatch.setenv("BONUS_AMOUNT", "2000")
    cfg.get_settings.cache_clear()
    try:
        assert perf.salary_for(100) == 100 * 250 + 2 * 2000   # 100//50=2
        g, a = perf.bonus_params()
        assert (g, a) == (50, 2000)
        g9, _ = perf.bonus_params("2026-09")       # 9月起按月规则 75
        assert g9 == 75
        g8, _ = perf.bonus_params("2026-08")       # 8月未命中按月规则 → 50
        assert g8 == 50
    finally:
        monkeypatch.delenv("BONUS_GROUP", raising=False)
        monkeypatch.delenv("BONUS_AMOUNT", raising=False)
        cfg.get_settings.cache_clear()
    assert perf.salary_for(100) == 100 * 250 + 1 * 3000       # 恢复默认 68/3000
    assert perf.salary_for(160, month="2026-09") == 160 * 250 + 2 * 1250  # 9月起门槛75奖1250
    assert perf.salary_for(150, month="2026-08") == 150 * 250 + 2 * 3000  # 8月仍68(150//68=2)


def test_sys_config_chain(client):
    """系统配置全局单值：最新一条(id最大)生效，warm 后计算读取。"""
    import app.db as appdb
    from app.models import SysConfig
    from app.services import perf
    db = appdb.SessionLocal()
    db.add(SysConfig(config_month=None, per_point=250, bonus_group=75,
                     bonus_amount=1250, updated_by=1))
    db.commit()
    perf.clear_config_cache()
    perf.warm_config(db, "2026-09")
    assert perf.bonus_params("2026-09") == (75, 1250)
    assert perf.month_per_point(db, "2026-09") == 250
    assert perf.salary_for(160, 250, "2026-09") == 160 * 250 + 2 * 1250
    # 全局：任何月份都读最新一条
    assert perf.bonus_params("2026-08") == (75, 1250)
    # 更新（新行 id 更大）→ 最新值
    db.add(SysConfig(config_month=None, per_point=300, bonus_group=80,
                     bonus_amount=2000, updated_by=1))
    db.commit()
    perf.clear_config_cache()
    perf.warm_config(db, "2026-09")
    assert perf.bonus_params("2026-09") == (80, 2000)
    assert perf.month_per_point(db, "2026-09") == 300
    perf.clear_config_cache()
    db.close()


def test_record_payment_no_double_deduct_after_seq1(client):
    """8月结转 -500：9月 seq1 扣完 → seq2 不得再扣（防重复扣款）。

    回归：absorbed 是负的已扣额，remaining 应为 carry - absorbed（=0），
    而非 carry + absorbed（=-1000 重复扣）。"""
    db = appdb.SessionLocal()
    db.add(Person(code="111", display_name="甲"))
    db.add(PayrollPeriodRow(month="2026-08", person_code="111",
                            half1_points=0, half2_points=0,
                            settle_points=0, prev_adjust_points=0,
                            diff_points=2, adjust_points=2,
                            diff_amount=500, adjust_amount=500,
                            prev_adjust_amount=-500,
                            updated_at=datetime.utcnow()))
    # 结转的**权威来源是账表**（record_payment 读它，不再读链子）
    db.add(PayrollAdjust(source_month="2026-08", person_code="111",
                         adjust_amount=-500, settled_amount=0, remaining=-500,
                         status="in_progress", updated_at=datetime.utcnow()))
    db.add(PayrollPeriodRow(month="2026-09", person_code="111",
                            half1_points=10, half2_points=10,
                            settle_points=0, prev_adjust_points=0,
                            diff_points=0, adjust_points=0,
                            half1_amount=2500, half2_amount=2500,
                            diff_amount=0, adjust_amount=0,
                            prev_adjust_amount=0,
                            updated_at=datetime.utcnow()))
    db.commit()
    # seq1 发薪：8月结转 -500 → 应扣 -500（工资 2500 足够）
    assert period.record_payment(db, "2026-09", "111", 1) is True
    p1 = (db.query(PayrollPayment)
          .filter_by(month="2026-09", person_code="111", seq=1).one())
    assert p1.adjust_applied == -500
    # seq2 发薪：结转已在 seq1 扣完 → 不得再扣（旧 bug 会扣 -1000）
    assert period.record_payment(db, "2026-09", "111", 2) is True
    p2 = (db.query(PayrollPayment)
          .filter_by(month="2026-09", person_code="111", seq=2).one())
    assert p2.adjust_applied == 0          # 不再重复扣
    assert p2.adjust_leftover is None
    db.close()


def test_record_payment_deduces_remaining_after_partial_seq1(client):
    """8月结转 -3000、seq1 工资 1000 只能扣 1000 → seq2 应扣剩余 -2000。"""
    db = appdb.SessionLocal()
    db.add(Person(code="111", display_name="甲"))
    db.add(PayrollPeriodRow(month="2026-08", person_code="111",
                            half1_points=0, half2_points=0,
                            settle_points=0, prev_adjust_points=0,
                            diff_points=12, adjust_points=12,
                            diff_amount=3000, adjust_amount=3000,
                            prev_adjust_amount=-3000,
                            updated_at=datetime.utcnow()))
    db.add(PayrollAdjust(source_month="2026-08", person_code="111",
                         adjust_amount=-3000, settled_amount=0, remaining=-3000,
                         status="in_progress", updated_at=datetime.utcnow()))
    db.add(PayrollPeriodRow(month="2026-09", person_code="111",
                            half1_points=4, half2_points=20,
                            settle_points=0, prev_adjust_points=0,
                            diff_points=0, adjust_points=0,
                            half1_amount=1000, half2_amount=5000,
                            diff_amount=0, adjust_amount=0,
                            prev_adjust_amount=0,
                            updated_at=datetime.utcnow()))
    db.commit()
    period.record_payment(db, "2026-09", "111", 1)
    p1 = (db.query(PayrollPayment)
          .filter_by(month="2026-09", person_code="111", seq=1).one())
    assert p1.adjust_applied == -1000          # 工资只够扣 1000
    period.record_payment(db, "2026-09", "111", 2)
    p2 = (db.query(PayrollPayment)
          .filter_by(month="2026-09", person_code="111", seq=2).one())
    assert p2.adjust_applied == -2000          # 剩余 -2000 在 seq2 扣
    db.close()


def test_carry_comes_from_ledger_not_stale_chain(client):
    """**回归（2026-10-17 生产）**：8 月那笔已在 9 月发薪时扣掉，10 月不得再扣一次。

    线上实况：账表里 8 月 `remaining=0 / status=settled`（已扣 ¥58,250），
    可 9 月行的链子 `prev_adjust_amount` 还是生成时的旧快照 **−102,500**
    （= 9 月自己 −44,250 + 8 月**已经扣过的** −58,250，全公司合计虚高 ¥592,500）。
    旧实现 `record_payment` 读链子 → 小川逸 10 月上半月 **88,000 工资会被整笔划走**。
    修后读账表 → 只扣 9 月真正还剩的 −44,250。
    """
    db = appdb.SessionLocal()
    db.add(Person(code="222", display_name="小川逸"))
    # 8 月：已结清
    db.add(PayrollAdjust(source_month="2026-08", person_code="222",
                         adjust_amount=-58250, settled_amount=-58250,
                         remaining=0, status="settled",
                         updated_at=datetime.utcnow()))
    # 9 月：还没扣
    db.add(PayrollAdjust(source_month="2026-09", person_code="222",
                         adjust_amount=-44250, settled_amount=0,
                         remaining=-44250, status="in_progress",
                         updated_at=datetime.utcnow()))
    # 9 月行的链子是**旧快照**（含已扣的 8 月）
    db.add(PayrollPeriodRow(month="2026-09", person_code="222",
                            half1_points=1407, half2_points=1370,
                            settle_points=2615, prev_adjust_points=-197,
                            diff_points=-359, adjust_points=0,
                            half1_amount=374250, half2_amount=366250,
                            settle_amount=696250,
                            diff_amount=-44250, adjust_amount=-44250,
                            prev_adjust_amount=-102500,
                            updated_at=datetime.utcnow()))
    db.add(PayrollPeriodRow(month="2026-10", person_code="222",
                            half1_points=332, half2_points=0,
                            settle_points=0, prev_adjust_points=-359,
                            diff_points=-691, adjust_points=0,
                            half1_amount=88000, half2_amount=0,
                            settle_amount=0, diff_amount=0, adjust_amount=0,
                            prev_adjust_amount=0,
                            updated_at=datetime.utcnow()))
    db.commit()

    assert period.record_payment(db, "2026-10", "222", 1) is True
    p1 = (db.query(PayrollPayment)
          .filter_by(month="2026-10", person_code="222", seq=1).one())
    assert p1.adjust_applied == -44250, "只扣 9 月真实剩余的 44,250"
    assert p1.adjust_applied != -88000, "绝不能按旧链子把整月工资扣光"
    # 账表被冲销：9 月结清
    a9 = (db.query(PayrollAdjust)
          .filter_by(source_month="2026-09", person_code="222").one())
    assert a9.remaining == 0 and a9.status == "settled"
    db.close()


def test_carry_map_keeps_amount_already_recovered_this_month(client):
    """🐞 2026-10-17 生产事故回归：**已被本月发放实扣的找平，不能从「应付」里消失**。

    场景照抄王俊岩：9 月偏差 −3,250（账表 `payroll_adjusts`，in_progress）；
    10 月上半月毛额 262,750；10-16 登记发薪（`record_payment` →
    `allocate_settlement` 把那笔冲成 settled / remaining=0）。

    - 修前：`carry_map("2026-10")` 只读 `remaining` → 该人直接消失 →
      绩效页/发薪表「找平 ¥0、应付=毛额 262,750」（虚增）。
    - 修后：已实扣的 −3,250 仍在 → 应付 **259,500**，与台账
      （`payroll_payments.amount + adjust_applied`）一致。

    生产实测（2026-10 全部 56 人）：修前应付合计 5,891,000 / 修后 5,069,250
    = 台账实付，**逐人一致 0 人不符**。
    """
    CODE = "2188240645887825"
    db = appdb.SessionLocal()
    db.add(Person(code=CODE, display_name="王俊岩"))
    # 9 月对账偏差（要扣 3,250）
    db.add(PayrollAdjust(source_month="2026-09", person_code=CODE,
                         adjust_amount=-3250, settled_amount=0, remaining=-3250,
                         status="in_progress", updated_at=datetime.utcnow()))
    # 9 月行（carry 的「余量点数」参考取上月行 diff−adjust）
    db.add(PayrollPeriodRow(month="2026-09", person_code=CODE,
                            diff_points=13, adjust_points=13,
                            diff_amount=-3250, adjust_amount=-3250,
                            half2_amount=40000, updated_at=datetime.utcnow()))
    # 10 月上半月：986 点 × 250 + 奖金 16,250 = 262,750
    db.add(PayrollPeriodRow(month="2026-10", person_code=CODE,
                            half1_points=986, half1_amount=262750,
                            half1_bonus=16250, updated_at=datetime.utcnow()))
    db.commit()

    # ① 登记发薪之前：是"仍未回收的余量"
    assert period.prior_debt_map(db, "2026-10")[CODE] == -3250
    assert period.carry_map(db, "2026-10")[CODE][1] == -3250

    # ② 登记 10 月上半月发薪 → 实扣 3,250，9 月那笔被冲成 settled
    assert period.record_payment(db, "2026-10", CODE, seq=1,
                                 amount=262750, points=986, bonus=16250) is True
    pay = db.query(PayrollPayment).filter(
        PayrollPayment.month == "2026-10",
        PayrollPayment.person_code == CODE).one()
    assert pay.adjust_applied == -3250
    assert db.query(PayrollAdjust).filter(
        PayrollAdjust.source_month == "2026-09").one().status == "settled"

    # ③ 关键：扣完之后"余量"确实是 0，但**这个月仍必须显示 −3,250**
    assert period.prior_debt_map(db, "2026-10").get(CODE, 0) == 0
    cm = period.carry_map(db, "2026-10")
    assert cm.get(CODE) is not None, \
        "已被本月发放实扣的找平从 carry_map 消失了 → 应付虚增回毛额（本次事故）"
    assert cm[CODE][1] == -3250, "应显示本月实际扣掉的 3,250，而不是 0"
    assert cm[CODE][0] == 0                       # 余量点数只作参考

    # ④ 应付 = 毛额 + 找平 = 台账实付（绩效页/发薪表就是 hamt + carry）
    assert 262750 + cm[CODE][1] == 259500 == pay.amount + pay.adjust_applied
    db.close()


def test_carry_map_does_not_double_count_partial_recovery(client):
    """部分回收：已实扣 + 未回收余量 = 原始偏差，**不重复计**（文伊琪型）。

    9 月偏差 −294,000；10 月上半月毛额只有 106,500，被整笔划走 → remaining=−187,500。
    carry 必须是 **−294,000**（= 已扣 106,500 + 余 187,500），
    这样 106,500 + (−294,000) < 0 → 该期应付 0、下期继续扣 187,500
    （与台账 adjust_leftover 一致）。
    """
    CODE = "P1"
    db = appdb.SessionLocal()
    db.add(Person(code=CODE, display_name="文伊琪"))
    db.add(PayrollAdjust(source_month="2026-09", person_code=CODE,
                         adjust_amount=-294000, settled_amount=0,
                         remaining=-294000, status="in_progress",
                         updated_at=datetime.utcnow()))
    db.add(PayrollPeriodRow(month="2026-09", person_code=CODE,
                            diff_points=0, adjust_points=0,
                            diff_amount=-294000, adjust_amount=-294000,
                            updated_at=datetime.utcnow()))
    db.add(PayrollPeriodRow(month="2026-10", person_code=CODE,
                            half1_points=426, half1_amount=106500,
                            updated_at=datetime.utcnow()))
    db.commit()

    assert period.record_payment(db, "2026-10", CODE, seq=1,
                                 amount=106500, points=426) is True
    pay = db.query(PayrollPayment).filter(
        PayrollPayment.month == "2026-10",
        PayrollPayment.person_code == CODE).one()
    assert pay.adjust_applied == -106500
    assert pay.adjust_leftover == -187500          # 递延到下半月
    assert period.carry_map(db, "2026-10")[CODE][1] == -294000   # 不重复也不漏
    # 绩效页/发薪表的算法：毛额 + carry → 负数则该期发 0、余数结转
    p1 = 106500 + (-294000)
    assert p1 < 0 and 0 == max(p1, 0)
    assert -p1 == 187500 == -pay.adjust_leftover
    db.close()


def test_pay_plan_is_the_single_rule_for_both_halves(client):
    """两期发薪规则**唯一实现**：上半月不够扣 → 发 0、余数转下半月。

    绩效页与发薪表导出共用 `period.pay_plan`（2026-10-17 前各写一套，
    导出还**不夹取** → 表里会出现负工资；且下半月一律不显示找平 → 余数凭空消失）。
    """
    db = appdb.SessionLocal()
    # 甲：小额找平（王俊岩型）——上半月毛额够扣
    # 乙：大额找平（文伊琪型）——上半月被整笔划走，余数转下半月
    db.add_all([Person(code="A", display_name="甲"), Person(code="B", display_name="乙")])
    db.add_all([
        PayrollAdjust(source_month="2026-09", person_code="A", adjust_amount=-3250,
                      settled_amount=0, remaining=-3250, status="in_progress",
                      updated_at=datetime.utcnow()),
        PayrollAdjust(source_month="2026-09", person_code="B",
                      adjust_amount=-294000, settled_amount=0,
                      remaining=-294000, status="in_progress",
                      updated_at=datetime.utcnow()),
    ])
    db.commit()
    half_map = {"A": (0, 262750, 16250, 0, 50000, 0),
                "B": (0, 106500, 0, 0, 200000, 0)}
    plan = period.pay_plan(db, "2026-10", half_map)

    # 甲：扣 3,250 → 259,500；下半月没有余数
    assert plan["A"]["h1"] == (-3250, 259500)
    assert plan["A"]["h2"][0] == 0 and plan["A"]["h2"][1] == 50000
    # 乙：找平 294,000 > 毛额 106,500 → 上半月发 0（**不是负数**），
    #     余 187,500 转下半月 → 下半月应付 200,000 − 187,500 = 12,500
    assert plan["B"]["h1"] == (-294000, 0)
    assert plan["B"]["h2"] == (-187500, 12500)
    # 铁律：两期应付都不许为负（导出原来会把负数直接落表）
    for c, v in plan.items():
        for k in ("h1", "h2"):
            assert v[k][1] >= 0, f"{c}.{k} 出现负工资 {v[k][1]}"

    # 登记上半月发薪后，下半月的找平必须等于台账剩余（record_payment 读的那个数）
    assert period.record_payment(db, "2026-10", "B", seq=1,
                                 amount=106500, points=0) is True
    b = db.query(PayrollPayment).filter(
        PayrollPayment.month == "2026-10", PayrollPayment.person_code == "B").one()
    assert b.adjust_applied == -106500 and b.adjust_leftover == -187500
    plan2 = period.pay_plan(db, "2026-10", half_map)
    assert plan2["B"]["h2"][0] == b.adjust_leftover == -187500
    assert plan2["B"]["h2"][0] == period.prior_debt_map(db, "2026-10")["B"]
    db.close()


def test_payroll_export_matches_ledger(client):
    """发薪表导出（**导出发薪表即登记**，运营照它付钱）里的「找平/应付」必须等于台账。

    2026-10-17 事故前这张表有三个毛病：① 找平只读 `remaining` → 已实扣的消失（应付虚增）；
    ② 不做夹取，`毛额+找平` 直接落表 → **出现负工资**；③ 下半月一律不显示找平
    → 上半月扣不完的余数凭空消失，可 `record_payment` 照扣（表与台账又不一致）。
    """
    from io import BytesIO

    from openpyxl import load_workbook

    from app.models import MonthPerfRecord
    from app.services import report

    db = appdb.SessionLocal()
    db.add_all([Person(code="A", display_name="甲"),
                Person(code="B", display_name="乙")])
    for code, salary in (("A", 262750), ("B", 206500)):
        db.add(MonthPerfRecord(month="2026-10", person_code=code, records=1,
                               p1=1, p2=0, points=1, salary=salary,
                               per_point=250))
    db.add_all([
        # 甲（王俊岩型）：下半月 50,000
        PayrollPeriodRow(month="2026-10", person_code="A",
                         half1_points=986, half2_points=0,
                         half1_records=1, half1_p1=1, half1_p2=0,
                         half2_records=0, half2_p1=0, half2_p2=0,
                         half1_bonus=16250, half2_bonus=0,
                         half1_amount=262750, half2_amount=50000,
                         settle_amount=0, prev_adjust_points=0,
                         prev_adjust_amount=0, diff_points=0, diff_amount=0,
                         adjust_points=0, adjust_amount=0),
        # 乙（文伊琪型）：下半月 200,000，上半月不够扣
        PayrollPeriodRow(month="2026-10", person_code="B",
                         half1_points=426, half2_points=0,
                         half1_records=1, half1_p1=1, half1_p2=0,
                         half2_records=0, half2_p1=0, half2_p2=0,
                         half1_bonus=0, half2_bonus=0,
                         half1_amount=106500, half2_amount=200000,
                         settle_amount=0, prev_adjust_points=0,
                         prev_adjust_amount=0, diff_points=0, diff_amount=0,
                         adjust_points=0, adjust_amount=0),
    ])
    db.add_all([
        PayrollAdjust(source_month="2026-09", person_code="A",
                      adjust_amount=-3250, settled_amount=0, remaining=-3250,
                      status="in_progress", updated_at=datetime.utcnow()),
        PayrollAdjust(source_month="2026-09", person_code="B",
                      adjust_amount=-294000, settled_amount=0,
                      remaining=-294000, status="in_progress",
                      updated_at=datetime.utcnow()),
    ])
    db.commit()
    # 运营点导出那一刻 = 登记上半月发薪（台账在此刻落定）
    assert period.record_payment(db, "2026-10", "A", seq=1,
                                 amount=262750, points=986, bonus=16250)
    assert period.record_payment(db, "2026-10", "B", seq=1,
                                 amount=106500, points=426)
    db.commit()

    def sheet(period_):
        data, _ = report.build_payroll_workbook(db, "2026-10", period_)
        return [list(r) for r in
                load_workbook(BytesIO(data)).active.iter_rows(values_only=True)]

    def row_of(rows, code):
        for r in rows:
            if r and r[0] == code:
                return r
        raise AssertionError(f"表里没有 {code}")

    h1 = sheet("half1")
    a, b = row_of(h1, "A"), row_of(h1, "B")
    # 甲：原本已实扣 3,250 → 表里必须显示 −3,250、应付 259,500（修前是 0 / 262,750）
    assert a[7] == 262750 and a[8] == -3250 and a[9] == 259500
    # 乙：找平 294,000 > 毛额 106,500 → 应付 0（**不是 −187,500**）
    assert b[7] == 106500 and b[8] == -294000 and b[9] == 0
    # 铁律：整张表不许出现负工资
    for r in h1:
        if r and isinstance(r[0], str) and r[0] not in ("合计",) and r[9] is not None:
            assert r[9] is None or not isinstance(r[9], (int, float)) or r[9] >= 0

    h2 = sheet("half2")
    a2, b2 = row_of(h2, "A"), row_of(h2, "B")
    assert a2[8] == 0 and a2[9] == 50000                    # 甲：余数 0
    assert b2[8] == -187500 and b2[9] == 12500              # 乙：余数转下半月
    # 与台账一致：下半月显示的找平 = record_payment 记的 adjust_leftover
    pay_b = db.query(PayrollPayment).filter(
        PayrollPayment.month == "2026-10", PayrollPayment.person_code == "B").one()
    assert b2[8] == pay_b.adjust_leftover == -187500
    db.close()
