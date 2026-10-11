# -*- coding: utf-8 -*-
"""AI 表头解析（对账文件表头自适应，模型优先、规则兜底）。"""
import pytest

import app.db as appdb
from app.models import Person
from app.services import recon
from tests.helpers import write_workbook

ALIPAY_HDR = ["Statement Date", "ISO PID", "ISO Name", "Store Basic ID",
              "Shop Name", "Agent Name", "Action Type", "Action Type",
              "Store Total Amount"]


def test_ai_parse_returns_empty_when_unconfigured(monkeypatch, tmp_path):
    """未配置模型 → ai_parse_daily_cols 返回 {}（回退规则，不崩）。"""
    monkeypatch.setattr(recon, "ai_configured", lambda: False)
    p = str(tmp_path / "x.xlsx")
    write_workbook(p, [("s", [ALIPAY_HDR], [["2026-08-01", "1", "n", "S1",
                                             "店1", "甲", "POSM", "PMS", "1"]])])
    assert recon.ai_parse_daily_cols(p) == {}


def test_ai_parse_daily_cols_and_parse(client, monkeypatch, tmp_path):
    """模型返回列映射 → ai_parse_daily_cols 给出映射；parse 按 AI 列解析成功。"""
    fake = ('{"date":0,"store_id":3,"store_name":4,"person_name":5,'
            '"person_code":null,"points":8,"visible":null,"deploy":null}')
    monkeypatch.setattr(recon, "ai_configured", lambda: True)
    monkeypatch.setattr(recon, "_chat", lambda prompt: fake)
    p = str(tmp_path / "alipay.xlsx")
    write_workbook(p, [("ISO Billing Export", [ALIPAY_HDR], [
        ["2026-08-01", "1", "MarsNavi", "S1", "店1", "甲", "POSM", "PMS", "1"],
        ["2026-08-01", "1", "MarsNavi", "S2", "店2", "乙", "POSM", "PMS", "2"],
        ["2026-08-02", "1", "MarsNavi", "S3", "店3", "甲", "POSM", "PMS", "2"],
    ])])
    cols = recon.ai_parse_daily_cols(p)
    assert cols.get("date") == 0
    assert cols.get("store_id") == 3
    assert cols.get("person_name") == 5
    assert cols.get("points") == 8
    # 建 Person 供姓名匹配
    db = appdb.SessionLocal()
    db.add(Person(code="111", display_name="甲"))
    db.add(Person(code="222", display_name="乙"))
    db.commit()
    parsed = recon.parse_daily_records(p, "2026-08", db=db, ai_cols=cols)
    db.close()
    assert parsed is not None
    total_cnt = sum(dd["cnt"] for v in parsed.values()
                    for dd in v["days"].values())
    total_pts = sum(dd["pts"] for v in parsed.values()
                    for dd in v["days"].values())
    assert total_cnt == 3 and total_pts == 5   # 1 + 2 + 2
    assert set(parsed) == {"111", "222"}


def test_ai_parse_invalid_mapping_returns_empty(monkeypatch, tmp_path):
    """模型返回不可信映射（越界/缺店）→ {} 回退规则。"""
    monkeypatch.setattr(recon, "ai_configured", lambda: True)
    monkeypatch.setattr(recon, "_chat",
                        lambda prompt: '{"date":0,"store_id":99,"points":8}')
    p = str(tmp_path / "y.xlsx")
    write_workbook(p, [("s", [ALIPAY_HDR], [["2026-08-01", "1", "n", "S1",
                                             "店1", "甲", "POSM", "PMS", "1"]])])
    assert recon.ai_parse_daily_cols(p) == {}


def test_interpret_prompt_contains_key_data(client, tmp_path, monkeypatch):
    """对账分析提示词包含关键事实（归因/反向名单/特殊标记），不依赖外部模型。"""
    from app.models import Person, ReconTask
    from app.services import recon
    import app.db as appdb
    _seed = appdb.SessionLocal()
    _seed.add(Person(code="111", display_name="甲"))
    _seed.commit()
    db = appdb.SessionLocal()
    # 手工造一个任务并塞 summary/params（走 build_interpret_prompt 只读）
    t = ReconTask(kind="monthly_v3", status="done", created_by=1,
                  params={"month": "2026-08", "file": "测试.xlsx",
                          "parser": "rules", "single_date": True,
                          "single_date_val": "2026-08-01",
                          "unmatched": [["未知员工", "S9", "2026-08-01"]]},
                  summary={"compared": 34, "diff_count": 3,
                           "monthly_diff": 2, "sys_only": ["111"],
                           "attribution": {"consistent": 100,
                                           "only_system": 5,
                                           "only_report": 2,
                                           "only_report_reason": {"from_sub": 2},
                                           "net": 3}})
    db.add(t)
    db.commit()
    prompt = recon.build_interpret_prompt(db, t.id)
    for key in ("2026-08", "测试.xlsx", "34", "仅系统有 5", "仅对账有 2",
                "反向名单", "账单日文件", "rules", "未知员工",
                "总体结论", "差异构成", "风险与需人工核对", "建议动作",
                "350字"):
        assert key in prompt, f"prompt 缺少: {key}"
    db.delete(t)
    db.commit()
    db.close()
    _seed.close()


def test_unmatched_collected_in_parse(client, tmp_path):
    """对账文件含未匹配员工 → parse 收集到 unmatched（供页面/分析提示）。"""
    import app.db as appdb
    from app.models import Person
    from app.services import recon
    from tests.helpers import write_workbook
    db = appdb.SessionLocal()
    db.add(Person(code="111", display_name="甲"))
    db.commit()
    hdr = ["Statement Date", "Store Basic ID", "Shop Name", "Agent Name",
           "Store Total Amount"]
    p = str(tmp_path / "u.xlsx")
    write_workbook(p, [("s", [hdr], [
        ["2026-08-01", "S1", "店1", "甲", "1"],
        ["2026-08-01", "S2", "店2", "不认识的人", "2"],
        ["2026-08-01", "S3", "店3", "甲", "2"],
    ])])
    um = []
    parsed = recon.parse_daily_records(p, "2026-08", db=db, unmatched=um)
    db.close()
    assert parsed is not None
    assert ("不认识的人", "S2", "2026-08-01") in um
    # 未匹配行未计入
    total = sum(dd["cnt"] for v in parsed.values() for dd in v["days"].values())
    assert total == 2


def test_report_ai_notes_in_explain_column(client, monkeypatch):
    """对账报告 Excel：员工×日与差异明细的说明列写入 AI 解读。"""
    import io
    from openpyxl import load_workbook
    from app.models import (Person, ReconDayRow, ReconResult, ReconTask)
    from app.services import recon
    import app.db as appdb
    db = appdb.SessionLocal()
    db.add(Person(code="111", display_name="甲"))
    db.commit()
    t = ReconTask(kind="monthly_v3", status="done", created_by=1,
                  params={"month": "2026-08", "file": "t.xlsx",
                          "single_date": True, "single_date_val": "2026-08-01"},
                  summary={"compared": 2, "diff_count": 1, "monthly_diff": 1})
    db.add(t)
    db.commit()
    db.add(ReconDayRow(task_id=t.id, ref_date=__import__("datetime").date(2026, 8, 1),
                       person_code="111", sys_points=28, rep_points=2129,
                       diff=-2101, side="both"))
    db.add(ReconResult(task_id=t.id, submitter_code="111",
                       status="diff", family="month",
                       system_value=28, report_value=2129, diff=-2101,
                       note="甲 点数差异"))
    db.commit()
    monkeypatch.setattr(recon, "ai_configured", lambda: True)
    monkeypatch.setattr(recon, "_chat",
                        lambda prompt: '{"111": "账单日聚合导致日级差异大，人月应以对账全月点数为准"}')
    wb = recon.build_report(db, t.id, "管理员")
    # 员工×日 说明列
    ws = wb["员工×日对账明细"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[1][6] == "点数不一致｜AI:账单日聚合导致日级差异大，人月应以对账全月点数为准"
    # 差异明细 AI说明列
    ws2 = wb["差异明细"]
    r2 = list(ws2.iter_rows(values_only=True))
    assert r2[1][6] == "账单日聚合导致日级差异大，人月应以对账全月点数为准"
    db.delete(t)
    db.commit()
    db.close()


def test_report_ai_notes_unconfigured(client, monkeypatch):
    """未配置模型 → 报告说明列保持原文本，不崩。"""
    import io
    from openpyxl import load_workbook
    from app.models import (Person, ReconDayRow, ReconTask)
    from app.services import recon
    import app.db as appdb
    db = appdb.SessionLocal()
    db.add(Person(code="111", display_name="甲"))
    db.commit()
    t = ReconTask(kind="monthly_v3", status="done", created_by=1,
                  params={"month": "2026-08", "file": "t.xlsx"},
                  summary={"compared": 1, "diff_count": 1})
    db.add(t)
    db.commit()
    db.add(ReconDayRow(task_id=t.id, ref_date=__import__("datetime").date(2026, 8, 1),
                       person_code="111", sys_points=28, rep_points=29,
                       diff=-1, side="both"))
    db.commit()
    monkeypatch.setattr(recon, "ai_configured", lambda: False)
    wb = recon.build_report(db, t.id, "管理员")
    ws = wb["员工×日对账明细"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[1][6] == "点数不一致"   # 无 AI 前缀
    db.delete(t)
    db.commit()
    db.close()


# ---------------- store_list（阿里结算月度账单）解析 ----------------
STORE_LIST_HDR = ["Settle Month", "ISO PID", "ISO Name", "ISO Email",
                  "ISO Country", "Store ID", "Store Name", "Store Country",
                  "Store Address", "CD Name", "Employee Name", "Action Type",
                  "Store Total Amount"]


def test_guess_kind_store_list(tmp_path):
    """阿里结算账单表头 → store_list（店+员工+点数直给+结算月）。"""
    from app.services import recon
    from tests.helpers import write_workbook
    p = str(tmp_path / "alipay.xlsx")
    write_workbook(p, [("MarsNavi", [STORE_LIST_HDR], [
        ["2026-09", "1", "n", "e", "JP", "S1", "店1", "JP", "addr", "CD",
         "甲", "POSM", "1"],
    ]), ("规则", [["计算规则", "M列，1为1点店铺，2为2点店铺"]], [])])
    assert recon.guess_kind(p) == "store_list"


def test_parse_store_list(client, tmp_path):
    """store_list 解析：按 员工姓名 汇总人月点数（点数列 1/2 直给）。"""
    import app.db as appdb
    from app.models import Person
    from app.services import recon
    from tests.helpers import write_workbook
    db = appdb.SessionLocal()
    db.add(Person(code="111", display_name="甲"))
    db.add(Person(code="222", display_name="乙"))
    db.commit()
    p = str(tmp_path / "sl.xlsx")
    write_workbook(p, [("MarsNavi", [STORE_LIST_HDR], [
        ["2026-09", "1", "n", "e", "JP", "S1", "店1", "JP", "a", "CD", "甲", "POSM", "1"],
        ["2026-09", "1", "n", "e", "JP", "S2", "店2", "JP", "a", "CD", "甲", "POSM", "2"],
        ["2026-09", "1", "n", "e", "JP", "S3", "店3", "JP", "a", "CD", "乙", "POSM", "2"],
        ["2026-08", "1", "n", "e", "JP", "S4", "店4", "JP", "a", "CD", "乙", "POSM", "1"],
    ])])
    unmatched = []
    out = recon.parse_store_list(p, "2026-09", db=db, unmatched=unmatched)
    assert out == {"111": {"name": "甲", "pts": 3, "cnt": 2},
                   "222": {"name": "乙", "pts": 2, "cnt": 1}}
    # 不在结算月的行被过滤；未匹配员工收集
    p2 = str(tmp_path / "sl2.xlsx")
    write_workbook(p2, [("MarsNavi", [STORE_LIST_HDR], [
        ["2026-09", "1", "n", "e", "JP", "S9", "店9", "JP", "a", "CD", "陌生人", "POSM", "1"],
    ])])
    unmatched2 = []
    out2 = recon.parse_store_list(p2, "2026-09", db=db, unmatched=unmatched2)
    assert out2 is None and unmatched2 == [("陌生人", "", "2026-09")]
    db.close()


def test_run_task_store_list_full(client, tmp_path, monkeypatch):
    """完整链路：store_list 上传 → 解析 → recon_data_rows 落库 → 人月差异。"""
    import app.db as appdb
    from app.models import MonthPerfRecord, Person, ReconDataRow, ReconTask
    from app.services import recon
    from tests.helpers import write_workbook
    db = appdb.SessionLocal()
    db.add(Person(code="111", display_name="甲"))
    db.add(Person(code="222", display_name="乙"))
    db.add(MonthPerfRecord(month="2026-09", person_code="111", points=10))
    db.commit()
    p = str(tmp_path / "sl3.xlsx")
    write_workbook(p, [("MarsNavi", [STORE_LIST_HDR], [
        ["2026-09", "1", "n", "e", "JP", "S1", "店1", "JP", "a", "CD", "甲", "POSM", "1"],
        ["2026-09", "1", "n", "e", "JP", "S2", "店2", "JP", "a", "CD", "甲", "POSM", "2"],
        ["2026-09", "1", "n", "e", "JP", "S3", "店3", "JP", "a", "CD", "乙", "POSM", "1"],
    ])])
    monkeypatch.setattr(recon, "ai_configured", lambda: False)
    t, older = recon.submit_task(db, "2026-09", "sl3.xlsx",
                                 open(p, "rb").read(), 1, sync=True)
    rows = db.query(ReconDataRow).filter(ReconDataRow.task_id == t.id).all()
    assert len(rows) == 2
    by = {r.person_code: r for r in rows}
    assert by["111"].points == 3 and by["222"].points == 1
    res = db.query(recon.ReconResult).filter(
        recon.ReconResult.task_id == t.id).all()
    # 甲：系统 10 点 vs 对账 3 点 → 差 +7；乙：系统 0 点 vs 对账 1 点 → 差 -1
    m = {r.submitter_code: r.diff for r in res}
    assert m.get("111") == 7 and m.get("222") == -1
    assert (t.summary or {}).get("kind") == "store_list"
    db.close()
